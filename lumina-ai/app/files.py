"""✦ LUMINA AI — Đọc nội dung tệp đính kèm (PDF / Word / Excel / PPTX / HTML / văn bản).

Người dùng gửi tệp → tách chữ ra → đưa vào ngữ cảnh cho bộ não đọc. Không lưu
tệp gốc lại (chỉ dùng cho lượt chat hiện tại), giống cách xử lý ảnh ở media.py.

Bộ đọc 2 tầng:
  1. markitdown (TÙY CHỌN) — nếu đã `pip install markitdown`, dùng nó để ra Markdown
     sạch (giữ tiêu đề/bảng/danh sách) cho MỌI định dạng, kể cả .pptx/.html/.epub.
     Ta truyền sẵn extension+mimetype nên nó KHÔNG cần magika lúc chạy.
  2. Fallback nhẹ — pypdf/python-docx/openpyxl (đã có sẵn, ~vài trăm KB). App mặc
     định vẫn nhẹ; muốn đọc xịn hơn thì tự cài markitdown (kéo ~35MB onnxruntime).
"""

import base64
import io
import os
import re

_DATA_URL_RE = re.compile(r"^data:(?P<mt>[\w./+-]+);base64,(?P<data>.+)$", re.DOTALL)

# Giới hạn ký tự đưa vào ngữ cảnh — tệp dài sẽ bị cắt bớt (tránh tốn quá nhiều token).
MAX_CHARS_PER_FILE = 15000
MAX_FILES = 3


def _decode(data_url: str) -> bytes | None:
    m = _DATA_URL_RE.match((data_url or "").strip())
    if not m:
        return None
    try:
        return base64.b64decode(m.group("data"))
    except Exception:  # noqa: BLE001
        return None


def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw))
    parts = []
    for page in reader.pages[:60]:  # chặn PDF quá dài (sách trăm trang) làm chậm
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _extract_docx(raw: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(raw))
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_xlsx(raw: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets[:10]:
        lines.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows(max_row=500, values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_text(raw: bytes) -> str:
    return raw.decode("utf-8", errors="ignore")


_EXTRACTORS = {
    "application/pdf": _extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _extract_xlsx,
    "text/plain": _extract_text,
    "text/csv": _extract_text,
    "text/markdown": _extract_text,
}

# Đuôi file → mimetype (để nhận diện khi data URL chỉ ghi "application/octet-stream").
_EXT_MIME = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".html": "text/html", ".htm": "text/html", ".epub": "application/epub+zip",
    ".txt": "text/plain", ".md": "text/markdown", ".csv": "text/csv", ".json": "application/json",
}
# Định dạng CHỈ markitdown đọc được (fallback nhẹ không có) — .pptx/.html/.epub.
_MARKITDOWN_ONLY = {
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "text/html", "application/epub+zip",
}

# Cho phép TẮT markitdown kể cả khi đã cài (MARKITDOWN_DISABLE=1) — deploy siêu nhẹ.
_MD_DISABLED = os.getenv("MARKITDOWN_DISABLE", "").strip() in ("1", "true", "yes")
_MD_INSTANCE = None
_MD_TRIED = False


def _markitdown():
    """Trả instance MarkItDown nếu đã cài, else None (lazy, chỉ import 1 lần)."""
    global _MD_INSTANCE, _MD_TRIED
    if _MD_TRIED:
        return _MD_INSTANCE
    _MD_TRIED = True
    if _MD_DISABLED:
        return None
    try:
        from markitdown import MarkItDown
        _MD_INSTANCE = MarkItDown(enable_plugins=False)
    except Exception:  # noqa: BLE001 — chưa cài markitdown → dùng fallback nhẹ
        _MD_INSTANCE = None
    return _MD_INSTANCE


def _extract_markitdown(raw: bytes, mimetype: str, ext: str) -> str:
    """Dùng markitdown → Markdown sạch. Truyền sẵn extension+mimetype nên không gọi magika."""
    md = _markitdown()
    if md is None:
        return ""
    from markitdown._stream_info import StreamInfo
    info = StreamInfo(extension=ext or None, mimetype=mimetype or None)
    return (md.convert_stream(io.BytesIO(raw), stream_info=info).text_content or "").strip()


def extract_text(name: str, data_url: str) -> dict:
    """Trả về {name, text, error}. Không bao giờ raise — lỗi thì trả error, chat vẫn tiếp tục."""
    m = _DATA_URL_RE.match((data_url or "").strip())
    if not m:
        return {"name": name, "text": "", "error": "Định dạng tệp không hợp lệ."}
    media_type = m.group("mt").lower()
    raw = _decode(data_url)
    if raw is None:
        return {"name": name, "text": "", "error": "Không đọc được tệp."}

    ext = os.path.splitext(name.lower())[1]
    # data URL nhiều khi ghi octet-stream → suy mimetype từ đuôi file.
    if media_type in ("application/octet-stream", "") and ext in _EXT_MIME:
        media_type = _EXT_MIME[ext]

    text = ""
    md_supported = media_type in _EXTRACTORS or media_type in _MARKITDOWN_ONLY or ext in _EXT_MIME
    # Tầng 1: markitdown (nếu có) — ra Markdown sạch cho mọi định dạng trên.
    if md_supported and _markitdown() is not None:
        try:
            text = _extract_markitdown(raw, media_type, ext)
        except Exception:  # noqa: BLE001 — markitdown lỗi → thử fallback nhẹ bên dưới
            text = ""

    # Tầng 2: fallback nhẹ (pypdf/docx/xlsx/text) — khi chưa cài markitdown hoặc nó fail.
    if not text:
        extractor = _EXTRACTORS.get(media_type)
        if not extractor and ext in (".txt", ".md", ".csv", ".json"):
            extractor = _extract_text
        if not extractor:
            hint = " — cài `markitdown` để đọc .pptx/.html/.epub" if media_type in _MARKITDOWN_ONLY else ""
            return {"name": name, "text": "",
                    "error": f"LUMINA chưa đọc được định dạng này ({media_type or 'không rõ'}){hint}."}
        try:
            text = extractor(raw).strip()
        except Exception as exc:  # noqa: BLE001 — tệp hỏng/định dạng lạ không được làm sập chat
            return {"name": name, "text": "", "error": f"Không đọc được nội dung tệp: {exc}"}

    if not text:
        return {"name": name, "text": "", "error": "Tệp rỗng hoặc không trích được chữ (có thể là ảnh scan)."}

    truncated = len(text) > MAX_CHARS_PER_FILE
    return {"name": name, "text": text[:MAX_CHARS_PER_FILE], "error": "", "truncated": truncated}


def build_context(files: list[dict]) -> str:
    """Ghép nội dung các tệp đã đọc thành khối ngữ cảnh cho bộ não."""
    blocks = []
    for f in files:
        if f.get("error"):
            continue
        note = " (đã cắt bớt vì quá dài)" if f.get("truncated") else ""
        blocks.append(f"--- Tệp: {f['name']}{note} ---\n{f['text']}")
    if not blocks:
        return ""
    return "\n\n[NỘI DUNG TỆP NGƯỜI DÙNG GỬI:\n" + "\n\n".join(blocks) + "\n]"
