"""✦ LUMINA AI — Sinh FILE tài liệu từ Markdown (docx / pdf / html).

Bổ sung cho files.py (chỉ ĐỌC): module này SINH ra 1 file mới để tải về. Dùng khi
bộ não đã đọc + tổng hợp nhiều tệp thành một bản Markdown có cấu trúc (một "cuốn
sách") — ta chuyển bản Markdown đó thành file thật cho người dùng tải.

- docx  : python-docx (đã có sẵn) — hỗ trợ MỌI ngôn ngữ (Việt, Trung, Nhật...), sửa được.
- pdf   : fpdf2 + font DejaVu — tốt cho chữ Latin/Việt; chữ CJK (Hán/Nhật) có thể
          không hiện (font DejaVu không phủ) → nên chọn docx/html cho nội dung CJK.
- html  : trang tự chứa, mở trên điện thoại rồi "Lưu thành PDF" — phủ mọi ngôn ngữ.

Markdown hỗ trợ ở mức thực dụng: # / ## / ### tiêu đề, - hoặc * gạch đầu dòng,
> trích dẫn, --- ngăn cách, đoạn văn thường, **đậm** (chỉ lược bỏ dấu ở PDF/docx).
"""

from __future__ import annotations

import html as _html
import io
import re

# Font Unicode có sẵn trên hệ thống (phủ Latin mở rộng + tiếng Việt).
_DEJAVU = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
_DEJAVU_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_ORDERED_RE = re.compile(r"^\d+[.)]\s+(.*)$")
_QUOTE_RE = re.compile(r"^>\s?(.*)$")


def _strip_inline(text: str) -> str:
    """Bỏ ký hiệu inline (**đậm**, *nghiêng*, `code`) cho đầu ra không phải HTML."""
    text = _BOLD_RE.sub(r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*", r"\1", text)  # *nghiêng*
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text


class _Block:
    __slots__ = ("kind", "level", "text")

    def __init__(self, kind: str, text: str, level: int = 0):
        self.kind = kind      # h | p | bullet | ordered | quote | hr
        self.text = text
        self.level = level


def parse_markdown(md: str) -> list[_Block]:
    """Bổ Markdown thành danh sách khối đơn giản (đủ cho tài liệu văn xuôi)."""
    blocks: list[_Block] = []
    para: list[str] = []

    def flush_para():
        if para:
            blocks.append(_Block("p", " ".join(para).strip()))
            para.clear()

    for raw in (md or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush_para()
            continue
        if line.strip() in ("---", "***", "___"):
            flush_para()
            blocks.append(_Block("hr", ""))
            continue
        m = _HEADING_RE.match(line)
        if m:
            flush_para()
            blocks.append(_Block("h", m.group(2).strip(), level=len(m.group(1))))
            continue
        m = _BULLET_RE.match(line)
        if m:
            flush_para()
            blocks.append(_Block("bullet", m.group(1).strip()))
            continue
        m = _ORDERED_RE.match(line)
        if m:
            flush_para()
            blocks.append(_Block("ordered", m.group(1).strip()))
            continue
        m = _QUOTE_RE.match(line)
        if m:
            flush_para()
            blocks.append(_Block("quote", m.group(1).strip()))
            continue
        para.append(line.strip())
    flush_para()
    return blocks


# ─────────────────────────── DOCX ───────────────────────────

def to_docx(title: str, md: str) -> bytes:
    """Markdown → .docx (python-docx). Hỗ trợ mọi ngôn ngữ, mở/sửa bằng Word/Docs."""
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    if title:
        doc.add_heading(title, level=0)

    for b in parse_markdown(md):
        if b.kind == "h":
            doc.add_heading(_strip_inline(b.text), level=min(max(b.level, 1), 4))
        elif b.kind == "bullet":
            doc.add_paragraph(_strip_inline(b.text), style="List Bullet")
        elif b.kind == "ordered":
            doc.add_paragraph(_strip_inline(b.text), style="List Number")
        elif b.kind == "quote":
            p = doc.add_paragraph(_strip_inline(b.text))
            p.paragraph_format.left_indent = Pt(24)
            if p.runs:
                p.runs[0].italic = True
        elif b.kind == "hr":
            doc.add_paragraph("─" * 20)
        elif b.kind == "p":
            doc.add_paragraph(_strip_inline(b.text))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─────────────────────────── PDF ───────────────────────────

def to_pdf(title: str, md: str) -> bytes:
    """Markdown → PDF (fpdf2 + DejaVu). Tốt cho Latin/Việt; CJK có thể không hiện."""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    def cell(h, txt):
        pdf.multi_cell(0, h, txt, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(20, 18, 20)
    pdf.add_page()

    import os
    have_font = os.path.exists(_DEJAVU)
    if have_font:
        pdf.add_font("DejaVu", "", _DEJAVU)
        if os.path.exists(_DEJAVU_BOLD):
            pdf.add_font("DejaVu", "B", _DEJAVU_BOLD)
        base = "DejaVu"
    else:  # dự phòng cực hiếm — font lõi (không dấu tiếng Việt)
        base = "Helvetica"

    def setf(style: str = "", size: int = 12):
        try:
            pdf.set_font(base, style, size)
        except Exception:  # noqa: BLE001 — style B thiếu file → về thường
            pdf.set_font(base, "", size)

    if title:
        setf("B", 20)
        cell(10, _strip_inline(title))
        pdf.ln(3)

    for b in parse_markdown(md):
        if b.kind == "h":
            pdf.ln(2)
            setf("B", max(16 - b.level, 12))
            cell(8, _strip_inline(b.text))
            pdf.ln(1)
        elif b.kind in ("bullet", "ordered"):
            setf("", 12)
            marker = "•  " if b.kind == "bullet" else "–  "
            cell(7, marker + _strip_inline(b.text))
        elif b.kind == "quote":
            setf("I" if have_font and os.path.exists(_DEJAVU_BOLD) else "", 12)
            cell(7, "    " + _strip_inline(b.text))
        elif b.kind == "hr":
            y = pdf.get_y() + 2
            pdf.line(20, y, 190, y)
            pdf.ln(5)
        elif b.kind == "p":
            setf("", 12)
            cell(7, _strip_inline(b.text))
            pdf.ln(1)

    out = pdf.output()
    return bytes(out)


# ─────────────────────────── HTML ───────────────────────────

_HTML_TMPL = """<!doctype html>
<html lang="vi"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
  :root {{ color-scheme: light dark; }}
  body {{ font-family: 'Times New Roman', Georgia, serif; line-height: 1.7;
    max-width: 720px; margin: 0 auto; padding: 40px 24px; color: #1a1a1a;
    background: #fff; }}
  @media (prefers-color-scheme: dark) {{ body {{ color: #e6e6e6; background: #1a1a1a; }} }}
  h1 {{ font-size: 2em; border-bottom: 2px solid currentColor; padding-bottom: .3em; }}
  h2 {{ font-size: 1.5em; margin-top: 1.6em; }}
  h3 {{ font-size: 1.2em; }}
  blockquote {{ border-left: 3px solid #999; margin: 1em 0; padding-left: 1em;
    font-style: italic; opacity: .9; }}
  hr {{ border: none; border-top: 1px solid #ccc; margin: 2em 0; }}
  p {{ text-align: justify; }}
  @media print {{ body {{ max-width: none; }} }}
</style></head><body>
{body}
</body></html>"""


def to_html(title: str, md: str) -> bytes:
    """Markdown → HTML tự chứa (mở trên điện thoại → 'Lưu thành PDF'). Phủ mọi ngôn ngữ."""
    def esc(s: str) -> str:
        s = _html.escape(s)
        return _BOLD_RE.sub(r"<strong>\1</strong>", s)

    parts: list[str] = []
    if title:
        parts.append(f"<h1>{_html.escape(title)}</h1>")

    list_open = None  # 'ul' | 'ol' | None

    def close_list():
        nonlocal list_open
        if list_open:
            parts.append(f"</{list_open}>")
            list_open = None

    for b in parse_markdown(md):
        if b.kind in ("bullet", "ordered"):
            want = "ul" if b.kind == "bullet" else "ol"
            if list_open != want:
                close_list()
                parts.append(f"<{want}>")
                list_open = want
            parts.append(f"<li>{esc(b.text)}</li>")
            continue
        close_list()
        if b.kind == "h":
            lvl = min(max(b.level, 1), 6)
            parts.append(f"<h{lvl}>{esc(b.text)}</h{lvl}>")
        elif b.kind == "quote":
            parts.append(f"<blockquote>{esc(b.text)}</blockquote>")
        elif b.kind == "hr":
            parts.append("<hr>")
        elif b.kind == "p":
            parts.append(f"<p>{esc(b.text)}</p>")
    close_list()

    doc = _HTML_TMPL.format(title=_html.escape(title or "Tài liệu"), body="\n".join(parts))
    return doc.encode("utf-8")


# ─────────────────────────── Điều phối ───────────────────────────

_MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "html": "text/html; charset=utf-8",
}


def generate(fmt: str, title: str, md: str) -> tuple[bytes, str, str]:
    """Trả (bytes, mime, phần mở rộng). fmt ∈ {docx, pdf, html}; mặc định docx."""
    fmt = (fmt or "docx").lower()
    if fmt == "pdf":
        return to_pdf(title, md), _MIME["pdf"], "pdf"
    if fmt == "html":
        return to_html(title, md), _MIME["html"], "html"
    return to_docx(title, md), _MIME["docx"], "docx"
