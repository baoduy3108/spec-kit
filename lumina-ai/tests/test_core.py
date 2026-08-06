"""✦ LUMINA AI — unit tests cho phần lõi (không cần API key / mạng)."""

import os
import sys
import tempfile
import time
import uuid

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

os.environ.setdefault("DEV_MODE", "true")
# DB riêng cho test — không đụng vào lumina.db thật (phải set TRƯỚC khi import app.config)
os.environ.setdefault("DB_PATH", os.path.join(tempfile.mkdtemp(prefix="lumina_test_"), "test.db"))

from app import db  # noqa: E402
from app.cache import ResponseCache  # noqa: E402
from app.circuit_breaker import BreakerState, CircuitBreaker  # noqa: E402
from app.memory import trim_history  # noqa: E402
from app.ratelimit import UserRateLimiter  # noqa: E402
from app.router import decide_route  # noqa: E402
from app.config import CONFIG  # noqa: E402


def _new_user_id() -> str:
    uid = f"test-{uuid.uuid4().hex[:8]}"
    db.upsert_user(uid, f"{uid}@example.com", "Test User", "")
    return uid


# ── Router: "mô hình bù trừ" chọn đúng chế độ ───────────────────────────────

def test_router_search_intent():
    route = decide_route("Tin tức AI mới nhất hôm nay là gì?")
    assert route.mode == "search"
    assert route.use_web_search is True


def test_router_deep_intent():
    route = decide_route("Viết hàm Python kiểm tra số nguyên tố và giải thích thuật toán")
    assert route.mode == "deep"
    assert route.model == CONFIG["CLAUDE_MODEL_DEEP"]
    assert route.effort == "high"


def test_router_fast_for_greeting():
    route = decide_route("xin chào")
    assert route.mode == "fast"
    assert route.model == CONFIG["CLAUDE_MODEL_FAST"]


def test_router_balanced_default():
    route = decide_route(
        "Kể cho mình nghe một câu chuyện ngắn về tình bạn giữa chú mèo và chú chó nhé, giọng văn ấm áp",
        history_len=2,
    )
    assert route.mode == "balanced"


def test_router_apex_disabled_by_default():
    route = decide_route("Đây là một bài toán cực khó cần suy luận sâu")
    assert route.mode != "apex"  # ENABLE_FABLE mặc định false → không dùng Fable


# ── Router: chế độ đa phương thức (vẽ ảnh / nghiên cứu) ─────────────────────

def test_router_image_intent():
    for q in ["vẽ cho tôi một con mèo", "tạo ảnh hoàng hôn trên biển", "draw a robot"]:
        assert decide_route(q).mode == "image_gen", q


def test_router_research_intent():
    route = decide_route("nghiên cứu sâu về tác động của AI tới việc làm")
    assert route.mode == "research"
    assert route.use_web_search is True


def test_router_force_mode_overrides():
    # Nút bấm giao diện ép chế độ dù nội dung không có từ khóa
    assert decide_route("con mèo", force_mode="image").mode == "image_gen"
    assert decide_route("biến đổi khí hậu", force_mode="research").mode == "research"


def test_router_no_false_positive_image():
    for q in ["Đà Lạt có gì chơi", "cách tạo file excel", "tạo thói quen tốt"]:
        assert decide_route(q).mode != "image_gen", q


def test_router_apex_locked_for_free_plan():
    """Bật ENABLE_FABLE tạm thời để kiểm tra gói Miễn phí bị khóa Đỉnh cao."""
    original = CONFIG["ENABLE_FABLE"]
    CONFIG["ENABLE_FABLE"] = True
    try:
        route = decide_route("Đây là một bài toán cực khó cần suy luận sâu", apex_allowed=False)
        assert route.mode == "deep"
        assert route.apex_locked is True

        route_paid = decide_route("Đây là một bài toán cực khó cần suy luận sâu", apex_allowed=True)
        assert route_paid.mode == "apex"
        assert route_paid.apex_locked is False
    finally:
        CONFIG["ENABLE_FABLE"] = original


# ── Memory: cắt cửa sổ ngữ cảnh ─────────────────────────────────────────────

def test_trim_history_keeps_recent():
    msgs = [{"role": "user" if i % 2 == 0 else "assistant", "content": "x" * 400} for i in range(50)]
    trimmed = trim_history(msgs, max_tokens=1000)
    assert len(trimmed) < 50
    assert trimmed[-1] == msgs[-1]          # luôn giữ tin nhắn cuối
    assert trimmed[0]["role"] == "user"     # tin đầu phải là user


def test_trim_history_always_keeps_last():
    msgs = [{"role": "user", "content": "y" * 100000}]
    trimmed = trim_history(msgs, max_tokens=10)
    assert len(trimmed) == 1


# ── Cache ───────────────────────────────────────────────────────────────────

def test_cache_roundtrip_and_ttl():
    cache = ResponseCache(ttl=1)
    key = ResponseCache.make_key([{"role": "user", "content": "hi"}], "claude-opus-4-8")
    cache.set(key, {"content": "xin chào", "usage": {"output_tokens": 3}})
    assert cache.get(key)["content"] == "xin chào"
    time.sleep(1.1)
    assert cache.get(key) is None  # hết hạn TTL


# ── Circuit breaker ─────────────────────────────────────────────────────────

def test_circuit_breaker_opens_and_recovers():
    breaker = CircuitBreaker(threshold=3, timeout=1)
    assert breaker.allow_request()
    for _ in range(3):
        breaker.record_failure()
    assert breaker.state == BreakerState.OPEN
    assert not breaker.allow_request()
    time.sleep(1.1)
    assert breaker.allow_request()  # HALF_OPEN thử lại
    breaker.record_success()
    assert breaker.state == BreakerState.CLOSED


# ── Rate limiter ────────────────────────────────────────────────────────────

def test_rate_limiter_blocks_after_burst():
    limiter = UserRateLimiter(rate_per_minute=60, burst=3)
    for _ in range(3):
        ok, _ = limiter.check("user-1")
        assert ok
    ok, wait = limiter.check("user-1")
    assert not ok
    assert wait >= 1
    # Người dùng khác không bị ảnh hưởng
    ok, _ = limiter.check("user-2")
    assert ok


def test_rate_limiter_upgrades_capacity_on_plan_change():
    """Khi gói của người dùng đổi (nâng cấp), bucket phải cập nhật dung lượng mới ngay."""
    limiter = UserRateLimiter(rate_per_minute=60, burst=15)
    ok, _ = limiter.check("user-3", rate_per_minute=10, burst=2)  # gói free: burst nhỏ
    assert ok
    ok, _ = limiter.check("user-3", rate_per_minute=10, burst=2)
    assert ok
    ok, _ = limiter.check("user-3", rate_per_minute=10, burst=2)
    assert not ok  # hết burst của gói free

    # Nâng cấp gói trả phí — burst cao hơn phải cho phép thêm ngay
    ok, _ = limiter.check("user-3", rate_per_minute=30, burst=50)
    assert ok


# ── Gói & đơn hàng / thanh toán (DB) ─────────────────────────────────────────

def test_new_user_defaults_to_free_plan():
    uid = _new_user_id()
    plan = db.get_effective_plan(uid)
    assert plan["key"] == "free"
    assert plan["apex_allowed"] is False
    assert plan["expires_at"] == 0


def test_create_order_and_mark_paid_activates_plan():
    uid = _new_user_id()
    order = db.create_order(uid, "monthly", "sepay", 500000, 0.0)
    assert order["status"] == "pending"
    assert order["id"].startswith("LUM")

    activated, plan = db.mark_order_paid(order["id"], provider_ref="tx1")
    assert activated is True
    assert plan["key"] == "monthly"
    assert plan["apex_allowed"] is True
    assert plan["expires_at"] > time.time()

    assert db.get_order(order["id"])["status"] == "paid"
    assert db.get_effective_plan(uid)["key"] == "monthly"


def test_mark_order_paid_idempotent():
    uid = _new_user_id()
    order = db.create_order(uid, "monthly", "sepay", 500000, 0.0)
    activated1, plan1 = db.mark_order_paid(order["id"])
    expires1 = plan1["expires_at"]
    # Webhook bắn trùng — không kích hoạt lại, không cộng dồn thời hạn
    activated2, plan2 = db.mark_order_paid(order["id"])
    assert activated1 is True
    assert activated2 is False
    assert plan2 is None
    assert db.get_effective_plan(uid)["expires_at"] == expires1


def test_mark_nonexistent_order():
    activated, plan = db.mark_order_paid("LUMZZZZZZ")
    assert activated is False
    assert plan is None


def test_sepay_auth_and_content_matching():
    from app import payments
    from app.config import CONFIG
    CONFIG["SEPAY_API_KEY"] = "secret-key"
    try:
        assert payments.verify_sepay_authorization("Apikey secret-key") is True
        assert payments.verify_sepay_authorization("Apikey wrong") is False
        assert payments.verify_sepay_authorization("") is False
    finally:
        CONFIG["SEPAY_API_KEY"] = ""

    # Tìm mã đơn trong nội dung chuyển khoản (kể cả lẫn chữ khác)
    pending = ["LUMABC123", "LUMDEF456"]
    assert payments.extract_order_id_from_content("chuyen tien LUMABC123 cam on", pending) == "LUMABC123"
    assert payments.extract_order_id_from_content("khong co ma", pending) is None


def test_vietqr_url_builder():
    from app import payments
    url = payments.build_vietqr_url("970436", "0123456789", "NGUYEN VAN A", 500000, "LUMABC123")
    assert "970436-0123456789" in url
    assert "amount=500000" in url
    assert "LUMABC123" in url


def test_premium_then_free_then_blocked():
    """2 lượt đầu = cao cấp, tới total_cap thì tụt free, vượt total_cap thì chặn."""
    uid = _new_user_id()
    # premium_cap=2, total_cap=4
    a1, p1, _ = db.consume_daily_usage(uid, premium_cap=2, total_cap=4)
    a2, p2, _ = db.consume_daily_usage(uid, premium_cap=2, total_cap=4)
    assert (a1, p1) == (True, True)   # lượt 1 — cao cấp
    assert (a2, p2) == (True, True)   # lượt 2 — cao cấp
    a3, p3, _ = db.consume_daily_usage(uid, premium_cap=2, total_cap=4)
    a4, p4, _ = db.consume_daily_usage(uid, premium_cap=2, total_cap=4)
    assert (a3, p3) == (True, False)  # lượt 3 — free (hết cao cấp)
    assert (a4, p4) == (True, False)  # lượt 4 — free
    a5, p5, used5 = db.consume_daily_usage(uid, premium_cap=2, total_cap=4)
    assert a5 is False               # lượt 5 — chạm total_cap, chặn
    assert used5 == 4

    premium_used, total_used = db.get_daily_usage(uid)
    assert premium_used == 2
    assert total_used == 4


def test_daily_usage_unlimited_when_cap_zero():
    uid = _new_user_id()
    for _ in range(5):
        allowed, use_premium, _ = db.consume_daily_usage(uid, premium_cap=0, total_cap=0)
        assert allowed is True
        assert use_premium is True  # premium_cap<=0 → luôn cao cấp


def test_orchestrator_chain_tiers():
    """Chuỗi cao cấp có Claude đứng đầu; chuỗi free không có Claude."""
    from app.orchestrator import orchestrator
    premium_chain = orchestrator._chain_for(use_premium=True)
    free_chain = orchestrator._chain_for(use_premium=False)
    assert premium_chain[0] == "claude"
    assert "claude" not in free_chain
    # Các engine free phải nằm trong danh sách đã đăng ký
    for name in free_chain:
        assert name in orchestrator.engines


def test_local_models_registered_as_fallback():
    """LOCAL_MODELS được đăng ký thành ≥5 bộ não local (dự phòng khi hết token API),
    nằm trong chuỗi free, và KHÔNG kích hoạt khi chưa tự host (thiếu OLLAMA_BASE_URL)."""
    from app.orchestrator import Orchestrator
    from app.config import CONFIG
    o = Orchestrator()
    local_names = [n for n in o.engines if getattr(o.engines[n], "is_local", False)]
    # ollama chính + ≥5 model local từ LOCAL_MODELS
    assert len(local_names) >= 5
    assert len(CONFIG["LOCAL_MODELS"]) >= 5
    # mỗi model local có model riêng và nằm trong chuỗi free
    for n in local_names:
        assert o.engines[n].model
    assert any(n.startswith("ollama-") for n in o.free_chain)
    # Không có OLLAMA_BASE_URL trong test → tất cả local đều không available (không tốn tài nguyên)
    if not CONFIG["OLLAMA_BASE_URL"]:
        assert all(not o.engines[n].available() for n in local_names)


def test_mistral_engine_registered():
    """Mistral API (cloud) được đăng ký, nằm trong chuỗi free, tắt khi thiếu key."""
    from app.orchestrator import Orchestrator
    from app.config import CONFIG
    o = Orchestrator()
    assert "mistral" in o.engines
    assert "mistral" in o.free_chain
    assert o.engines["mistral"].model
    if not CONFIG["MISTRAL_API_KEY"]:
        assert not o.engines["mistral"].available()


def test_kimi_engine_registered():
    """Kimi K3 (Moonshot, API tương thích OpenAI) được đăng ký trong chuỗi free,
    dùng qua API (KHÔNG host trọng số 2.8T), tắt khi thiếu key."""
    from app.orchestrator import Orchestrator
    from app.config import CONFIG
    o = Orchestrator()
    assert "kimi" in o.engines
    assert "kimi" in o.free_chain
    assert o.engines["kimi"].model == CONFIG["KIMI_MODEL"]
    assert "kimi" in o.engines["kimi"].base_url or "moonshot" in o.engines["kimi"].base_url
    if not CONFIG["KIMI_API_KEY"]:
        assert not o.engines["kimi"].available()


def test_local_only_mode_uses_only_local_engines():
    """LOCAL_ONLY=true: LUMINA chỉ dùng model local (Ollama), KHÔNG gọi API ngoài nào
    (không Claude/Gemini/Groq/Kimi…) → 0 token, bộ não riêng chạy trên máy người dùng."""
    from app.config import CONFIG
    from app.orchestrator import Orchestrator
    orig = CONFIG.get("LOCAL_ONLY")
    try:
        CONFIG["LOCAL_ONLY"] = True
        o = Orchestrator()
        chain = o._chain_for(use_premium=True)  # kể cả 'cao cấp' cũng KHÔNG được có API ngoài
        assert "claude" not in chain
        for ext in ("gemini", "groq", "github", "openrouter", "deepseek", "mistral", "kimi", "openai"):
            assert ext not in chain, ext
        assert chain and all(n == "ollama" or n.startswith("ollama-") for n in chain)
    finally:
        CONFIG["LOCAL_ONLY"] = orig
    # Mặc định (LOCAL_ONLY off): vẫn có API ngoài + cao cấp bắt đầu bằng Claude.
    o2 = Orchestrator()
    assert o2._chain_for(use_premium=True)[0] == "claude"
    assert "gemini" in o2.free_chain


# ── Đa phương thức: xử lý ảnh (media) ───────────────────────────────────────

def test_parse_data_url_valid():
    from app.media import parse_data_url
    assert parse_data_url("data:image/png;base64,AAAA") == ("image/png", "AAAA")
    assert parse_data_url("data:image/jpeg;base64,/9j/xyz") == ("image/jpeg", "/9j/xyz")


def test_parse_data_url_rejects_non_image():
    from app.media import parse_data_url
    assert parse_data_url("data:text/plain;base64,AAAA") is None
    assert parse_data_url("không phải data url") is None
    assert parse_data_url("") is None


def test_has_images_detection():
    from app.media import has_images
    assert has_images([{"role": "user", "content": "hi"}]) is False
    assert has_images([{"role": "user", "content": "hi", "images": ["data:image/png;base64,AA"]}]) is True


# ── Đa phương thức: tạo ảnh (imagegen) ──────────────────────────────────────

def test_imagegen_pollinations_url():
    import asyncio
    from app.imagegen import generate_image
    result = asyncio.get_event_loop().run_until_complete(generate_image("con rồng lửa"))
    assert result["url"].startswith("https://image.pollinations.ai/prompt/")
    assert result["prompt"] == "con rồng lửa"


def test_imagegen_empty_prompt_has_fallback():
    import asyncio
    from app.imagegen import generate_image
    result = asyncio.get_event_loop().run_until_complete(generate_image("   "))
    assert result["url"].startswith("https://image.pollinations.ai/")


# ── Kho tri thức nội bộ (RAG-lite — "học dần", giảm token) ──────────────────

def test_knowledge_remember_and_lookup():
    from app import knowledge
    knowledge.remember("trà xanh", "Trà xanh là loại trà làm từ lá Camellia sinensis...",
                       url="https://vi.wikipedia.org/wiki/Tr%C3%A0_xanh", source="wikipedia")
    hits = knowledge.lookup_local("lợi ích của trà xanh là gì")
    assert hits and hits[0]["topic"] == "trà xanh"
    assert "Camellia" in hits[0]["summary"]


def test_knowledge_remember_updates_not_duplicates():
    from app import knowledge
    knowledge.remember("hà nội", "Bản cũ.", source="wikipedia")
    knowledge.remember("hà nội", "Bản mới hơn.", source="wikipedia")
    hits = knowledge.lookup_local("hà nội")
    same_topic = [h for h in hits if h["topic"] == "hà nội"]
    assert len(same_topic) == 1
    assert same_topic[0]["summary"] == "Bản mới hơn."


def test_knowledge_keywords_skip_stopwords():
    from app.knowledge import extract_keywords
    kws = extract_keywords("trà xanh là gì và có lợi ích gì cho sức khỏe")
    assert "trà" in kws and "xanh" in kws
    assert "là" not in kws and "gì" not in kws


def test_learn_from_user_teaches_and_recalls():
    """'Dạy LUMINA': lệnh ghi nhớ của người dùng được tiếp thu vào kho và tra lại được
    ở lượt sau (workflow tiến hoá theo kiến thức người dùng)."""
    from app import knowledge
    learned = knowledge.learn_from_user(
        "ghi nhớ rằng dự án Zenith của tôi dùng cổng 8842 cho máy chủ nội bộ")
    assert learned is not None
    assert "8842" in learned
    # Lượt sau: hỏi lại → kho nội bộ trả về đúng điều đã dạy
    hits = knowledge.lookup_local("dự án zenith dùng cổng nào")
    assert any("8842" in h["summary"] for h in hits)


def test_learn_from_user_ignores_non_teaching():
    """Câu thường (không phải mệnh lệnh dạy) KHÔNG bị bắt nhầm để ghi nhớ."""
    from app import knowledge
    assert knowledge.learn_from_user("tôi không nhớ nổi mật khẩu wifi") is None
    assert knowledge.learn_from_user("hôm nay trời đẹp quá") is None
    assert knowledge.learn_from_user("bạn có nhớ tên tôi không") is None


def test_knowledge_build_context_warns_cross_check():
    from app.knowledge import build_context
    ctx = build_context([{"topic": "trà xanh", "summary": "abc", "url": "https://x", "source": "wikipedia"}])
    assert "ĐỐI CHIẾU" in ctx          # bắt bộ não đối chiếu chéo (chống nguồn bị sửa bịp)
    assert "https://x" in ctx           # kèm link nguồn để trích dẫn


def test_knowledge_source_trust_tiers():
    """Tầng 1 chống bịp: mỗi nguồn có mức tin riêng; nguồn lạ bị đánh dấu thận trọng."""
    from app.knowledge import source_trust
    # Wikipedia & tin tức & user = độ tin thấp/vừa (cảnh báo)
    assert "THẤP" in source_trust("news")[0] or "VỪA" in source_trust("news")[0]
    assert "VỪA" in source_trust("wikipedia")[0]
    assert "THẤP" in source_trust("user")[0]
    # Nguồn không xác định → cảnh báo rõ
    tier, desc = source_trust("nguồn-lạ-hoắc")
    assert "KHÔNG RÕ" in tier or "thận trọng" in desc


def test_knowledge_is_news_query():
    """Nhận diện câu thời sự để bật nguồn tin tức thời gian thực."""
    from app.knowledge import is_news_query
    assert is_news_query("tin tức mới nhất về bầu cử")
    assert is_news_query("diễn biến hôm nay ra sao")
    assert is_news_query("chuyện gì xảy ra năm 2025")   # có năm gần đây
    assert not is_news_query("giải thích định lý pytago")
    assert not is_news_query("cách viết hàm kiểm tra email")


def test_knowledge_build_context_three_layer_anti_deception():
    """Tầng 2+3: khối ngữ cảnh nêu rõ 3 tầng chống bịp + nhãn nguồn từng mẩu."""
    from app.knowledge import build_context
    ctx = build_context([
        {"topic": "sự kiện X", "summary": "tiêu đề tin...", "url": "https://n", "source": "news"},
        {"topic": "khái niệm Y", "summary": "định nghĩa...", "url": "https://w", "source": "wikipedia"},
    ])
    assert "3 TẦNG" in ctx or "3 tầng" in ctx.lower()
    assert "ĐỐI CHIẾU CHÉO" in ctx        # tầng 2
    assert "PHÁN ĐOÁN ĐỘC LẬP" in ctx      # tầng 3
    assert "chưa kiểm chứng" in ctx.lower()  # tin tức phải gắn nhãn


def test_knowledge_gather_prefers_local_no_network():
    import asyncio
    from app import knowledge
    knowledge.remember("số nguyên tố", "Số nguyên tố là số tự nhiên lớn hơn 1...", source="wikipedia")
    items = asyncio.get_event_loop().run_until_complete(knowledge.gather("số nguyên tố là gì"))
    assert items and items[0]["topic"] == "số nguyên tố"


def test_knowledge_gather_never_raises_offline():
    import asyncio
    from app import knowledge
    # Chủ đề không có trong kho + mạng bị chặn → phải trả [] êm, không nổ lỗi
    items = asyncio.get_event_loop().run_until_complete(knowledge.gather("zzzz-khong-ton-tai-9999"))
    assert items == []


# ── Thư viện Kỹ năng nội bộ (tuyển chọn từ Claude Skills công khai) ──────────

def test_skills_library_loads_non_empty():
    from app import skills
    assert len(skills._SKILLS) >= 40  # đã tuyển chọn 53 skill, cho biên độ an toàn


def test_skills_find_matching_skill_hits_expected():
    from app import skills
    skill = skills.find_matching_skill("làm sao viết test trước khi code cho tính năng mới")
    assert skill is not None
    assert skill.slug == "test-driven-development"


def test_skills_acquire_codebase_knowledge_matches_and_warns_no_real_files():
    from app import skills
    skill = skills.find_matching_skill("giúp mình tài liệu hóa codebase này")
    assert skill is not None
    assert skill.slug == "acquire-codebase-knowledge"
    ctx = skills.build_skill_context(skill)
    assert "không khẳng định đã tạo/ghi file thật" in ctx


def test_skills_dashmotion_matches_diagram_request():
    from app import skills
    skill = skills.find_matching_skill("vẽ sơ đồ kiến trúc động cho hệ thống này")
    assert skill is not None
    assert skill.slug == "dashmotion"


def test_skills_library_has_at_least_65():
    from app import skills
    assert len(skills._SKILLS) >= 65


def test_skills_agent_architecture_topics_match():
    from app import skills
    cases = {
        "có nên dùng nhiều AI agent phối hợp không": "multi-agent-patterns",
        "làm sao cho AI agent nhớ được across session": "memory-systems",
        "viết mô tả tool cho MCP server sao cho AI hiểu đúng": "tool-design",
        "đừng làm quá tay, chỉ sửa đúng cái tôi yêu cầu thôi": "anti-over-engineering",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_library_has_at_least_72():
    from app import skills
    assert len(skills._SKILLS) >= 72


def test_skills_platform_and_language_topics_match():
    from app import skills
    cases = {
        "review giao diện web này có đạt chuẩn accessibility không": "platform-design-web",
        "thiết kế app android theo material design 3": "platform-design-android",
        "nên dùng model nào cho tác vụ rẻ, model nào cho tác vụ khó": "model-hierarchy",
        "lỗi typescript any type sửa sao đây": "typescript-magician",
        "implement oauth pkce cho app này": "oauth-security",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_library_has_at_least_78():
    from app import skills
    assert len(skills._SKILLS) >= 78


def test_skills_design_and_security_topics_match():
    from app import skills
    cases = {
        "đánh giá giao diện này theo nguyên tắc usability": "nielsen-usability-heuristics",
        "kiểm tra code này có bị sql injection hay xss không": "vibesec",
        "chọn bảng màu oklch cho design system": "color-expert",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_library_has_at_least_82():
    from app import skills
    assert len(skills._SKILLS) >= 82


def test_skills_writing_and_platform_topics_match():
    from app import skills
    cases = {
        "sửa lại đoạn văn này cho bớt giống văn phong AI viết": "unslop",
        "review code swiftui này giúp tôi": "swiftui-expert",
        "thiết lập eslint 9 cho dự án javascript": "eslint-neostandard-linting",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_library_has_at_least_90():
    from app import skills
    assert len(skills._SKILLS) >= 90


def test_skills_ux_patterns_match():
    from app import skills
    cases = {
        "dạy tôi học python từ đầu, đi từng bước một": "teaching-framework",
        "thiết kế trạng thái loading cho trang này": "loading-states-perceived-performance",
        "màu báo lỗi nên dùng sao cho đúng": "status-colors-and-errors",
        "validate form này khi nào thì hợp lý": "form-design",
        "thêm animation nút bấm cho mượt hơn": "micro-interactions",
        "responsive design cho mobile với desktop khác nhau sao": "responsive-paradigms",
        "thiết kế toast notification cho app": "notifications-and-recovery",
        "làm bảng dữ liệu có chọn nhiều item được không": "data-display-and-selection",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_find_matching_skill_irrelevant_returns_none():
    from app import skills
    assert skills.find_matching_skill("hôm nay trời đẹp không, đi chơi đâu nhỉ") is None


def test_skills_build_skill_context_has_caveat_and_cap():
    from app import skills
    skill = skills.Skill(
        slug="test-slug", name="Test Skill", description="", category="engineering",
        keywords=("test",), body="x" * 20000,
    )
    ctx = skills.build_skill_context(skill)
    assert "KHÔNG có công cụ chạy bash/git/browser/MCP thật" in ctx
    assert len(ctx) < 20000  # đã cắt bớt, không tiêm nguyên bản dài


# ── Router: chế độ 📝 Phụ đề (video → transcript SRT) ────────────────────────

def test_router_subtitle_force_mode():
    route = decide_route("bất kỳ nội dung gì", force_mode="subtitle")
    assert route.mode == "subtitle"
    assert route.label == "📝 Phụ đề"


# ── Router: chế độ ⚙️ Lumina Forge (quy trình 6 giai đoạn) ──────────────────

def test_router_agent_force_mode():
    route = decide_route("bất kỳ nội dung gì", force_mode="agent")
    assert route.mode == "agent"
    assert route.label == "⚙️ Lumina Forge"
    assert route.use_web_search is True
    assert route.effort == "high"


def test_router_agent_not_auto_detected():
    # Lumina Forge CHỈ kích hoạt qua nút bấm (force_mode) — không có regex tự đoán,
    # để tránh vô tình chạy quy trình 6 giai đoạn nặng nề trên câu hỏi thường.
    for q in ["phân tích giúp mình đoạn code này", "thiết kế hệ thống mới", "review code giúp mình"]:
        assert decide_route(q).mode != "agent", q


def test_router_critique_force_mode():
    route = decide_route("nên chọn kiến trúc nào", force_mode="critique")
    assert route.mode == "critique" and route.label == "🔎 Phản biện"
    assert route.use_web_search is True and route.effort == "high"
    # không tự kích hoạt khi không bấm nút
    assert decide_route("nên chọn kiến trúc nào").mode != "critique"


def test_versions_crud():
    """🕑 Version hóa: lưu v1/v2, liệt kê, đọc, xóa; nhãn tự tăng."""
    from fastapi.testclient import TestClient
    from app import auth
    from app.main import app

    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}
    try:
        client = TestClient(app)
        cid = "conv-" + uid
        assert client.post("/api/versions", json={"conversation_id": cid, "content": ""}).status_code == 400
        v1 = client.post("/api/versions", json={"conversation_id": cid, "content": "ý tưởng thô"}).json()
        v2 = client.post("/api/versions", json={"conversation_id": cid, "content": "bản tối ưu", "label": "production"}).json()
        assert v1["label"] == "v1" and v2["label"] == "production"
        lst = client.get(f"/api/versions?conversation_id={cid}").json()["versions"]
        assert len(lst) == 2
        assert client.get(f"/api/versions/{v1['id']}").json()["content"] == "ý tưởng thô"
        assert client.delete(f"/api/versions/{v1['id']}").status_code == 200
        assert len(client.get(f"/api/versions?conversation_id={cid}").json()["versions"]) == 1
        assert client.get("/api/versions/khongco").status_code == 404
    finally:
        app.dependency_overrides.pop(auth.require_user, None)


# ── 📚 Tổng hợp tệp → FILE mới (docgen + /api/compose) ───────────────────────

def test_docgen_generates_all_formats_with_vietnamese():
    from app import docgen
    md = ("# Chương 1\n\nĐạo khả đạo, phi thường đạo.\n\n## Mục con\n\n"
          "- Vô danh thiên địa\n- Hữu danh vạn vật\n\n> Trích dẫn.\n\n---\n\nĐoạn kết.")
    docx_bytes, mime_d, ext_d = docgen.generate("docx", "Đạo Đức Kinh", md)
    pdf_bytes, mime_p, ext_p = docgen.generate("pdf", "Đạo Đức Kinh", md)
    html_bytes, mime_h, ext_h = docgen.generate("html", "Đạo Đức Kinh", md)
    assert ext_d == "docx" and len(docx_bytes) > 500
    assert ext_p == "pdf" and pdf_bytes[:4] == b"%PDF"
    assert ext_h == "html" and b"<h1>" in html_bytes and "Chương".encode() in html_bytes
    # docx đọc lại được và giữ tiếng Việt
    import io as _io
    import docx as _docx
    doc = _docx.Document(_io.BytesIO(docx_bytes))
    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "Đạo khả đạo" in texts and "Chương 1" in texts


def test_files_ext_detection_and_optional_markitdown():
    """files.py: nhận mimetype từ đuôi khi octet-stream; text luôn đọc được (2 tầng)."""
    import base64
    from app import files

    def durl(mime, data):
        return f"data:{mime};base64," + base64.b64encode(data).decode()

    # text/plain luôn đọc được (markitdown hay fallback đều xong)
    r = files.extract_text("a.txt", durl("text/plain", "Xin chào LUMINA".encode()))
    assert r["error"] == "" and "LUMINA" in r["text"]
    # octet-stream + đuôi .md → suy ra mimetype và đọc được
    r = files.extract_text("note.md", durl("application/octet-stream", "# Tiêu đề\n\nnội dung".encode()))
    assert r["error"] == "" and "Tiêu đề" in r["text"]
    # định dạng lạ không có đuôi hỗ trợ → lỗi nhẹ nhàng, không raise
    r = files.extract_text("x.bin", durl("application/x-thing", b"\x00\x01\x02"))
    assert r["text"] == "" and r["error"]


def test_docgen_unknown_format_defaults_docx():
    from app import docgen
    data, mime, ext = docgen.generate("weird", "T", "# H\n\ndoan")
    assert ext == "docx"


def test_compose_endpoint_synthesizes_file():
    """/api/compose: đọc tệp → (bộ não giả) tổng hợp → trả FILE tải về."""
    import base64
    from fastapi.testclient import TestClient
    from app import auth
    from app.main import app, orchestrator

    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}

    async def _fake_run(messages, route, use_premium=True, system_extra=""):
        yield {"type": "text", "text": "# Sách tổng hợp\n\nNội dung tổng hợp từ hai tệp nguồn."}

    orig = orchestrator.run
    orchestrator.run = _fake_run
    try:
        client = TestClient(app)
        src = base64.b64encode("Đạo khả đạo, phi thường đạo.".encode()).decode()
        data_url = f"data:text/plain;base64,{src}"
        # docx
        r = client.post("/api/compose", json={
            "files": [{"name": "nguon1.txt", "data_url": data_url},
                      {"name": "nguon2.txt", "data_url": data_url}],
            "instruction": "Tổng hợp thành sách", "title": "Đạo Đức Kinh Tổng Hợp", "format": "docx"})
        assert r.status_code == 200, r.text
        assert "wordprocessingml" in r.headers["content-type"]
        assert ".docx" in r.headers["content-disposition"]
        assert len(r.content) > 500
        # pdf
        r2 = client.post("/api/compose", json={
            "files": [{"name": "n.txt", "data_url": data_url}], "format": "pdf", "title": "T"})
        assert r2.status_code == 200 and r2.content[:4] == b"%PDF"
        # không tệp → 400
        assert client.post("/api/compose", json={"files": [], "format": "docx"}).status_code == 400
        # định dạng sai → 400
        assert client.post("/api/compose", json={
            "files": [{"name": "n.txt", "data_url": data_url}], "format": "xls"}).status_code == 400
    finally:
        orchestrator.run = orig
        app.dependency_overrides.pop(auth.require_user, None)


# ── 🔑 API riêng + ~30 LLM phụ ───────────────────────────────────────────────

def test_orchestrator_has_many_openrouter_brains():
    from app.orchestrator import Orchestrator
    o = Orchestrator()
    orn = [n for n in o.engines if n.startswith("openrouter")]
    assert len(orn) >= 30  # base + ~30 model free
    assert len(o.free_chain) >= 40  # tổng bộ não free rất nhiều


def test_api_key_lifecycle_and_openai_endpoint():
    from fastapi.testclient import TestClient
    from app import auth, db
    from app.main import app, orchestrator

    uid = _new_user_id()
    db.upsert_user(uid, f"{uid}@x.com", "APIUser", "")
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "APIUser", "picture": "", "is_admin": False}

    async def _fake_run(messages, route, use_premium=True, system_extra=""):
        yield {"type": "text", "text": "Xin chào từ LUMINA API."}

    orig = orchestrator.run
    orchestrator.run = _fake_run
    try:
        client = TestClient(app)
        # tạo key
        r = client.post("/api/keys", json={"name": "test"})
        assert r.status_code == 200
        raw = r.json()["key"]
        assert raw.startswith("lum_")
        assert len(client.get("/api/keys").json()["keys"]) == 1
        # gọi /v1/chat/completions bằng Bearer key (không cookie)
        r = client.post("/v1/chat/completions",
                        headers={"Authorization": f"Bearer {raw}"},
                        json={"model": "lumina", "messages": [{"role": "user", "content": "hi"}]})
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["model"] == "lumina" and j["object"] == "chat.completion"
        assert "LUMINA API" in j["choices"][0]["message"]["content"]
        # key sai → 401
        assert client.post("/v1/chat/completions", headers={"Authorization": "Bearer lum_bad"},
                           json={"messages": [{"role": "user", "content": "x"}]}).status_code == 401
        # thu hồi key
        kid = client.get("/api/keys").json()["keys"][0]["id"]
        assert client.delete(f"/api/keys/{kid}").status_code == 200
        assert client.post("/v1/chat/completions", headers={"Authorization": f"Bearer {raw}"},
                           json={"messages": [{"role": "user", "content": "x"}]}).status_code == 401
    finally:
        orchestrator.run = orig
        app.dependency_overrides.pop(auth.require_user, None)


# ── 🌐 Chợ Agent (cộng đồng: chia sẻ + cài) ──────────────────────────────────

def test_agent_marketplace_share_and_install():
    from app import db
    author = _new_user_id()
    buyer = _new_user_id()
    db.upsert_user(author, f"{author}@x.com", "Tác Giả", "")
    db.upsert_user(buyer, f"{buyer}@x.com", "Người Cài", "")
    # tác giả tạo agent, ban đầu KHÔNG lên chợ
    aid = db.create_agent(author, "Trợ lý Toán", "Bạn là gia sư Toán chuyên nghiệp.", "🧮")
    assert not any(a["id"] == aid for a in db.list_shared_agents())
    # chia sẻ → xuất hiện trên chợ với tên tác giả
    assert db.set_agent_shared(aid, author, True, author="Tác Giả")
    shop = db.list_shared_agents(exclude_user=buyer)
    row = next((a for a in shop if a["id"] == aid), None)
    assert row and row["author"] == "Tác Giả" and row["mine"] is False and "gia sư" in row["preview"]
    # người khác cài → có bản sao trong bộ sưu tập của họ, installs tăng
    inst = db.install_shared_agent(aid, buyer)
    assert inst and inst["name"] == "Trợ lý Toán"
    assert any(a["id"] == inst["id"] for a in db.list_agents(buyer))
    assert db.get_shared_agent(aid)["installs"] == 1
    # cài lại → không tạo trùng, trả bản đã có
    inst2 = db.install_shared_agent(aid, buyer)
    assert inst2["id"] == inst["id"]
    assert db.get_shared_agent(aid)["installs"] == 1  # không tăng nữa
    # không cho tự cài agent của chính mình
    assert db.install_shared_agent(aid, author) is None
    # gỡ chia sẻ → biến khỏi chợ
    db.set_agent_shared(aid, author, False)
    assert not any(a["id"] == aid for a in db.list_shared_agents())


def test_agent_marketplace_endpoints():
    from fastapi.testclient import TestClient
    from app import auth, db
    from app.main import app
    author = _new_user_id(); buyer = _new_user_id()
    db.upsert_user(author, f"{author}@x.com", "AuthorX", "")
    aid = db.create_agent(author, "Agent Chia Sẻ", "Persona test", "🤝")

    def as_user(uid, name):
        app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                               "name": name, "picture": "", "is_admin": False}
    client = TestClient(app)
    try:
        as_user(author, "AuthorX")
        assert client.post(f"/api/agents/{aid}/share", json={"shared": True}).status_code == 200
        as_user(buyer, "Buyer")
        shop = client.get("/api/marketplace").json()["agents"]
        assert any(a["id"] == aid for a in shop)
        r = client.post(f"/api/marketplace/{aid}/install")
        assert r.status_code == 200 and r.json()["name"] == "Agent Chia Sẻ"
        assert any(a["name"] == "Agent Chia Sẻ" for a in client.get("/api/agents").json()["agents"])
    finally:
        app.dependency_overrides.pop(auth.require_user, None)


# ── 📊 GraphRAG: đồ thị tri thức ─────────────────────────────────────────────

def test_graph_rag_ingest_and_search():
    from app import graph_rag as g
    g.reset()
    g.ingest_text("Lão Tử viết Đạo Đức Kinh. Lão Tử sống ở nước Sở.")
    g.ingest_text("Trang Tử kế thừa Lão Tử. Trang Tử viết Nam Hoa Kinh.")
    st = g.stats()
    assert st["nodes"] >= 4 and st["edges"] >= 3
    # local: quanh "Lão Tử" phải thấy Đạo Đức Kinh + Trang Tử
    local = g.local_search("Lão Tử liên quan gì")
    ctx = g.build_context(local)
    assert "Lão Tử" in ctx and ("Đạo Đức Kinh" in ctx or "Trang Tử" in ctx)
    # global: hub hàng đầu là Lão Tử hoặc Đạo Đức Kinh
    hubs = g.global_overview(5)["hubs"]
    assert any(h["entity"] in ("Lão Tử", "Đạo Đức Kinh") for h in hubs)
    # entity extraction KHÔNG bắt mảnh giữa từ (không có "ết"/"ống")
    names = {h["entity"] for h in hubs}
    assert not any(n.startswith(("ế", "ố", "ừ", "ó ")) for n in names)
    g.reset()
    assert g.stats() == {"nodes": 0, "edges": 0}


def test_graph_endpoints():
    from fastapi.testclient import TestClient
    from app import auth, graph_rag
    from app.main import app
    graph_rag.reset()
    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}
    try:
        client = TestClient(app)
        r = client.post("/api/graph/ingest", json={"text": "Hà Nội là thủ đô Việt Nam. Việt Nam ở Đông Nam Á."})
        assert r.status_code == 200 and r.json()["stats"]["nodes"] >= 3
        r = client.get("/api/graph/query", params={"q": "Việt Nam", "scope": "local"})
        assert r.status_code == 200 and "Việt Nam" in r.json()["context"]
        r = client.get("/api/graph/query", params={"q": "", "scope": "global"})
        assert r.status_code == 200 and r.json()["result"]["scope"] == "global"
    finally:
        graph_rag.reset()
        app.dependency_overrides.pop(auth.require_user, None)


# ── 🌍 Worldmonitor: số liệu quốc gia cập nhật ───────────────────────────────

def test_worldmonitor_detection_and_parse():
    from app import worldmonitor as w
    # phát hiện tên nước Việt/Anh, không để "hà" nuốt "hàn quốc"
    assert w.detect_country("dân số việt nam bao nhiêu") == "Vietnam"
    assert w.detect_country("thủ đô của pháp") == "France"
    assert w.detect_country("kinh tế hàn quốc") == "South Korea"
    assert w.detect_country("vẽ con mèo dễ thương") is None
    # is_world_query cần cả tên nước + từ khóa số liệu
    assert w.is_world_query("dân số nhật bản") is True
    assert w.is_world_query("hôm nay trời đẹp") is False
    assert w.is_world_query("nhật bản có anime hay") is False  # có nước nhưng không hỏi số liệu
    # parse JSON REST Countries → summary có số liệu
    sample = {"name": {"common": "Vietnam"}, "capital": ["Hanoi"], "population": 98000000,
              "area": 331212.0, "region": "Asia", "subregion": "South-Eastern Asia",
              "currencies": {"VND": {"name": "Vietnamese dong"}}, "languages": {"vie": "Vietnamese"}}
    fact = w._parse_country(sample)
    assert fact["source"] == "worldmonitor" and "Hanoi" in fact["summary"] and "98.000.000" in fact["summary"]


# ── Media: video đính kèm ────────────────────────────────────────────────────

def test_media_parse_video_data_url():
    from app.media import parse_video_data_url
    import base64
    small = base64.b64encode(b"x" * 100).decode()
    assert parse_video_data_url(f"data:video/mp4;base64,{small}") == ("video/mp4", small)
    assert parse_video_data_url("data:image/png;base64,AAAA") is None  # không phải video


def test_media_parse_video_rejects_oversized():
    from app.media import parse_video_data_url, MAX_VIDEO_BYTES
    import base64
    huge = base64.b64encode(b"x" * (MAX_VIDEO_BYTES + 1000)).decode()
    assert parse_video_data_url(f"data:video/mp4;base64,{huge}") is None


def test_media_has_videos():
    from app.media import has_videos
    assert has_videos([{"role": "user", "content": "hi"}]) is False
    assert has_videos([{"role": "user", "content": "hi", "videos": ["data:video/mp4;base64,AA"]}]) is True


def test_media_decode_video_roundtrip():
    from app.media import decode_video
    import base64
    raw = b"video bytes here"
    b64 = base64.b64encode(raw).decode()
    assert decode_video(f"data:video/mp4;base64,{b64}") == raw
    assert decode_video("không phải data url") is None


# ── Tệp đính kèm: đọc PDF/Word/Excel/txt ─────────────────────────────────────

def test_files_extract_docx_and_xlsx_and_txt():
    import base64
    import io
    from app.files import build_context, extract_text

    import docx
    d = docx.Document()
    d.add_paragraph("Xin chào LUMINA, đây là tài liệu thử nghiệm.")
    buf = io.BytesIO(); d.save(buf)
    r_docx = extract_text(
        "bao_cao.docx",
        "data:application/vnd.openxmlformats-officedocument.wordprocessingml.document;base64,"
        + base64.b64encode(buf.getvalue()).decode(),
    )
    assert r_docx["error"] == ""
    assert "LUMINA" in r_docx["text"]

    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Tên", "Điểm"]); ws.append(["An", 9])
    buf2 = io.BytesIO(); wb.save(buf2)
    r_xlsx = extract_text(
        "diem.xlsx",
        "data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,"
        + base64.b64encode(buf2.getvalue()).decode(),
    )
    assert r_xlsx["error"] == ""
    assert "An" in r_xlsx["text"]

    r_txt = extract_text("note.txt", "data:text/plain;base64," + base64.b64encode(b"ghi chu test").decode())
    assert r_txt["text"] == "ghi chu test"

    ctx = build_context([r_docx, r_xlsx, r_txt])
    assert "bao_cao.docx" in ctx and "diem.xlsx" in ctx and "note.txt" in ctx


def test_files_extract_rejects_bad_input():
    from app.files import extract_text
    assert extract_text("x.pdf", "không phải data url")["error"] != ""
    assert extract_text("x.bin", "data:application/octet-stream;base64,AAAA")["error"] != ""


def test_files_build_context_skips_errors():
    from app.files import build_context
    ok = {"name": "a.txt", "text": "nội dung", "error": ""}
    bad = {"name": "b.pdf", "text": "", "error": "hỏng"}
    ctx = build_context([ok, bad])
    assert "a.txt" in ctx
    assert "b.pdf" not in ctx  # tệp lỗi không được đưa vào ngữ cảnh


# ── Lồng tiếng & phụ đề video: các hàm thuần túy + ffmpeg thật ──────────────

def test_video_dub_srt_formatting():
    from app.video_dub import _fmt_srt_time, build_srt
    assert _fmt_srt_time(0) == "00:00:00,000"
    assert _fmt_srt_time(65.5) == "00:01:05,500"
    segs = [{"start": 0, "end": 2.5, "original": "hi", "translated": "chào"}]
    srt = build_srt(segs)
    assert "00:00:00,000 --> 00:00:02,500" in srt
    assert "chào" in srt


def test_video_dub_extract_json_array_variants():
    from app.video_dub import _extract_json_array
    fenced = '```json\n[{"start": 0.0, "end": 1.0, "original": "a", "translated": "b"}]\n```'
    assert _extract_json_array(fenced)[0]["translated"] == "b"
    bare = 'noise [{"start":0,"end":1,"original":"a","translated":"b"}] noise'
    assert _extract_json_array(bare)[0]["translated"] == "b"
    assert _extract_json_array("không có json") == []


def test_video_dub_ffmpeg_pipeline_real():
    """Kiểm thử thật bằng ffmpeg: tạo video test → ghép audio mới → không bị cắt ngắn."""
    import asyncio
    import os
    import tempfile
    from app.video_dub import _build_dubbed_audio, _mux_final, _probe_duration, _run_ffmpeg

    async def go():
        work = tempfile.mkdtemp()
        test_video = os.path.join(work, "test.mp4")
        await _run_ffmpeg(["-f", "lavfi", "-i", "testsrc=duration=3:size=160x120:rate=10",
                           "-f", "lavfi", "-i", "sine=frequency=1000:duration=3",
                           "-c:v", "libx264", "-c:a", "aac", "-y", test_video])
        duration = await _probe_duration(test_video)
        assert abs(duration - 3.0) < 0.2

        seg_audio = os.path.join(work, "seg_0.mp3")
        await _run_ffmpeg(["-f", "lavfi", "-i", "sine=frequency=440:duration=1.0", "-y", seg_audio])
        segments = [{"start": 0.0, "end": 1.0, "original": "hi", "translated": "chào"}]
        # Đoạn 0 rỗng (giả lập lỗi TTS) rồi mới tới đoạn có audio thật — kiểm tra
        # KHÔNG bị lệch chỉ số input khi ghép (bug đã phát hiện + sửa khi phát triển).
        dubbed = await _build_dubbed_audio(segments, [seg_audio], work, duration)
        audio_dur = await _probe_duration(dubbed)
        assert abs(audio_dur - duration) < 0.3  # audio phải khớp ĐỦ độ dài video, không bị cắt

        out_path = os.path.join(work, "output.mp4")
        await _mux_final(test_video, dubbed, None, out_path)
        assert os.path.getsize(out_path) > 1000
        out_dur = await _probe_duration(out_path)
        assert abs(out_dur - duration) < 0.3  # video cuối KHÔNG bị cắt ngắn so với gốc

    asyncio.get_event_loop().run_until_complete(go())


def test_extract_video_frames_real():
    """🎞 Tách khung hình video thật bằng ffmpeg → ảnh JPEG data URL cho bộ não nhìn."""
    import asyncio
    import base64
    import os
    import tempfile
    from app import media
    from app.video_dub import _run_ffmpeg

    async def go():
        work = tempfile.mkdtemp()
        vid = os.path.join(work, "clip.mp4")
        await _run_ffmpeg(["-f", "lavfi", "-i", "testsrc=duration=4:size=320x240:rate=10",
                           "-pix_fmt", "yuv420p", "-y", vid])
        data_url = "data:video/mp4;base64," + base64.b64encode(open(vid, "rb").read()).decode()

        frames = await media.extract_video_frames(data_url, count=4)
        assert len(frames) == 4
        for f in frames:
            assert f.startswith("data:image/jpeg;base64,")
            raw = base64.b64decode(f.split(",", 1)[1])
            assert raw[:2] == b"\xff\xd8"          # JPEG magic
            assert len(raw) > 300                    # là ảnh thật, không rỗng

        # Video hỏng / không hợp lệ → trả [] (lùi an toàn về hành vi cũ, không vỡ luồng).
        assert await media.extract_video_frames("data:video/mp4;base64,QUJD") == []
        assert await media.extract_video_frames("không phải data url") == []

    asyncio.get_event_loop().run_until_complete(go())


def test_video_link_detection_and_vtt_parse():
    """🎬 Nhận diện link video + tách phụ đề WebVTT → văn bản sạch."""
    from app import video_link

    assert video_link.is_video_link("https://www.youtube.com/watch?v=abc123")
    assert video_link.is_video_link("https://youtu.be/abc123")
    assert video_link.is_video_link("https://vimeo.com/12345")
    assert video_link.is_video_link("https://cdn.example.com/clip.mp4")
    assert not video_link.is_video_link("https://en.wikipedia.org/wiki/Cat")
    assert not video_link.is_video_link("không phải url")

    vtt = (
        "WEBVTT\n\n"
        "00:00:00.000 --> 00:00:02.000\n"
        "Xin chào các bạn\n\n"
        "00:00:02.000 --> 00:00:04.000\n"
        "Xin chào các bạn\n\n"          # dòng lặp (auto-caption) → phải gộp
        "00:00:04.000 --> 00:00:06.000\n"
        "<c>hôm nay</c> trời đẹp\n"
    )
    text = video_link._vtt_to_text(vtt)
    assert "Xin chào các bạn" in text
    assert text.count("Xin chào các bạn") == 1     # đã bỏ trùng lặp liên tiếp
    assert "hôm nay" in text and "trời đẹp" in text
    assert "-->" not in text and "WEBVTT" not in text and "<c>" not in text


def test_widget_endpoint_knowledge():
    """📊 Endpoint widget knowledge trả tri thức đã học; loại đã bỏ (status)/không hỗ trợ → 404."""
    from fastapi.testclient import TestClient
    from app import auth, knowledge
    from app.main import app

    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}
    try:
        client = TestClient(app)

        # knowledge: dạy 1 mẩu rồi widget phải tìm thấy.
        knowledge.remember("thủ đô nước Pháp", "Paris là thủ đô của Pháp.", source="manual")
        r = client.get("/api/widget/knowledge", params={"q": "thủ đô nước Pháp"})
        assert r.status_code == 200
        kd = r.json()
        assert kd["type"] == "knowledge"
        assert any("paris" in (it.get("summary", "") + it.get("topic", "")).lower() for it in kd["items"])

        # widget "status" đã bị bỏ → 404; loại không hỗ trợ → 404
        assert client.get("/api/widget/status").status_code == 404
        assert client.get("/api/widget/khongco").status_code == 404
    finally:
        app.dependency_overrides.pop(auth.require_user, None)


def test_projects_crud_and_widgets_and_conv_assign():
    """◧ Project: tạo/sửa/xóa + lưu widget mặt bàn + gán hội thoại; cô lập theo người dùng."""
    from fastapi.testclient import TestClient
    from app import auth, db
    from app.main import app

    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}
    try:
        client = TestClient(app)
        # tạo
        r = client.post("/api/projects", json={"name": "Dự án A"})
        assert r.status_code == 200
        pid = r.json()["id"]
        assert r.json()["name"] == "Dự án A" and r.json()["widgets"] == []
        assert any(p["id"] == pid for p in client.get("/api/projects").json()["projects"])

        # lưu widget (loại hợp lệ giữ lại, loại rác bị loại)
        r = client.put(f"/api/projects/{pid}", json={"widgets": [
            {"type": "news", "query": "giá vàng"}, {"type": "clock", "mode": "clock"},
            {"type": "HACK", "query": "x"},
        ]})
        assert r.status_code == 200
        widgets = r.json()["widgets"]
        assert len(widgets) == 2 and {w["type"] for w in widgets} == {"news", "clock"}

        # đổi tên
        assert client.put(f"/api/projects/{pid}", json={"name": "Dự án B"}).json()["name"] == "Dự án B"

        # gán hội thoại vào project rồi liệt kê
        conv_id = db.create_conversation(uid, "Chat trong dự án", project_id=pid)
        detail = client.get(f"/api/projects/{pid}").json()
        assert any(c["id"] == conv_id for c in detail["conversations"])

        # người dùng KHÁC không thấy / không sửa được project này
        other = _new_user_id()
        app.dependency_overrides[auth.require_user] = lambda: {"id": other, "email": f"{other}@x.com",
                                                               "name": "O", "picture": "", "is_admin": False}
        assert client.get(f"/api/projects/{pid}").status_code == 404
        assert client.delete(f"/api/projects/{pid}").status_code == 404

        # chủ sở hữu xóa được; hội thoại KHÔNG bị xóa, chỉ gỡ khỏi project
        app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                               "name": "T", "picture": "", "is_admin": False}
        assert client.delete(f"/api/projects/{pid}").status_code == 200
        assert db.get_conversation(conv_id, uid) is not None
        assert db.get_conversation(conv_id, uid).get("project_id") in (None, "")
    finally:
        app.dependency_overrides.pop(auth.require_user, None)


def test_agents_crud_and_isolation():
    """🤖 Agent tùy chỉnh: tạo/sửa/xóa, cần instructions, cô lập theo người dùng."""
    from fastapi.testclient import TestClient
    from app import auth
    from app.main import app

    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}
    try:
        client = TestClient(app)
        # thiếu instructions → 400
        assert client.post("/api/agents", json={"name": "X", "instructions": ""}).status_code == 400
        # tạo
        r = client.post("/api/agents", json={"name": "Trợ lý MKT", "emoji": "📣",
                                             "instructions": "Bạn là chuyên gia marketing, trả lời ngắn gọn."})
        assert r.status_code == 200
        aid = r.json()["id"]
        assert r.json()["emoji"] == "📣"
        assert any(a["id"] == aid for a in client.get("/api/agents").json()["agents"])
        # sửa
        assert client.put(f"/api/agents/{aid}", json={"name": "MKT Pro"}).json()["name"] == "MKT Pro"

        # người dùng khác không thấy / không xóa được
        other = _new_user_id()
        app.dependency_overrides[auth.require_user] = lambda: {"id": other, "email": f"{other}@x.com",
                                                               "name": "O", "picture": "", "is_admin": False}
        assert client.delete(f"/api/agents/{aid}").status_code == 404
        assert all(a["id"] != aid for a in client.get("/api/agents").json()["agents"])
        # chủ xóa được
        app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                               "name": "T", "picture": "", "is_admin": False}
        assert client.delete(f"/api/agents/{aid}").status_code == 200
    finally:
        app.dependency_overrides.pop(auth.require_user, None)


def test_journal_and_evolution_report():
    """🕰 Time Capsule: ghi mục nhật ký + liệt kê theo thời gian; evolution report lùi an toàn khi chưa có feedback."""
    from fastapi.testclient import TestClient
    from app import auth
    from app.main import app

    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}
    try:
        client = TestClient(app)
        assert client.post("/api/journal", json={"title": ""}).status_code == 400
        a = client.post("/api/journal", json={"kind": "idea", "title": "Ý tưởng A"}).json()
        client.post("/api/journal", json={"kind": "rejected", "title": "Bỏ A", "note": "vì phức tạp"})
        entries = client.get("/api/journal").json()["entries"]
        assert len(entries) == 2 and entries[0]["title"] == "Ý tưởng A" and entries[0]["kind"] == "idea"
        assert client.delete(f"/api/journal/{a['id']}").status_code == 200
        assert len(client.get("/api/journal").json()["entries"]) == 1

        # evolution report: chưa có feedback → thông báo, không raise
        rep = client.post("/api/evolution/report").json()
        assert rep["summary"]["total"] == 0 and "phản hồi" in rep["report"].lower()
    finally:
        app.dependency_overrides.pop(auth.require_user, None)


def test_project_memory_and_feedback():
    """📌 Trí nhớ dự án lưu/đọc được; 👍/👎 feedback ghi nhận + tổng hợp."""
    from fastapi.testclient import TestClient
    from app import auth
    from app.main import app

    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}
    try:
        client = TestClient(app)
        pid = client.post("/api/projects", json={"name": "Dự án X"}).json()["id"]
        mem = "Kiến trúc: FastAPI + SQLite. Đã bỏ Stripe vì chưa có merchant."
        r = client.put(f"/api/projects/{pid}", json={"memory": mem})
        assert r.status_code == 200 and r.json()["memory"] == mem
        assert client.get(f"/api/projects/{pid}").json()["project"]["memory"] == mem

        # feedback
        assert client.post("/api/feedback", json={"rating": 1}).json()["ok"] is True
        s = client.post("/api/feedback", json={"rating": -1, "note": "hơi dài"}).json()["summary"]
        assert s["up"] == 1 and s["down"] == 1 and s["total"] == 2
    finally:
        app.dependency_overrides.pop(auth.require_user, None)


def test_goals_crud_and_progress():
    """🎯 Goal Engine: tạo mục tiêu, đặt bước, tiến độ % tự tính, breakdown lùi an toàn khi không có engine."""
    from fastapi.testclient import TestClient
    from app import auth
    from app.main import app

    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}
    try:
        client = TestClient(app)
        assert client.post("/api/goals", json={"title": ""}).status_code == 400
        r = client.post("/api/goals", json={"title": "Học IELTS 7.5", "target_date": "2027-06-01"})
        assert r.status_code == 200
        gid = r.json()["id"]
        assert r.json()["progress"] == 0 and r.json()["steps"] == []

        # đặt 4 bước, 1 done → progress 25%
        steps = [{"text": f"Bước {i}", "done": i == 0} for i in range(4)]
        pr = client.put(f"/api/goals/{gid}", json={"steps": steps}).json()
        assert len(pr["steps"]) == 4 and pr["progress"] == 25

        # tick hết → 100%
        steps2 = [{"text": s["text"], "done": True} for s in pr["steps"]]
        assert client.put(f"/api/goals/{gid}", json={"steps": steps2}).json()["progress"] == 100

        # breakdown không có engine → không raise, giữ nguyên (best-effort)
        assert client.post(f"/api/goals/{gid}/breakdown").status_code == 200

        assert any(g["id"] == gid for g in client.get("/api/goals").json()["goals"])
        assert client.delete(f"/api/goals/{gid}").status_code == 200
        # người khác không xóa được (đã xóa) / mục tiêu lạ → 404
        assert client.delete("/api/goals/khongco").status_code == 404
    finally:
        app.dependency_overrides.pop(auth.require_user, None)


def test_agent_sessions_tagging_and_listing():
    """🤖 Agent như workspace: hội thoại gắn agent_id (kể cả 'forge') + liệt kê session của agent."""
    from fastapi.testclient import TestClient
    from app import auth, db
    from app.main import app

    uid = _new_user_id()
    app.dependency_overrides[auth.require_user] = lambda: {"id": uid, "email": f"{uid}@x.com",
                                                           "name": "T", "picture": "", "is_admin": False}
    try:
        client = TestClient(app)
        aid = db.create_agent(uid, "Trợ lý", "Hãy nói ngắn gọn.", "🤖")
        c1 = db.create_conversation(uid, "session A", agent_id=aid)
        c2 = db.create_conversation(uid, "session B", agent_id=aid)
        cf = db.create_conversation(uid, "forge session", agent_id="forge")
        db.create_conversation(uid, "ngoài agent")  # không gắn agent

        got = {c["id"] for c in client.get(f"/api/agents/{aid}/conversations").json()["conversations"]}
        assert got == {c1, c2}
        assert cf in {c["id"] for c in client.get("/api/agents/forge/conversations").json()["conversations"]}
        # agent không tồn tại (không phải forge) → 404
        assert client.get("/api/agents/khongco/conversations").status_code == 404
    finally:
        app.dependency_overrides.pop(auth.require_user, None)


def test_orchestrator_injects_agent_instructions():
    """system_extra (hướng dẫn agent) được chèn vào system prompt của orchestrator."""
    import asyncio
    from app.orchestrator import orchestrator
    from app.router import decide_route

    captured = {}

    class _FakeEngine:
        def available(self): return True
        class breaker:
            @staticmethod
            def allow_request(): return True
            @staticmethod
            def record_success(): pass
        async def stream_chat(self, messages, route, system_prompt):
            captured["sys"] = system_prompt
            yield {"type": "text", "text": "ok"}

    route = decide_route("chào bạn", history_len=0, apex_allowed=False, force_mode=None)
    orig = dict(orchestrator.engines)
    orig_chain = list(orchestrator.free_chain)
    try:
        orchestrator.engines["claude"] = _FakeEngine()
        orchestrator.free_chain = []  # chỉ dùng claude giả
        async def go():
            async for _ in orchestrator.run([{"role": "user", "content": "chào bạn"}], route,
                                            use_premium=True, system_extra="Bạn LUÔN nói như hải tặc."):
                pass
        asyncio.get_event_loop().run_until_complete(go())
    finally:
        orchestrator.engines = orig
        orchestrator.free_chain = orig_chain
    assert "hải tặc" in captured.get("sys", "")
    assert "HƯỚNG DẪN AGENT TÙY CHỈNH" in captured.get("sys", "")


def test_orchestrator_system_prompt_has_workflow_directive():
    """SYSTEM_PROMPT hướng dẫn xuất khối lumina-workflow và nói rõ đây chỉ là bản thiết kế,
    LUMINA không tự chạy/lập lịch (giữ đúng ranh giới thành thật)."""
    from app.orchestrator import SYSTEM_PROMPT
    assert "lumina-workflow" in SYSTEM_PROMPT
    assert "KHÔNG tự chạy" in SYSTEM_PROMPT


def test_video_link_build_context_and_whisper_gate():
    from app import video_link, transcribe

    assert video_link.build_context({"transcript": ""}) == ""
    ctx = video_link.build_context({
        "title": "Hướng dẫn nấu phở", "transcript": "đầu tiên ninh xương",
        "transcript_source": "phụ đề gốc",
    })
    assert "Hướng dẫn nấu phở" in ctx and "ninh xương" in ctx and "phụ đề gốc" in ctx

    # Không có key Whisper nào → chép lời trả "" (lùi an toàn về chỉ khung hình).
    from app.config import CONFIG
    old_g, old_o = CONFIG.get("GROQ_API_KEY"), CONFIG.get("OPENAI_API_KEY")
    CONFIG["GROQ_API_KEY"] = ""
    CONFIG["OPENAI_API_KEY"] = ""
    try:
        assert transcribe.whisper_enabled() is False
        import asyncio
        assert asyncio.get_event_loop().run_until_complete(
            transcribe.transcribe_audio("/nonexistent/path.mp3")) == ""
    finally:
        CONFIG["GROQ_API_KEY"], CONFIG["OPENAI_API_KEY"] = old_g, old_o


# ── Trí nhớ dài hạn: nhớ lại hội thoại CŨ khi mở hội thoại MỚI ──────────────

def test_recall_finds_relevant_past_conversation():
    from app import recall
    uid = _new_user_id()
    conv1 = db.create_conversation(uid, "Dự án LUMINA")
    db.add_message(conv1, "user", "Dự án LUMINA AI của tôi deadline là ngày 20 tháng 8 nhé")
    db.add_message(conv1, "assistant", "Đã ghi nhận, deadline là 20/8.")
    conv2 = db.create_conversation(uid, "Nấu ăn")
    db.add_message(conv2, "user", "Cách nấu phở bò ngon")

    conv3 = db.create_conversation(uid, "Hội thoại mới")
    items = recall.gather(uid, "deadline dự án LUMINA AI là khi nào", exclude_conv_id=conv3)
    assert items
    assert any("20" in it["snippet"] for it in items)
    assert not any("phở" in it["snippet"].lower() for it in items)


def test_recall_isolated_between_users():
    """Bảo mật quan trọng nhất: KHÔNG BAO GIỜ được lộ dữ liệu giữa 2 người dùng khác nhau."""
    from app import recall
    uid_a = _new_user_id()
    uid_b = _new_user_id()
    conv_a = db.create_conversation(uid_a, "Bí mật A")
    db.add_message(conv_a, "user", "Dự án LUMINA bí mật của tôi có mật khẩu SECRET-AAA")
    conv_b = db.create_conversation(uid_b, "Bí mật B")
    db.add_message(conv_b, "user", "Dự án LUMINA bí mật của tôi có mật khẩu SECRET-BBB")

    items_b = recall.gather(uid_b, "mật khẩu dự án LUMINA của tôi là gì", exclude_conv_id="")
    joined = " ".join(it["snippet"] for it in items_b)
    assert "SECRET-BBB" in joined       # thấy đúng dữ liệu của chính mình
    assert "SECRET-AAA" not in joined   # TUYỆT ĐỐI không thấy dữ liệu của người khác


def test_recall_excludes_current_conversation():
    from app import recall
    uid = _new_user_id()
    conv1 = db.create_conversation(uid, "Hội thoại hiện tại")
    db.add_message(conv1, "user", "Từ khóa đặc biệt XYZKEYWORD nằm trong chính hội thoại này")
    items = recall.gather(uid, "XYZKEYWORD", exclude_conv_id=conv1)
    assert items == []  # bị loại vì exclude_conv_id — history của hội thoại hiện tại đã có sẵn rồi


def test_recall_build_context_has_anti_hallucination_directive():
    from app import recall
    ctx = recall.build_context([{"conversation_id": "c1", "title": "T", "role": "user",
                                 "snippet": "nội dung", "created_at": 0}])
    assert "không bịa" in ctx
    assert "T" in ctx


def test_recall_no_keywords_returns_empty():
    from app import recall
    uid = _new_user_id()
    assert recall.gather(uid, "là gì và có", exclude_conv_id="") == []  # toàn từ dừng


def test_db_search_messages_empty_keywords():
    assert db.search_messages("any-user", [], exclude_conv_id="") == []


# ── 🌐 Đọc trang web khi người dùng dán link (mọi bộ não, không chỉ Claude) ──

def test_webpage_extract_urls():
    from app.webpage import extract_urls
    msg = "đọc giúp tôi https://example.com/page?a=1 nói về gì, còn cả http://test.org/x nữa."
    urls = extract_urls(msg)
    assert urls == ["https://example.com/page?a=1", "http://test.org/x"]


def test_webpage_extract_urls_dedupe_and_limit():
    from app.webpage import extract_urls
    msg = "https://a.com https://a.com https://b.com https://c.com https://d.com"
    urls = extract_urls(msg, limit=2)
    assert urls == ["https://a.com", "https://b.com"]  # loại trùng + giới hạn số lượng


def test_webpage_extract_urls_none():
    from app.webpage import extract_urls
    assert extract_urls("không có link nào ở đây cả") == []


def test_webpage_parses_html_strips_boilerplate():
    """Kiểm tra tách nội dung: loại script/nav/header/footer, giữ đúng nội dung chính."""
    from bs4 import BeautifulSoup
    html = """
    <html><head><title>Bai viet hay</title></head>
    <body>
    <nav>Menu dieu huong</nav>
    <script>alert("xau")</script>
    <header>Header quang cao</header>
    <article><h1>Tieu de</h1><p>Noi dung chinh ve LUMINA AI.</p></article>
    <footer>Footer ban quyen</footer>
    </body></html>
    """
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "svg"]):
        tag.decompose()
    text = soup.get_text(separator="\n", strip=True)
    assert "dieu huong" not in text
    assert "xau" not in text
    assert "quang cao" not in text
    assert "ban quyen" not in text
    assert "LUMINA AI" in text


def test_webpage_fetch_page_never_raises_on_bad_url():
    import asyncio
    from app.webpage import fetch_page
    result = asyncio.get_event_loop().run_until_complete(fetch_page("https://khong-ton-tai-9999.invalid"))
    assert result["error"] != ""  # báo lỗi thân thiện, không raise


def test_webpage_build_context_skips_errors():
    from app.webpage import build_context
    ok = {"url": "https://a.com", "title": "Trang A", "text": "nội dung A", "error": ""}
    bad = {"url": "https://b.com", "title": "", "text": "", "error": "hỏng"}
    ctx = build_context([ok, bad])
    assert "Trang A" in ctx
    assert "b.com" not in ctx  # trang lỗi không được đưa vào ngữ cảnh


def test_webpage_build_context_empty_when_no_pages():
    from app.webpage import build_context
    assert build_context([]) == ""
    assert build_context([{"url": "x", "title": "", "text": "", "error": "lỗi"}]) == ""


def test_skills_library_has_at_least_110():
    from app import skills
    assert len(skills._SKILLS) >= 110


def test_skills_library_has_at_least_115():
    from app import skills
    assert len(skills._SKILLS) >= 115


def test_skills_library_has_at_least_119():
    from app import skills
    assert len(skills._SKILLS) >= 119


def test_skills_library_has_at_least_126():
    from app import skills
    assert len(skills._SKILLS) >= 126


def test_skills_library_has_at_least_127():
    from app import skills
    assert len(skills._SKILLS) >= 127


def test_skills_library_has_at_least_129():
    from app import skills
    assert len(skills._SKILLS) >= 129


def test_skills_library_has_at_least_131():
    from app import skills
    assert len(skills._SKILLS) >= 131


def test_skills_library_has_at_least_136():
    from app import skills
    assert len(skills._SKILLS) >= 136


def test_skills_library_has_at_least_138():
    from app import skills
    assert len(skills._SKILLS) >= 138


def test_skills_library_has_at_least_142():
    from app import skills
    assert len(skills._SKILLS) >= 142


def test_skills_library_has_at_least_144():
    from app import skills
    assert len(skills._SKILLS) >= 144


def test_skills_library_has_at_least_146():
    from app import skills
    assert len(skills._SKILLS) >= 146


def test_skills_library_has_at_least_156():
    from app import skills
    assert len(skills._SKILLS) >= 156


def test_skills_library_has_at_least_168():
    from app import skills
    assert len(skills._SKILLS) >= 168


def test_skills_library_has_at_least_182():
    from app import skills
    assert len(skills._SKILLS) >= 182


def test_skills_library_has_at_least_194():
    from app import skills
    assert len(skills._SKILLS) >= 194


def test_skills_library_has_at_least_207():
    from app import skills
    assert len(skills._SKILLS) >= 207


def test_skills_library_has_at_least_221():
    from app import skills
    assert len(skills._SKILLS) >= 221


def test_skills_library_has_at_least_235():
    from app import skills
    assert len(skills._SKILLS) >= 235


def test_skills_library_has_at_least_250():
    from app import skills
    assert len(skills._SKILLS) >= 250


def test_skills_library_has_at_least_264():
    from app import skills
    assert len(skills._SKILLS) >= 264


def test_skills_library_has_at_least_280():
    from app import skills
    assert len(skills._SKILLS) >= 280


def test_skills_library_has_at_least_300():
    from app import skills
    assert len(skills._SKILLS) >= 300


def test_skills_library_has_at_least_310():
    from app import skills
    assert len(skills._SKILLS) >= 310


def test_skills_library_has_at_least_321():
    from app import skills
    assert len(skills._SKILLS) >= 321


def test_skills_library_has_at_least_335():
    from app import skills
    assert len(skills._SKILLS) >= 335


def test_skills_library_has_at_least_350():
    from app import skills
    assert len(skills._SKILLS) >= 350


def test_skills_library_has_at_least_365():
    from app import skills
    assert len(skills._SKILLS) >= 365


def test_skills_library_has_at_least_380():
    from app import skills
    assert len(skills._SKILLS) >= 380


def test_skills_library_has_at_least_395():
    from app import skills
    assert len(skills._SKILLS) >= 395


def test_skills_library_has_at_least_410():
    from app import skills
    assert len(skills._SKILLS) >= 410


def test_skills_library_has_at_least_425():
    from app import skills
    assert len(skills._SKILLS) >= 425


def test_skills_library_has_at_least_440():
    from app import skills
    assert len(skills._SKILLS) >= 440


def test_skills_library_has_at_least_455():
    from app import skills
    assert len(skills._SKILLS) >= 455


def test_skills_library_has_at_least_470():
    from app import skills
    assert len(skills._SKILLS) >= 470


def test_skills_library_has_at_least_485():
    from app import skills
    assert len(skills._SKILLS) >= 485


def test_skills_library_has_at_least_500():
    from app import skills
    assert len(skills._SKILLS) >= 500


def test_skills_library_has_at_least_514():
    from app import skills
    assert len(skills._SKILLS) >= 514


def test_skills_library_has_at_least_529():
    from app import skills
    assert len(skills._SKILLS) >= 529


def test_skills_library_has_at_least_544():
    from app import skills
    assert len(skills._SKILLS) >= 544


def test_skills_library_has_at_least_559():
    from app import skills
    assert len(skills._SKILLS) >= 559


def test_skills_library_has_at_least_574():
    from app import skills
    assert len(skills._SKILLS) >= 574


def test_skills_library_has_at_least_589():
    from app import skills
    assert len(skills._SKILLS) >= 589


def test_skills_library_has_at_least_604():
    from app import skills
    assert len(skills._SKILLS) >= 604


def test_skills_library_has_at_least_608():
    from app import skills
    assert len(skills._SKILLS) >= 608


def test_skills_library_has_at_least_618():
    from app import skills
    assert len(skills._SKILLS) >= 618


def test_skills_library_has_at_least_628():
    from app import skills
    assert len(skills._SKILLS) >= 628


def test_skills_library_has_at_least_638():
    from app import skills
    assert len(skills._SKILLS) >= 638


def test_skills_library_has_at_least_646():
    from app import skills
    assert len(skills._SKILLS) >= 646


def test_skills_library_has_at_least_654():
    from app import skills
    assert len(skills._SKILLS) >= 654


def test_skills_library_has_at_least_660():
    from app import skills
    assert len(skills._SKILLS) >= 660


def test_skills_library_has_at_least_669():
    from app import skills
    assert len(skills._SKILLS) >= 669


def test_skills_library_has_at_least_677():
    from app import skills
    assert len(skills._SKILLS) >= 677


def test_skills_library_has_at_least_685():
    from app import skills
    assert len(skills._SKILLS) >= 685


def test_skills_library_has_at_least_694():
    from app import skills
    assert len(skills._SKILLS) >= 694


def test_skills_library_has_at_least_696():
    from app import skills
    assert len(skills._SKILLS) >= 696


def test_skills_library_has_at_least_703():
    from app import skills
    assert len(skills._SKILLS) >= 703


def test_skills_library_has_at_least_711():
    from app import skills
    assert len(skills._SKILLS) >= 711


def test_skills_library_has_at_least_718():
    from app import skills
    assert len(skills._SKILLS) >= 718


def test_skills_library_has_at_least_726():
    from app import skills
    assert len(skills._SKILLS) >= 726


def test_skills_library_has_at_least_728():
    from app import skills
    assert len(skills._SKILLS) >= 728


def test_rag_semantic_lookup_ranks_by_meaning():
    """🔎 RAG thật: cosine + semantic_lookup xếp hạng theo NGỮ NGHĨA (không cần API key)."""
    from app import knowledge
    from app.embeddings import cosine

    # cosine thuần: đồng hướng = 1, vuông góc = 0, chiều lệch = 0.
    assert abs(cosine([1.0, 0.0], [1.0, 0.0]) - 1.0) < 1e-9
    assert abs(cosine([1.0, 0.0], [0.0, 1.0])) < 1e-9
    assert cosine([1.0, 0.0], []) == 0.0

    # Lưu 2 fact + gán vector thủ công (mô phỏng embeddings, không gọi mạng).
    knowledge.remember("mèo nhà", "Mèo là thú cưng phổ biến.", source="manual")
    knowledge.remember("động cơ đốt trong", "Nguyên lý xi-lanh và piston.", source="manual")
    conn = knowledge._get_conn()
    with knowledge._lock:
        cat = conn.execute("SELECT id FROM facts WHERE topic='mèo nhà'").fetchone()[0]
        eng = conn.execute("SELECT id FROM facts WHERE topic='động cơ đốt trong'").fetchone()[0]
    knowledge.store_vector(cat, [1.0, 0.0, 0.0])
    knowledge.store_vector(eng, [0.0, 1.0, 0.0])

    # Truy vấn gần "mèo" (vector [0.9,0.1,0]) → phải trả 'mèo nhà' đứng đầu.
    res = knowledge.semantic_lookup([0.9, 0.1, 0.0], limit=2, min_sim=0.5)
    assert res and res[0]["topic"] == "mèo nhà"
    assert "sim" in res[0]


def test_skills_match_english_queries_too():
    """Truy vấn TIẾNG ANH cũng khớp kỹ năng (fallback token tên slug), không phá khớp tiếng Việt."""
    from app import skills
    cases = {
        "how to do test driven development in python": "test-driven-development",
        "what is graphrag": "how-graphrag-works",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)
    # Câu tiếng Việt vẫn khớp đúng (giai đoạn 1 ưu tiên tuyệt đối).
    vi = skills.find_matching_skill("làm sao viết test trước khi code")
    assert vi is not None and vi.slug == "test-driven-development"


def test_skills_library_has_at_least_754():
    from app import skills
    assert len(skills._SKILLS) >= 754


def test_skills_repo_game_batch_topics_match():
    from app import skills
    cases = {
        "tạo asset game 2d bằng ai spritesheet chroma-key tách frame căn pivot tilemap": "ai-game-asset-generation",
        "lặp cải thiện dựa trên kết quả chạy thật đánh giá từ sản phẩm đang chạy không tin compile sạch": "proof-driven-iteration",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_763():
    from app import skills
    assert len(skills._SKILLS) >= 763


def test_skills_distributed_reliability_batch_topics_match():
    from app import skills
    cases = {
        "ghi log trước khi ghi dữ liệu wal đảm bảo bền vững sau khi sập máy redo và undo khi khởi động lại": "write-ahead-logging",
        "nhiều phiên bản dữ liệu mvcc cô lập ảnh chụp snapshot isolation đọc không chặn ghi write skew": "mvcc-and-snapshot-isolation",
        "giao thức lan truyền gossip kiểu dịch tễ phát hiện lỗi node bằng heartbeat swim phi-accrual": "gossip-and-failure-detection",
        "cây băm merkle tree hash phân cấp chống phân kỳ replica anti-entropy so sánh gốc rồi đi xuống nhánh khác nhau": "merkle-trees-and-anti-entropy",
        "độ trễ đuôi tail latency p99 request dự phòng hedged gửi bản sao thứ hai fan-out khuếch đại một thành phần chậm": "tail-latency-and-hedged-requests",
        "đồng hồ logic lamport và vector clock phát hiện cập nhật đồng thời concurrent theo quan hệ nhân quả happens-before": "vector-clocks-and-causality",
        "giảm tải chủ động load shedding từ chối bớt việc suy giảm nhẹ nhàng graceful degradation admission control": "load-shedding-and-graceful-degradation",
        "token phân định fencing tăng đơn điệu khoá phân tán không đủ an toàn não phân đôi split-brain hai leader cùng ghi": "fencing-tokens-and-split-brain",
        "đồng bộ đồng hồ vật lý ntp có sai số truetime commit-wait và hybrid logical clock hlc chống lệch đồng hồ clock skew": "clock-synchronization-and-hybrid-logical-clocks",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_771():
    from app import skills
    assert len(skills._SKILLS) >= 771


def test_skills_lowlevel_batch_topics_match():
    from app import skills
    cases = {
        "thuật toán không khoá lock-free và wait-free atomic compare-and-swap cas vấn đề aba con trỏ tái sử dụng tránh deadlock": "lock-free-and-wait-free-algorithms",
        "false sharing hai biến khác nhau chung một cache line giao thức nhất quán cache mesi ping-pong giữa nhân padding căn chỉnh tách cache line": "false-sharing-and-cache-line-contention",
        "đọc sao chép cập nhật rcu cho dữ liệu đọc nhiều ghi ít người đọc không khoá giai đoạn ân hạn grace period": "rcu-read-copy-update",
        "dạng gán đơn tĩnh ssa mỗi biến gán đúng một lần phi-node hợp giá trị tại điểm nhập luồng điều khiển truyền hằng số": "ssa-form-and-compiler-optimization",
        "cấp phát thanh ghi register allocation bằng tô màu đồ thị graph coloring biến sống cùng lúc giao nhau tràn thanh ghi spill khoảng sống live range": "register-allocation-and-graph-coloring",
        "nội tuyến inlining thay lời gọi bằng thân hàm phân tích thoát escape analysis đối tượng không rời phạm vi không thoát cấp phát trên stack thay vì heap": "escape-analysis-and-inlining",
        "lời gọi đuôi tail call và tối ưu tco chạy đệ quy trong không gian stack hằng số tiếp diễn continuation async await generator": "tail-calls-and-continuations",
        "profile-guided optimization pgo tối ưu theo hồ sơ chạy thật thu thập profile nhánh nóng tách mã nóng lạnh hot cold splitting tối ưu trường hợp phổ biến common case": "profile-guided-optimization",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_779():
    from app import skills
    assert len(skills._SKILLS) >= 779


def test_skills_net_stream_batch_topics_match():
    from app import skills
    cases = {
        "điều khiển tắc nghẽn tcp chia sẻ băng thông cửa sổ tắc nghẽn congestion window khởi động chậm slow start rồi tăng cộng giảm nhân aimd thuật toán cubic và bbr": "tcp-congestion-control",
        "chặn đầu hàng head-of-line blocking một mục kẹt làm nghẽn cả hàng segment tcp mất làm đơ mọi luồng http2 http3 chuyển sang quic udp nhiều hàng đợi riêng thay vì một hàng chung": "head-of-line-blocking",
        "phình đệm bufferbloat buffer router quá lớn giữ hàng đợi khổng lồ quản lý hàng đợi chủ động aqm codel fq-codel pie độ trễ tăng vọt khi tải nặng": "bufferbloat-and-active-queue-management",
        "thời gian sự kiện event-time và thời gian xử lý processing-time watermark ước lượng đã thấy hết sự kiện tới mốc t xử lý dữ liệu trễ allowed lateness": "event-time-and-watermarks",
        "cửa sổ hoá stream windowing gom luồng vô hạn thành khối hữu hạn cửa sổ cố định không chồng lấn tumbling cửa sổ trượt chồng lấn sliding cửa sổ phiên session window": "windowing-in-stream-processing",
        "xử lý luồng đúng một lần exactly-once khác biệt với at-least-once và at-most-once khử trùng lặp deduplication idempotent sink chốt checkpoint và offset nguyên tử": "exactly-once-stream-processing",
        "xử lý luồng có trạng thái stateful nhớ qua nhiều sự kiện trạng thái theo khoá keyed state chốt điểm checkpoint và ảnh chụp snapshot để chịu lỗi savepoint nâng cấp": "stateful-stream-processing-and-checkpointing",
        "mô hình dataflow batch và streaming là cùng một phép tính kiến trúc lambda gồm lớp batch và lớp tốc độ kiến trúc kappa một pipeline streaming tái xử lý bằng replay": "dataflow-and-stream-batch-unification",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_787():
    from app import skills
    assert len(skills._SKILLS) >= 787


def test_skills_applied_crypto_batch_topics_match():
    from app import skills
    cases = {
        "mã hoá có xác thực aead vừa bảo mật vừa toàn vẹn aes-gcm và chacha20-poly1305 encrypt-then-mac chế độ không xác thực aes-cbc nguy hiểm dữ liệu liên kết associated data": "authenticated-encryption-aead",
        "hàm dẫn xuất khoá kdf biến một bí mật thành nhiều khoá hkdf trích rồi mở rộng extract-then-expand không dùng thẳng shared secret hay mật khẩu làm khoá tách miền domain separation": "key-derivation-functions",
        "mã hoá phong bì envelope encryption khoá dữ liệu và khoá bọc data key dek mã hoá dữ liệu kek bọc lại dek khoá gốc không bao giờ rời kms hsm xoay khoá rẻ chỉ bọc lại dek": "envelope-encryption-and-kms",
        "trao đổi khoá diffie-hellman ecdh trên kênh công khai bí mật chuyển tiếp forward secrecy khoá tạm ephemeral lộ khoá riêng dài hạn không giải mã được lưu lượng cũ pfs tls 1.3": "forward-secrecy-and-key-exchange",
        "kênh phụ side-channel rò rỉ bí mật dù crypto đúng tấn công thời gian timing attack so sánh mac thoát sớm lập trình thời gian hằng số constant-time không rẽ nhánh theo bí mật": "constant-time-code-and-side-channels",
        "quản lý nonce và iv số dùng một lần trong crypto đối xứng tái dùng nonce cùng khoá phá vỡ aes-gcm nghiêm trọng nonce ngẫu nhiên hay bộ đếm counter giới hạn 96-bit gcm": "nonce-and-iv-management",
        "hạ tầng khoá công khai pki chuỗi tin cậy chứng chỉ tls ca gốc và ca trung gian ký chứng chỉ thu hồi chứng chỉ crl và ocsp minh bạch chứng chỉ certificate transparency log chỉ thêm": "certificate-transparency-and-pki",
        "linh hoạt mật mã cryptographic agility đổi thuật toán không phải xây lại mọi thuật toán rồi cũng yếu md5 sha-1 ciphertext tự mô tả có mã định danh thuật toán và phiên bản chuyển đổi hậu lượng tử post-quantum": "cryptographic-agility-and-algorithm-migration",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_795():
    from app import skills
    assert len(skills._SKILLS) >= 795


def test_skills_ml_systems_batch_topics_match():
    from app import skills
    cases = {
        "gộp lô liên tục continuous batching phục vụ llm thông lượng cao batch tĩnh phải chờ chuỗi dài nhất lãng phí gpu lập lịch theo từng bước iteration-level thêm và loại chuỗi mỗi bước decode": "continuous-batching-for-llm-serving",
        "paged attention quản lý kv cache như bộ nhớ ảo hệ điều hành kv cache là nút cổ chai bộ nhớ khối trang cố định thay vì bộ đệm liền mạch chia sẻ tiền tố prefix sharing": "paged-attention-and-kv-cache-memory",
        "flash attention chú ý nhận biết io nhanh và tiết kiệm bộ nhớ không tạo ma trận chú ý n bình phương đầy đủ chia ô tiling và softmax trực tuyến online tính lại recomputation": "flash-attention",
        "song song hoá mô hình quá lớn cho một gpu song song dữ liệu data parallel chia batch song song tensor chia ma trận trong một tầng song song đường ống pipeline phân mảnh trạng thái fsdp zero": "tensor-and-pipeline-parallelism",
        "bộ nhớ huấn luyện bị chi phối bởi activation không phải trọng số checkpointing gradient đánh đổi tính toán lấy bộ nhớ chỉ lưu vài activation rồi tính lại phần còn lại khi backprop": "gradient-checkpointing-and-activation-memory",
        "huấn luyện độ chính xác hỗn hợp mixed precision nhanh gấp đôi nửa bộ nhớ fp16 tràn dưới gradient cần nhân tỉ lệ loss scaling bf16 dễ hơn cùng dải mũ với fp32": "mixed-precision-training",
        "mở rộng độ dài ngữ cảnh llm bằng điều chỉnh rope rotary position mô hình hỏng ngoài độ dài đã huấn luyện nội suy vị trí position interpolation ntk-aware và yarn": "context-length-extension-and-rope-scaling",
        "hai pha suy luận llm prefill nghẽn tính toán và decode nghẽn bộ nhớ tách riêng disaggregation prefill và decode lên tài nguyên khác nhau thời gian tới token đầu ttft chunked prefill": "prefill-and-decode-disaggregation",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_803():
    from app import skills
    assert len(skills._SKILLS) >= 803


def test_skills_os_kernel_batch_topics_match():
    from app import skills
    cases = {
        "bộ đệm trang page cache kernel giữ trang tệp trong ram đọc lại tệp gần như tức thì không chạm đĩa ánh xạ tệp vào bộ nhớ mmap trang bẩn dirty page và ghi lại writeback": "page-cache-and-memory-mapped-files",
        "mô hình vào ra io blocking và non-blocking sẵn sàng readiness select poll epoll so với hoàn tất completion io_uring vấn đề c10k một luồng mỗi kết nối không mở rộng": "io-models-and-io-uring",
        "sao chép khi ghi copy-on-write cow tiến trình con dùng chung trang chỉ đọc fork tạo tiến trình mới gần như miễn phí chỉ sao chép trang khi có bên ghi vào redis fork để chụp nhanh": "copy-on-write-and-fork",
        "vào ra không sao chép zero-copy sendfile splice gửi tệp qua socket kiểu thường sao chép bốn lần loại bỏ bản sao thừa và vượt biên user kernel kafka nginx nhanh nhờ zero-copy": "zero-copy-io",
        "ngắt interrupt thiết bị báo xong thay vì cpu chờ thăm dò polling lãng phí cpu dma thiết bị tự chuyển dữ liệu vào ra bộ nhớ không cần cpu sao chép bão ngắt interrupt storm và napi": "interrupts-dma-and-device-io",
        "truy cập bộ nhớ không đồng nhất numa trên máy nhiều socket bộ nhớ node xa chậm hơn node cục bộ cấp phát theo lần chạm đầu first-touch ghim luồng và bộ nhớ cùng một node": "numa-and-memory-locality",
        "chi phí chuyển ngữ cảnh context switch giữa luồng nhiều luồng hơn số nhân làm chậm chương trình chi phí gián tiếp cache và tlb nguội sau khi chuyển quá tải luồng oversubscription": "context-switching-cost",
        "trang lớn huge page và bộ đệm dịch địa chỉ tlb trang 4kb khiến tập làm việc lớn tràn tlb gây miss đi bộ bảng trang page walk tốn kém transparent huge page thp": "huge-pages-and-tlb",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_811():
    from app import skills
    assert len(skills._SKILLS) >= 811


def test_skills_db_engine_batch_topics_match():
    from app import skills
    cases = {
        "bể đệm buffer pool cơ sở dữ liệu cache trang đĩa trong ram chính sách thay thế trang page replacement ghim trang pinning và trang bẩn dirty page lru hỏng khi quét lớn sequential flooding": "buffer-pool-and-page-replacement",
        "ước lượng lực lượng cardinality estimation đoán số dòng mỗi bước thống kê bảng histogram và số giá trị phân biệt distinct ước lượng sai gây kế hoạch thảm hoạ giả định độc lập sai": "cardinality-estimation-and-statistics",
        "thuật toán join vật lý nested-loop hash join sort-merge hash join tốt nhất cho equi-join lớn không sắp xếp nested loop trên bảng lớn là thảm hoạ": "join-algorithms-hash-merge-nested-loop",
        "đẩy điều kiện lọc predicate pushdown xuống gần nguồn dữ liệu nhất cắt tỉa cột projection pruning chỉ đọc cột cần cắt phân vùng partition pruning đọc ít byte từ kho cột parquet": "predicate-pushdown-and-projection-pruning",
        "khoá hai pha two-phase locking 2pl đảm bảo giao dịch tuần tự hoá pha tăng lấy khoá và pha giảm nhả khoá bế tắc deadlock và phát hiện bằng đồ thị chờ wait-for": "two-phase-locking-and-deadlocks",
        "mức cô lập isolation level và các bất thường đồng thời đọc bẩn dirty read đọc không lặp lại non-repeatable phantom mất cập nhật lost update và write skew read committed serializable": "isolation-levels-and-anomalies",
        "index bao phủ covering index và quét chỉ mục index-only scan index chứa đủ mọi cột truy vấn cần bỏ qua tra bảng cột include đính kèm payload ngoài khoá": "covering-indexes-and-index-only-scans",
        "biên dịch truy vấn và sinh mã code generation cho engine cơ sở dữ liệu mô hình volcano diễn giải từng dòng tốn chi phí thực thi vector hoá xử lý lô cột qua toán tử biên dịch jit sinh mã máy": "query-compilation-and-code-generation",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_819():
    from app import skills
    assert len(skills._SKILLS) >= 819


def test_skills_compression_search_batch_topics_match():
    from app import skills
    cases = {
        "mã hoá entropy gán mã ngắn cho ký hiệu hay gặp mã dài cho hiếm mã huffman xây cây theo tần suất entropy shannon là giới hạn nén lý thuyết mã tiền tố prefix-free": "huffman-and-entropy-coding",
        "nén từ điển lz77 lz78 thay chuỗi lặp bằng tham chiếu ngược cửa sổ trượt sliding window tham chiếu khoảng cách và độ dài distance length từ điển nạp sẵn preset dictionary": "lz77-and-dictionary-compression",
        "mã hoá số học arithmetic coding mã hoá bằng số bit phân số vượt giới hạn số bit nguyên của huffman biểu diễn cả thông điệp thành một số ans asymmetric numeral systems zstd": "arithmetic-and-range-coding",
        "mã hoá độ dài chạy run-length thay chuỗi giá trị lặp bằng giá trị và số đếm mã hoá delta lưu hiệu giữa các giá trị liên tiếp delta cộng zigzag cho số nguyên và dấu thời gian": "run-length-and-delta-encoding",
        "chỉ mục ngược inverted index ánh xạ mỗi từ tới danh sách tài liệu chứa nó danh sách postings truy vấn boolean và cụm từ bằng giao danh sách postings like phần trăm không mở rộng": "inverted-index-and-full-text-search",
        "xếp hạng liên quan tf-idf và bm25 cho tìm kiếm tần suất từ term frequency tần suất tài liệu nghịch idf từ hiếm quan trọng hơn bm25 bão hoà tf và chuẩn hoá độ dài tài liệu": "tf-idf-and-bm25-ranking",
        "nén danh sách postings chỉ mục tìm kiếm postings sắp xếp theo doc id lưu khoảng cách gap delta bit-packing variable-byte frame-of-reference pfordelta con trỏ nhảy skip pointer nén chỉ mục tăng tốc truy vấn nhờ ít i/o": "posting-list-compression",
        "phân tích văn bản tokenization tách token và stemming cho tìm kiếm chuẩn hoá lowercase bỏ dấu accent phân tích lúc lập chỉ mục và lúc truy vấn phải khớp tách từ ngôn ngữ cjk": "text-analysis-tokenization-and-stemming",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_827():
    from app import skills
    assert len(skills._SKILLS) >= 827


def test_skills_numerical_probabilistic_batch_topics_match():
    from app import skills
    cases = {
        "ổn định số học và triệt tiêu thảm hoạ catastrophic cancellation trừ hai số gần bằng nhau mất hết chữ số có nghĩa tái công thức ổn định log-sum-exp welford kahan summation": "numerical-stability-and-cancellation",
        "phương pháp monte carlo giải bằng lấy mẫu ngẫu nhiên ước lượng tích phân xác suất kỳ vọng bằng trung bình nhiều mẫu sai số giảm theo một trên căn n bất kể số chiều": "monte-carlo-methods",
        "count-min sketch đếm tần suất tỉ item trong bộ nhớ cố định nhỏ cấu trúc xác suất nhiều hàm băm và mảng đếm hai chiều tìm phần tử nổi trội heavy hitters top-k": "count-min-sketch",
        "hyperloglog đếm số phần tử phân biệt distinct trong luồng lớn bằng vài kilobyte ước lượng lực lượng cardinality từ số bit 0 dẫn đầu tối đa trong hash gộp merge nhiều shard": "hyperloglog-cardinality-estimation",
        "chỉ mục không gian spatial index cho truy vấn gần đây và trong vùng quadtree chia ô đệ quy và k-d tree chia theo trục r-tree bao đối tượng trong hình chữ nhật lồng nhau bounding box": "spatial-indexing-quadtrees-and-r-trees",
        "geohash biến toạ độ 2d thành chuỗi số một chiều sắp xếp được đan xen bit vĩ độ kinh độ theo đường z-order morton curve điểm gần nhau thường chung tiền tố prefix geohash": "geohashing-and-spatial-hashing",
        "minhash ước lượng độ tương đồng jaccard của tập từ chữ ký nhỏ locality-sensitive hashing lsh băm item giống nhau vào cùng thùng tìm gần trùng near-duplicate không so sánh mọi cặp": "minhash-and-locality-sensitive-hashing",
        "sketch phân vị và percentile xấp xỉ p99 p999 trong bộ nhớ nhỏ t-digest ddsketch gk sai số có chặn không được lấy trung bình của percentile gộp merge nhiều shard": "quantile-sketches-and-approximate-percentiles",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_835():
    from app import skills
    assert len(skills._SKILLS) >= 835


def test_skills_data_format_batch_topics_match():
    from app import skills
    cases = {
        "protocol buffers protobuf mã hoá dữ liệu có cấu trúc thành nhị phân nhỏ gọn định dạng wire tag-length-value theo số hiệu trường không phải tên số hiệu field number nằm trên wire": "protobuf-and-wire-format",
        "số nguyên độ dài biến đổi varint lưu số bằng ít byte nhất bảy bit mỗi byte cộng bit tiếp diễn continuation bit zigzag đan xen số dương và âm để số âm nhỏ vẫn ngắn": "varint-and-zigzag-encoding",
        "tuần tự hoá nhị phân tự mô tả không cần schema messagepack cbor nhỏ và nhanh hơn json mà vẫn linh hoạt gắn nhãn kiểu dữ liệu gọn trong byte đầu": "binary-serialization-msgpack-and-cbor",
        "giải tuần tự không sao chép zero-copy flatbuffers capnproto đọc dữ liệu thẳng trong buffer không cần bước parse giải mã truy cập trường theo offset qua bảng vtable": "flatbuffers-and-zero-copy-deserialization",
        "apache avro tuần tự hoá bản ghi nhỏ gọn và tiến hoá schema an toàn avro không gắn tag trường inline schema quyết định bố cục schema registry lưu schema theo id và ép tương thích": "avro-and-schema-registry",
        "json schema định nghĩa mô tả và ép kiểu dữ liệu json khai báo kiểu trường bắt buộc và ràng buộc kiểm định ở ranh giới tin cậy validate boundary additionalproperties kiểm soát trường thừa": "json-schema-and-validation",
        "thứ tự byte endianness big-endian và little-endian cùng byte khác số network byte order phải cố định thứ tự byte trong định dạng tệp và mạng bố cục struct căn chỉnh alignment và đệm padding": "endianness-and-binary-layout",
        "thương lượng nội dung content negotiation phục vụ đúng định dạng ngôn ngữ mã hoá header accept và accept-language kiểu media mime type content-type header vary quan trọng cho cache": "content-negotiation-and-media-types",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_843():
    from app import skills
    assert len(skills._SKILLS) >= 843


def test_skills_durable_game_batch_topics_match():
    from app import skills
    cases = {
        "thực thi bền vững durable execution quy trình chạy dài chịu được sập máy khởi động lại engine lưu từng bước để workflow tiếp tục đúng chỗ đã dừng workflow và activity temporal": "durable-execution-and-workflow-engines",
        "tính tất định determinism của workflow phát lại replay dựng lại trạng thái từ lịch sử sự kiện event-sourced mọi ngẫu nhiên đồng hồ i/o trực tiếp phá vỡ replay side effect phải nằm trong activity": "workflow-determinism-and-replay",
        "chính sách thử lại retry với backoff cho từng bước activity họ timeout schedule-to-start start-to-close heartbeat phát hiện bước chạy dài bị treo idempotency bắt buộc khi có retry": "retries-timeouts-and-heartbeats",
        "tương tác với workflow đang chạy signal query signal gửi sự kiện đầu vào bất đồng bộ query đọc trạng thái hiện tại mà không thay đổi mẫu chờ điều kiện wait-for-condition human-in-the-loop": "signals-queries-and-workflow-interaction",
        "vòng lặp game game loop vật lý cần bước thời gian cố định fixed timestep còn render chạy nhanh tuỳ máy di chuyển theo delta-time độc lập khung hình mẫu accumulator nội suy interpolation": "game-loop-and-fixed-timestep",
        "phát hiện va chạm collision detection và phản hồi kiểm tra hộp bao aabb broad phase narrow phase phân vùng không gian va chạm quét liên tục swept chống xuyên tường tunneling giải quyết tách trục axis-separated để trượt": "collision-detection-and-response",
        "chuyển động không tuyến tính đường cong easing và nhịp animation tự nhiên ease-out ease-in ease-in-out dùng khi nào cubic-bezier và vật lý lò xo spring so le stagger reduced-motion": "easing-and-animation-timing",
        "gpu vẽ mọi thứ trên web pipeline webgl opengl và shader vertex shader định vị mỗi đỉnh fragment shader tô màu mỗi điểm ảnh buffer attribute uniform varying chuỗi ma trận mvp gộp draw call": "webgl-and-shader-fundamentals",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_851():
    from app import skills
    assert len(skills._SKILLS) >= 851


def test_skills_web_security_batch_topics_match():
    from app import skills
    cases = {
        "giả mạo yêu cầu liên trang csrf cookie tự gửi kèm nên nguy hiểm cho request thay đổi trạng thái token chống csrf synchronizer và double-submit cookie samesite lax strict": "csrf-and-same-site-defenses",
        "giả mạo yêu cầu phía máy chủ ssrf server tự gọi url do người dùng nhập tấn công endpoint metadata đám mây 169.254.169.254 phòng thủ bằng allowlist và chặn egress mạng": "ssrf-and-url-fetch-safety",
        "chèn mã kịch bản liên trang xss dữ liệu kẻ tấn công thành script chạy trong trình duyệt mã hoá đầu ra output encoding theo ngữ cảnh không phải lọc đầu vào xss lưu trữ phản chiếu và dom": "xss-and-output-encoding",
        "tiêm sql injection dữ liệu kẻ tấn công biến câu truy vấn thành vũ khí sửa thật bằng truy vấn tham số hoá prepared statement tách mã khỏi dữ liệu nối chuỗi là gốc rễ lỗ hổng": "sql-injection-and-parameterized-queries",
        "cướp nhấp chuột clickjacking ui redressing nhúng site vào iframe trong suốt đè lên nội dung mồi phòng thủ x-frame-options và csp frame-ancestors samesite không ngăn được clickjacking": "clickjacking-and-frame-protection",
        "cố định phiên session fixation kẻ tấn công kiểm soát session id trước rồi chiếm đăng nhập tái tạo session id khi đổi quyền đăng nhập cờ cookie an toàn httponly secure samesite": "session-fixation-and-secure-sessions",
        "kiểm soát truy cập bị hỏng broken access control idor tham chiếu trực tiếp đối tượng đổi id trong request để xem dữ liệu người khác phân quyền phía máy chủ trên mọi request giấu id hay dùng uuid không phải bảo mật": "idor-and-broken-access-control",
        "giải tuần tự không an toàn insecure deserialization biến byte không tin cậy thành đối tượng có thể chạy mã chuỗi gadget chain thực thi mã từ xa rce serializer nhị phân java pickle php .net nguy hiểm": "insecure-deserialization-and-gadget-chains",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_859():
    from app import skills
    assert len(skills._SKILLS) >= 859


def test_skills_download_crawl_batch_topics_match():
    from app import skills
    cases = {
        "tải phân đoạn song song tăng tốc download chia tệp thành nhiều dải byte tải qua nhiều kết nối http range request ghi mỗi đoạn vào đúng offset trong tệp không cần bước gộp merge": "segmented-parallel-downloading",
        "tải tiếp resumable download tiếp tục từ chỗ đã dừng http range request tải từ offset đã lưu if-range etag last-modified phát hiện tệp đã thay đổi để khởi động lại không nối dữ liệu hỏng": "resumable-downloads-and-range-requests",
        "độ tin cậy trình quản lý tải xuống thử lại với backoff khi lỗi tạm thời hoàn tất nguyên tử ghi tệp tạm rồi đổi tên lưu trạng thái phần đã tải để sống sót qua khởi động lại app": "download-manager-reliability",
        "phát trực tuyến thích ứng hls m3u8 và mpeg-dash chia video thành nhiều segment nhỏ manifest playlist master và media tải các segment rồi ghép hoặc remux thành một tệp playlist trực tiếp live và vod": "hls-dash-segment-downloading",
        "giới hạn băng thông throttling tự giới hạn tốc độ tải bằng token bucket lập lịch tải hàng đợi giới hạn số tải đồng thời chia ngân sách băng thông chung cho nhiều kết nối song song": "bandwidth-throttling-and-download-scheduling",
        "sao lưu toàn bộ website để dùng offline mirror kiểu wget httrack thu thập đệ quy theo liên kết crawling viết lại url thành đường dẫn tương đối để bản sao chạy offline giới hạn phạm vi cùng tên miền và độ sâu": "website-mirroring-and-crawling",
        "crawler lịch sự tránh gây hại và bị chặn tuân thủ robots.txt allow disallow crawl-delay giới hạn tốc độ và số kết nối đồng thời mỗi máy chủ tôn trọng 429 và retry-after": "crawler-politeness-and-robots-txt",
        "chuẩn hoá url canonicalization để nhận ra các url khác nhau trỏ cùng một trang khử trùng lặp tránh tải lại vô tận và bẫy crawler bỏ fragment sắp xếp query thẻ canonical và băm nội dung": "url-normalization-and-crawl-dedup",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_867():
    from app import skills
    assert len(skills._SKILLS) >= 867


def test_skills_agent_infra_batch_topics_match():
    from app import skills
    cases = {
        "kiến trúc agent lập trình tự động coding agent vòng lặp quan sát suy nghĩ hành động observe-think-act trên workspace sandbox bộ công cụ đọc sửa file chạy lệnh chạy test điều kiện dừng": "coding-agent-architecture",
        "thực thi mã trong sandbox cho agent chạy code an toàn container microvm cách ly filesystem và mạng giới hạn tài nguyên không truy cập host môi trường tạm ephemeral mỗi tác vụ kiểm soát egress": "sandboxed-code-execution-for-agents",
        "giao diện agent-máy tính aci thiết kế bộ công cụ và quan sát cho model không phải cho người đầu ra công cụ ngắn gọn lan can guardrail chặn lỗi agent phản hồi sau mỗi hành động": "agent-computer-interface",
        "agent điều khiển trình duyệt thật làm tác vụ web browser-use playwright vòng lặp cảm nhận quyết định hành động bộ hành động click gõ điều hướng cuộn chờ nội dung động selector dễ vỡ chống bot": "browser-automation-for-agents",
        "cảm nhận trang web cho agent biến trang thành thứ llm hiểu html thô quá lớn và nhiễu dùng cây trợ năng accessibility tree trích và đánh số phần tử tương tác indexing interactive elements": "web-perception-and-dom-for-agents",
        "agent dùng máy tính bằng thị giác gui agent chụp màn hình suy luận rồi xuất thao tác chuột bàn phím theo toạ độ điểm ảnh định vị toạ độ grounding set-of-mark đánh dấu phần tử": "computer-use-and-gui-agents",
        "agent sửa file tin cậy bài toán biểu diễn chỉnh sửa ghi đè cả file so với diff hợp nhất so với khối tìm và thay thế search-replace định dạng diff theo số dòng dễ vỡ với llm công cụ edit nên từ chối sửa sai": "agent-file-editing-and-diffs",
        "an toàn agent hành động khác với an toàn nội dung đầu ra lan can hành động action guardrails phạm vi quyền tối thiểu least privilege người xác nhận human-in-the-loop cho hành động hệ trọng không thể hoàn tác": "agent-safety-and-action-guardrails",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_875():
    from app import skills
    assert len(skills._SKILLS) >= 875


def test_skills_build_and_deps_batch_topics_match():
    from app import skills
    cases = {
        "build tăng dần incremental chỉ dựng lại phần đã thay đổi và build kín hermetic cho kết quả giống nhau mọi nơi đồ thị build target và phụ thuộc phát hiện thay đổi theo nội dung băm khai báo đủ input": "incremental-and-hermetic-builds",
        "đánh đổi giữa một kho lớn monorepo và nhiều kho nhỏ polyrepo thay đổi nguyên tử xuyên nhiều project sở hữu độc lập monorepo cần công cụ phát hiện target bị ảnh hưởng build caching code owners": "monorepo-and-polyrepo-tooling",
        "đánh số phiên bản ngữ nghĩa semantic versioning major minor patch để truyền đạt tương thích thế nào là thay đổi phá vỡ breaking change định luật hyrum toán tử khoảng phiên bản caret tilde": "semantic-versioning-and-compatibility",
        "giải phụ thuộc dependency resolution biến các khoảng phiên bản mong muốn thành một tập gói cụ thể tái lập được bài toán thoả mãn ràng buộc kiểu sat lockfile ghim đồ thị đã giải chính xác commit lockfile": "dependency-resolution-and-lockfiles",
        "địa ngục phụ thuộc dependency hell và bài toán phụ thuộc kim cương diamond hai phụ thuộc cần phiên bản không tương thích của một thư viện chung native toàn cục chỉ một phiên bản thắng": "dependency-hell-and-diamond-dependencies",
        "tốc độ vòng lặp sửa-thấy-kết-quả quyết định năng suất lập trình hot reload thay module nóng hot module replacement biên dịch tăng dần giữ nguyên trạng thái khi sửa dev server vite webpack": "hot-reload-and-fast-feedback",
        "liên kết linking và loader biến các mã biên dịch riêng thành một chương trình chạy được liên kết tĩnh static và động dynamic phân giải ký hiệu symbol resolution lỗi symbol not found dll hell": "linking-and-loaders",
        "tiêm phụ thuộc dependency injection và đảo ngược điều khiển inversion of control ioc không để đối tượng tự tạo phụ thuộc mà truyền vào constructor injection lập trình theo giao diện interface di giúp test và mock": "dependency-injection-and-inversion-of-control",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_883():
    from app import skills
    assert len(skills._SKILLS) >= 883


def test_skills_library_has_at_least_891():
    from app import skills
    assert len(skills._SKILLS) >= 891


def test_skills_library_has_at_least_899():
    from app import skills
    assert len(skills._SKILLS) >= 899


def test_skills_library_has_at_least_907():
    from app import skills
    assert len(skills._SKILLS) >= 907


def test_skills_library_has_at_least_915():
    from app import skills
    assert len(skills._SKILLS) >= 915


def test_skills_library_has_at_least_923():
    from app import skills
    assert len(skills._SKILLS) >= 923


def test_skills_library_has_at_least_931():
    from app import skills
    assert len(skills._SKILLS) >= 931


def test_skills_library_has_at_least_939():
    from app import skills
    assert len(skills._SKILLS) >= 939


def test_skills_library_has_at_least_947():
    from app import skills
    assert len(skills._SKILLS) >= 947


def test_skills_library_has_at_least_955():
    from app import skills
    assert len(skills._SKILLS) >= 955


def test_skills_library_has_at_least_963():
    from app import skills
    assert len(skills._SKILLS) >= 963


def test_skills_library_has_at_least_971():
    from app import skills
    assert len(skills._SKILLS) >= 971


def test_skills_library_has_at_least_979():
    from app import skills
    assert len(skills._SKILLS) >= 979


def test_skills_library_has_at_least_987():
    from app import skills
    assert len(skills._SKILLS) >= 987


def test_skills_library_has_at_least_995():
    from app import skills
    assert len(skills._SKILLS) >= 995


def test_skills_library_has_at_least_1000():
    from app import skills
    assert len(skills._SKILLS) >= 1000


def test_skills_library_has_at_least_1013():
    from app import skills
    assert len(skills._SKILLS) >= 1013


def test_skills_library_has_at_least_1021():
    from app import skills
    assert len(skills._SKILLS) >= 1021


def test_skills_library_has_at_least_1029():
    from app import skills
    assert len(skills._SKILLS) >= 1029


def test_skills_library_has_at_least_1037():
    from app import skills
    assert len(skills._SKILLS) >= 1037


def test_skills_library_has_at_least_1061():
    from app import skills
    assert len(skills._SKILLS) >= 1061


def test_skills_library_has_at_least_1077():
    from app import skills
    assert len(skills._SKILLS) >= 1077


def test_skills_library_has_at_least_1087():
    from app import skills
    assert len(skills._SKILLS) >= 1087


def test_skills_library_has_at_least_1093():
    from app import skills
    assert len(skills._SKILLS) >= 1093


def test_skills_library_has_at_least_1099():
    from app import skills
    assert len(skills._SKILLS) >= 1099


def test_skills_prompt_optimization_and_briefing_topics_match():
    from app import skills
    for text, expected in [
        ("làm sao tối ưu prompt tự động theo thước đo bằng dspy", "automatic-prompt-optimization"),
        ("dùng prompt optimizer opro cải thiện chỉ dẫn", "automatic-prompt-optimization"),
        ("làm bản tin 30 ngày qua tổng hợp thay đổi gần đây", "last-30-days-briefing"),
        ("thiết kế vòng lặp agent plan act observe có ngân sách bước", "agent-loop-engineering"),
    ]:
        s = skills.find_matching_skill(text)
        assert s and s.slug == expected, (text, s.slug if s else None)


def test_skills_library_has_at_least_1100():
    from app import skills
    assert len(skills._SKILLS) >= 1100


def test_skills_library_has_at_least_1103():
    from app import skills
    assert len(skills._SKILLS) >= 1103


def test_skills_library_has_at_least_1106():
    from app import skills
    assert len(skills._SKILLS) >= 1106


def test_skills_library_has_at_least_1109():
    from app import skills
    assert len(skills._SKILLS) >= 1109


def test_skills_library_has_at_least_1114():
    from app import skills
    assert len(skills._SKILLS) >= 1114


def test_skills_advanced_gamedev_topics_match():
    from app import skills
    for text, expected in [
        ("tích hợp fmod wwise event bank rtpc âm thanh game", "audio-middleware-fmod-wwise"),
        ("sinh dungeon bằng bsp cellular automata nối phòng mst", "procedural-dungeon-and-level-generation"),
        ("tạo shader bằng node shader graph fresnel dissolve", "shader-graphs-and-node-based-materials"),
        ("game ai htn hierarchical task network compound task", "htn-hierarchical-task-network-planning"),
        ("tối ưu hiệu năng game giảm draw call gc spike profiling", "game-performance-optimization-and-profiling"),
    ]:
        s = skills.find_matching_skill(text)
        assert s and s.slug == expected, (text, s.slug if s else None)


def test_skills_unreal_pixi_gamemaker_topics_match():
    from app import skills
    for text, expected in [
        ("viết game unreal engine bằng c++ actor component gamemode pawn", "unreal-engine-cpp-gameplay"),
        ("render 2d webgl bằng pixijs stage sprite ticker batching", "pixijs-webgl-2d-rendering"),
        ("làm game 2d bằng gamemaker gml object event step draw alarm", "gamemaker-studio-and-gml"),
    ]:
        s = skills.find_matching_skill(text)
        assert s and s.slug == expected, (text, s.slug if s else None)


def test_skills_bevy_phaser_cocos_topics_match():
    from app import skills
    for text, expected in [
        ("làm game bằng rust bevy dùng system query commands spawn", "bevy-and-rust-game-development"),
        ("làm game web bằng phaser 3 preload create update arcade physics", "phaser-web-game-development"),
        ("làm game bằng cocos creator node component prefab typescript", "cocos-creator-game-development"),
    ]:
        s = skills.find_matching_skill(text)
        assert s and s.slug == expected, (text, s.slug if s else None)


def test_skills_game_engine_coding_topics_match():
    from app import skills
    for text, expected in [
        ("viết game bằng godot dùng gdscript node scene signal", "godot-engine-and-gdscript"),
        ("viết game unity dùng monobehaviour prefab coroutine rigidbody", "unity-engine-gameplay-scripting"),
        ("viết game bằng raylib tự viết game loop immediate mode", "raylib-and-immediate-mode-games"),
    ]:
        s = skills.find_matching_skill(text)
        assert s and s.slug == expected, (text, s.slug if s else None)


def test_recent_digest_detection_and_window():
    """'N ngày qua có gì mới' được nhận diện + rút đúng số ngày cửa sổ; câu số liệu
    quốc gia thường KHÔNG bị nhận nhầm là bản tin."""
    from app import knowledge as k
    assert k.is_recent_digest_query("AI có gì mới 30 ngày qua")
    assert k.is_recent_digest_query("what's new in AI last 7 days")
    assert not k.is_recent_digest_query("dân số việt nam bao nhiêu")
    assert k.recent_window_days("30 ngày qua") == 30
    assert k.recent_window_days("7 ngày qua") == 7
    assert k.recent_window_days("tuần qua") == 7
    assert k.recent_window_days("tháng qua") == 30
    assert k.recent_window_days("có gì mới") == 30          # mặc định
    assert 1 <= k.recent_window_days("999 ngày qua") <= 365  # chặn biên


def test_recent_digest_gather_offline_never_raises():
    """gather() với câu 'N ngày qua' KHÔNG được raise khi offline (nguồn phụ)."""
    import asyncio
    from app import knowledge
    out = asyncio.run(knowledge.gather("công nghệ AI có gì mới 30 ngày qua"))
    assert isinstance(out, list)


def test_skills_contest_and_motion_topics_match():
    from app import skills
    cases = {
        "lập trình thi đấu thi chuyên tin olympiad chiến lược giải bài và đọc ràng buộc độ phức tạp thuật toán và cấu trúc dữ liệu cốt lõi quy hoạch động thuật toán đồ thị tham lam greedy kỹ thuật thi codeforces": "competitive-programming-techniques",
        "hình học không gian nâng cao thi chuyên thpt đường thẳng và mặt phẳng trong không gian góc và khoảng cách khối đa diện lăng trụ chóp cầu thể tích và diện tích thiết diện phương pháp tọa độ": "advanced-solid-geometry",
        "hóa phân tích chuyên hóa phân tích định tính và định lượng chuẩn độ axit bazơ oxi hóa khử tạo phức phân tích khối lượng phổ và sắc ký đường chuẩn sai số và độ chính xác": "analytical-chemistry",
        "hoạt ảnh lottie và dotlottie định dạng lottie json xuất từ after effects bodymovin phát lottie trên web và di động định dạng dotlottie nhỏ gọn tương tác và điều khiển tối ưu": "lottie-and-dotlottie-animation",
        "kỹ thuật hoạt ảnh svg hoạt ảnh svg bằng css và js vẽ nét theo stroke dasharray biến hình morphing đường path so sánh smil css js chuyển động theo đường path hiệu năng": "svg-animation-techniques",
        "quy trình thiết kế chuyển động với ai từ thiết kế tĩnh hoặc prompt ra hoạt ảnh ui motion graphics chọn định dạng css gsap lottie video lập kế hoạch timeline và keyframe easing và biên đạo xuất json dotlottie": "ai-motion-design-workflow",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_calculus_science_topics_match():
    from app import skills
    cases = {
        "giới hạn và tính liên tục giải tích thpt khái niệm giới hạn hàm số giới hạn một bên và vô cực kỹ thuật tính giới hạn dạng vô định tính liên tục định lý giá trị trung gian": "limits-and-continuity",
        "đạo hàm và khảo sát hàm số thpt quy tắc tính đạo hàm hàm hợp tích thương tính đơn điệu và cực trị tính lồi lõm và điểm uốn vẽ đồ thị khảo sát hàm số bài toán tối ưu": "derivatives-and-function-analysis",
        "kỹ thuật tích phân giải tích thpt nguyên hàm định lý cơ bản giải tích tích phân xác định phương pháp đổi biến tích phân từng phần ứng dụng tính diện tích và thể tích": "integration-techniques",
        "phương trình vi phân nhập môn phương trình cấp một tách biến và tuyến tính điều kiện đầu mô hình tăng trưởng và phân rã theo hàm mũ phương trình vi phân cấp hai": "differential-equations-intro",
        "ma trận và hệ phương trình tuyến tính phép toán ma trận định thức ma trận nghịch đảo giải hệ phương trình khử gauss quy tắc cramer": "matrices-and-linear-systems",
        "xác suất tổ hợp thi chuyên không gian mẫu và biến cố xác suất bằng cách đếm xác suất có điều kiện và độc lập quy tắc cộng và nhân đếm phần bù kỳ vọng": "combinatorial-probability",
        "toán rời rạc và lý thuyết tập hợp tập hợp và các phép toán quan hệ và ánh xạ nguyên lý chuồng bồ câu tổ hợp cơ bản số học đồng dư nền tảng tin học": "discrete-math-and-set-theory",
        "địa chất học và khoa học trái đất cấu tạo trái đất và kiến tạo mảng đá và chu trình đá khoáng vật thời gian địa chất và hóa thạch động đất và núi lửa": "geology-and-earth-science",
        "khoa học môi trường và phát triển bền vững hệ sinh thái và đa dạng sinh học ô nhiễm biến đổi khí hậu cạn kiệt tài nguyên năng lượng tái tạo": "environmental-science-and-sustainability",
        "phương pháp khoa học và thí nghiệm giả thuyết biến số và nhóm đối chứng thiết kế thí nghiệm thu thập và phân tích dữ liệu bản chất tri thức khoa học lý thuyết và tái lặp": "scientific-method-and-experimentation",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_deep_bio_humanities_topics_match():
    from app import skills
    cases = {
        "di truyền phân tử và công nghệ sinh học chuyên sinh nhân đôi adn phiên mã và dịch mã giáo lý trung tâm điều hòa gen đột biến gen công nghệ pcr crispr": "molecular-genetics-and-biotech",
        "hô hấp tế bào và chuyển hóa chuyên sinh đường phân chu trình krebs chuỗi vận chuyển electron atp năng lượng của tế bào hô hấp hiếu khí và kị khí lên men enzyme": "cellular-respiration-and-metabolism",
        "miễn dịch học và bệnh tật chuyên sinh tác nhân gây bệnh miễn dịch tự nhiên và thu được kháng thể tế bào b và t trí nhớ miễn dịch và tiêm vắc xin": "immunology-and-disease",
        "di truyền quần thể và tiến hóa chuyên sinh tần số alen và kiểu gen nguyên lý hardy weinberg các nhân tố tiến hóa chọn lọc phiêu bạt dòng gen đột biến": "population-genetics-and-evolution",
        "hệ thần kinh và hệ nội tiết chuyên sinh tế bào thần kinh và điện thế hoạt động xináp và chất dẫn truyền thần kinh phản xạ hoocmon và các tuyến nội tiết": "nervous-and-endocrine-systems",
        "phân tích và phê bình văn học chuyên văn phân tích chủ đề nhân vật cốt truyện biện pháp tu từ ẩn dụ biểu tượng giọng điệu diễn giải ý nghĩa tác phẩm": "literary-analysis-and-criticism",
        "tổng quan lịch sử việt nam văn lang âu lạc và thời bắc thuộc các triều đại phong kiến độc lập các cuộc kháng chiến chống ngoại xâm thời pháp thuộc": "vietnamese-history-overview",
        "triết học và tư duy phản biện các nhánh triết học siêu hình nhận thức đạo đức logic các khung đạo đức thuyết vị lợi nghĩa vụ đức hạnh nhận diện ngụy biện và đánh giá bằng chứng": "philosophy-and-critical-thinking",
        "logic hình thức và lập luận logic mệnh đề và logic vị từ bảng chân trị tính hợp lệ và tính đúng đắn suy luận diễn dịch và quy nạp lượng từ": "formal-logic-and-argumentation",
        "thiên văn học và vũ trụ học hệ mặt trời các ngôi sao và vòng đời sao thiên hà vụ nổ lớn big bang và sự giãn nở vũ trụ vật chất tối": "astronomy-and-cosmology",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_deep_math_physics_chem_topics_match():
    from app import skills
    cases = {
        "phương trình hàm thi chuyên olympiad tìm tất cả hàm số thỏa mãn kỹ thuật thế giá trị đặc biệt tính đơn ánh toàn ánh điểm bất động phương trình cauchy chứng minh chặt không còn nghiệm khác": "functional-equations",
        "số phức thi chuyên toán nâng cao đơn vị ảo i và phép toán mặt phẳng phức mô đun và acgumen dạng lượng giác và mũ định lý moivre căn bậc n của đơn vị": "complex-numbers",
        "hình học phẳng olympiad thi chuyên các tâm tam giác tứ giác nội tiếp và góc nội tiếp phương tích của điểm định lý ptolemy ceva menelaus chứng minh đồng quy thẳng hàng": "plane-geometry-olympiad",
        "quy nạp toán học và phương pháp chứng minh thi chuyên quy nạp yếu và mạnh chứng minh phản chứng nguyên lý cực hạn bất biến và đơn biến chuồng bồ câu": "mathematical-induction-and-proof",
        "mạch điện xoay chiều chuyên lý dòng điện xoay chiều và giá trị hiệu dụng rms dung kháng cảm kháng và tổng trở mạch rlc nối tiếp hiện tượng cộng hưởng": "ac-circuits-and-impedance",
        "thuyết tương đối hẹp chuyên lý hai tiên đề einstein sự giãn nở thời gian sự co độ dài phép biến đổi lorentz động lượng và năng lượng tương đối tính": "special-relativity-deep",
        "tĩnh điện và điện trường chuyên lý định luật coulomb lực điện điện trường và đường sức định luật gauss điện thế tụ điện và điện dung": "electrostatics-and-electric-fields",
        "hấp dẫn và định luật kepler chuyên lý định luật vạn vật hấp dẫn newton chuyển động quỹ đạo và vệ tinh vận tốc thoát ba định luật kepler": "gravitation-and-kepler",
        "nhiệt hóa học chuyên hóa enthalpy và nhiệt phản ứng định luật hess nhiệt tạo thành chuẩn năng lượng liên kết entropy và năng lượng tự do gibbs": "thermochemistry-and-hess-law",
        "cơ chế phản ứng hữu cơ chuyên hóa tác nhân ái nhân và ái điện tử phản ứng thế sn1 sn2 phản ứng tách e1 e2 phản ứng cộng vào nối đôi markovnikov": "organic-reaction-mechanisms",
        "cấu hình electron và số lượng tử chuyên hóa các obitan s p d f bốn số lượng tử nguyên lý vững bền aufbau quy tắc hund nguyên lý loại trừ pauli": "electron-configuration-and-quantum-numbers",
        "dung dịch và tính chất tập hợp chuyên hóa độ tan tích số tan ksp tăng nhiệt độ sôi giảm nhiệt độ đông đặc áp suất thẩm thấu": "solutions-and-colligative-properties",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_advanced_chuyen_topics_match():
    from app import skills
    cases = {
        "số học và chia hết thi chuyên olympiad tính chia hết và ước bội số nguyên tố và phân tích thừa số ước chung lớn nhất euclid số học đồng dư modulo định lý fermat nhỏ phương trình nghiệm nguyên diophantine": "number-theory-and-divisibility",
        "bất đẳng thức và cực trị thi chuyên olympiad bất đẳng thức am gm cosi cauchy schwarz bunhiacopxki trung bình lũy thừa và sắp xếp lại jensen và tính lồi kỹ thuật đổi biến chuẩn hóa tổng bình phương sos": "inequalities-and-optimization",
        "tổ hợp và đếm thi chuyên olympiad quy tắc cộng và nhân đếm chỉnh hợp hoán vị tổ hợp nhị thức newton nguyên lý bù trừ song ánh đếm nguyên lý dirichlet chuồng bồ câu truy hồi và hàm sinh": "combinatorics-and-counting",
        "động lực học vật rắn và chuyển động quay chuyên lý mô men lực mô men quán tính mô men động lượng và bảo toàn động năng quay chuyển động lăn không trượt định lý trục song song": "rigid-body-and-rotational-dynamics",
        "thuyết động học phân tử và nhiệt động lực học chuyên lý phương trình khí lý tưởng nhiệt độ và động năng phân tử phân bố vận tốc maxwell boltzmann nguyên lý thứ nhất và nội năng động cơ nhiệt và hiệu suất": "kinetic-theory-and-gas-laws",
        "cân bằng hóa học và động hóa học chuyên hóa phản ứng thuận nghịch hằng số cân bằng k nguyên lý chuyển dịch le chatelier tốc độ phản ứng và định luật tốc độ năng lượng hoạt hóa và arrhenius chất xúc tác": "chemical-equilibrium-and-kinetics",
        "điện hóa học và phản ứng oxi hóa khử chuyên hóa số oxi hóa và cân bằng phản ứng oxi hóa khử pin điện hóa galvani thế điện cực chuẩn và suất điện động dãy điện hóa điện phân và định luật faraday": "electrochemistry-and-redox",
        "cân bằng axit bazơ chuyên hóa thuyết bronsted lowry cho nhận proton axit bazơ mạnh và yếu hằng số ka kb và pka thang ph poh và tính toán dung dịch đệm henderson hasselbalch thủy phân muối đường cong chuẩn độ": "acid-base-and-ph-equilibria",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_thpt_subjects_topics_match():
    from app import skills
    cases = {
        "lượng giác và hàm số thpt sin cos tan tỉ số lượng giác đường tròn lượng giác và radian công thức lượng giác biến đổi đồ thị và tính tuần hoàn giải phương trình lượng giác hàm số tập xác định": "trigonometry-and-functions",
        "mũ và logarit thpt hàm số mũ tăng giảm theo cấp số nhân số e cơ số tự nhiên định nghĩa và tính chất logarit giải phương trình mũ và logarit đồ thị hàm mũ ứng dụng lãi kép chu kỳ bán rã": "exponentials-and-logarithms",
        "dao động và sóng thpt dao động điều hòa biên độ chu kỳ tần số con lắc và lò xo tính chất sóng bước sóng tốc độ sóng ngang và sóng dọc giao thoa sóng dừng và cộng hưởng sóng âm": "oscillations-and-waves",
        "quang học và ánh sáng thpt phản xạ và khúc xạ ánh sáng định luật snell và phản xạ toàn phần thấu kính và gương tạo ảnh bản chất sóng của ánh sáng giao thoa nhiễu xạ quang phổ điện từ tán sắc": "optics-and-light",
        "hóa hữu cơ cơ bản thpt cacbon liên kết và mạch cacbon hydrocacbon ankan anken ankin thơm nhóm chức ancol axit este anđehit amin đồng phân danh pháp gọi tên các loại phản ứng hữu cơ": "organic-chemistry-basics",
        "sinh thái học và hệ sinh thái thpt quần thể và quần xã sinh vật chuỗi và lưới thức ăn bậc dinh dưỡng dòng năng lượng và tháp sinh khối chu trình sinh địa hóa quan hệ giữa các loài diễn thế đa dạng sinh học": "ecology-and-ecosystems",
        "tổng quan lịch sử thế giới các nền văn minh cổ đại và đế chế thời trung cổ thời đại khám phá các cuộc cách mạng khoa học công nghiệp chính trị hai cuộc thế chiến chiến tranh lạnh thời hiện đại": "world-history-overview",
        "địa lý tự nhiên và khí hậu thpt cấu tạo trái đất và kiến tạo mảng địa hình dạng đất khí quyển và thời tiết đới khí hậu và yếu tố hình thành vòng tuần hoàn nước tài nguyên thiên nhiên biến đổi khí hậu": "physical-geography-and-climate",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_knowledge_science_topics_match():
    from app import skills
    cases = {
        "cơ học newton và lực ba định luật chuyển động lực hấp dẫn ma sát pháp tuyến biểu đồ vật thể tự do động lượng và xung lượng công năng lượng và bảo toàn chuyển động tròn": "newtonian-mechanics-and-forces",
        "điện từ học cơ bản electromagnetism điện tích và điện trường điện áp dòng điện điện trở từ tính và nam châm điện cảm ứng điện từ hợp nhất maxwell sóng điện từ và ánh sáng": "electromagnetism-fundamentals",
        "cấu tạo nguyên tử và bảng tuần hoàn proton neutron electron số hiệu nguyên tử và đồng vị lớp electron và cấu hình tổ chức bảng tuần hoàn xu hướng tuần hoàn độ âm điện": "atomic-structure-and-periodic-table",
        "liên kết hóa học và phản ứng liên kết ion cộng hóa trị kim loại vì sao nguyên tử liên kết quy tắc bát tử hình dạng phân tử và phân cực phương trình hóa học cân bằng axit bazơ": "chemical-bonding-and-reactions",
        "sinh học tế bào và di truyền cấu trúc tế bào và bào quan tế bào nhân sơ và nhân thực phân bào nguyên phân giảm phân adn và nhiễm sắc thể di truyền mendel gen alen trội lặn": "cell-biology-and-heredity",
        "tiến hóa và chọn lọc tự nhiên biến dị di truyền sống sót khác biệt hậu duệ có biến đổi bằng chứng hóa thạch adn tương đồng hình thành loài hiểu lầm thường gặp đa dạng sự sống": "evolution-and-natural-selection",
        "trực giác giải tích calculus giới hạn đạo hàm là tốc độ thay đổi độ dốc tích phân là tích lũy diện tích định lý cơ bản liên hệ đạo hàm tích phân tốc độ thay đổi và tích lũy": "calculus-intuition",
        "trực giác đại số tuyến tính linear algebra vector và không gian vector ma trận là phép biến đổi tuyến tính nhân ma trận là hợp phép biến đổi định thức vector riêng trị riêng tích vô hướng": "linear-algebra-intuition",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_econ_psych_eng_topics_match():
    from app import skills
    cases = {
        "kinh tế học vi mô và hành vi thị trường microeconomics cung và cầu điểm cân bằng độ co giãn giá thặng dư tiêu dùng tư duy cận biên cấu trúc thị trường độc quyền ngoại ứng": "microeconomics-and-market-behavior",
        "kinh tế học vĩ mô macroeconomics tổng sản phẩm và tăng trưởng gdp lạm phát thất nghiệp chu kỳ kinh tế chính sách tiền tệ ngân hàng trung ương lãi suất chính sách tài khóa": "macroeconomics-fundamentals",
        "kinh tế học hành vi behavioral economics con người lệch khỏi mô hình duy lý lý thuyết triển vọng và ngại thua kế toán tinh thần thiên lệch hiện tại mặc định kiến trúc lựa chọn cú hích nudge": "behavioral-economics-and-nudges",
        "thiên lệch nhận thức và ra quyết định cognitive bias tư duy hệ 1 và hệ 2 thiên lệch xác nhận mỏ neo tính sẵn có nhận thức muộn tự tin thái quá chiến lược khử thiên lệch": "cognitive-biases-and-decision-making",
        "tâm lý học trí nhớ và học tập memory learning mã hóa lưu trữ truy xuất trí nhớ làm việc và dài hạn đường cong quên lặp lại ngắt quãng luyện truy xuất xen kẽ": "memory-and-learning-psychology",
        "tâm lý học xã hội social psychology tuân theo và phục tùng thuyết phục và ảnh hưởng nguyên tắc cialdini quy kết bất hòa nhận thức bằng chứng xã hội nội nhóm ngoại nhóm sức mạnh hoàn cảnh": "social-psychology-fundamentals",
        "mạch điện tử tương tự analog circuit định luật ohm và kirchhoff chia áp mạch rc rl và hằng số thời gian trở kháng và bộ lọc diode và transistor khuếch đại thuật toán op-amp": "analog-circuit-design-fundamentals",
        "tĩnh học và sức bền vật liệu statics cân bằng lực và biểu đồ vật thể tự do mô men và mô men xoắn ứng suất và biến dạng kéo nén cắt uốn và dầm hệ số an toàn": "statics-and-mechanics-of-materials",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_language_topics_match():
    from app import skills
    cases = {
        "viết tốt tiếng Việt và giao tiếp tự nhiên hệ đại từ xưng hô theo quan hệ giọng điệu và mức lịch sự dấu thanh và chính tả trật tự từ tự nhiên tránh dịch cứng": "writing-well-in-vietnamese",
        "viết tốt tiếng Anh và giao tiếp rõ ràng thể chủ động tiếng Anh đơn giản lỗi mạo từ giới từ thì thành ngữ anh-anh anh-mỹ súc tích": "writing-well-in-english",
        "viết tốt tiếng Trung mandarin chữ Hán giản thể phồn thể không biến hình cấu trúc chủ đề lượng từ thanh điệu thành ngữ tứ tự chengyu": "writing-well-in-chinese",
        "viết tốt tiếng Nhật ba bảng chữ hiragana katakana kanji kính ngữ keigo trợ từ trật tự sov chủ đề wa ga các mức trang trọng ngữ cảnh hàm ẩn": "writing-well-in-japanese",
        "viết tốt tiếng Hàn bảng chữ hangul các cấp độ nói kính ngữ jondaenmal banmal trợ từ trật tự sov dấu chủ đề eun neun i ga thứ bậc xã hội": "writing-well-in-korean",
        "viết tốt tiếng Pháp giống ngữ pháp và hợp giống tu so với vous chia động từ và thức giả định dấu và chính tả bạn giả false friend": "writing-well-in-french",
        "viết tốt tiếng Tây Ban Nha giống ngữ pháp tú usted vos chia động từ giả định ser so với estar dấu và chấm câu ngược biến thể vùng miền": "writing-well-in-spanish",
        "viết tốt tiếng Đức bốn cách và ba giống du so với Sie trật tự động từ cuối câu v2 danh từ ghép viết hoa danh từ động từ tách": "writing-well-in-german",
        "viết tốt tiếng Ả Rập chữ viết phải sang trái nối chữ hệ gốc và khuôn song thể diglossia msa phương ngữ giống số đôi dual chia động từ": "writing-well-in-arabic",
        "viết tốt tiếng Nga bảng chữ kirin sáu cách ba giống thể động từ hoàn thành chưa hoàn thành ty vy không có mạo từ trật tự từ linh hoạt": "writing-well-in-russian",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_creative_writing_topics_match():
    from app import skills
    cases = {
        "cấu trúc truyện và cốt truyện story structure plot ba hồi three act hành trình người hùng nhân quả biến cố khởi đầu cao trào cài cắm và hồi đáp setup payoff": "story-structure-and-plot",
        "phát triển nhân vật và cung nhân vật character development arc động cơ và khát khao muốn so với cần khuyết điểm và vết thương cung thay đổi trưởng thành tính chủ động agency giọng riêng": "character-development-and-arc",
        "văn phong và giọng văn prose style voice nhịp câu và biến đổi độ dài chọn từ chính xác cắt từ thừa ngôn ngữ cụ thể giọng điệu văn vực xây giọng văn riêng biên tập": "prose-style-and-voice",
        "xây dựng thế giới cho truyện worldbuilding fiction bối cảnh đáng tin và nhất quán nội tại nguyên tắc tảng băng hòa thế giới vào câu chuyện hệ phép thuật luật văn hóa lịch sử tránh nhồi thông tin": "worldbuilding-for-fiction",
        "viết hội thoại truyện cuốn hút compelling fiction dialogue giọng riêng từng nhân vật tầng nghĩa ẩn và điều không nói hội thoại làm nhiều việc tránh nói thẳng lộ liễu thẻ dẫn và cử chỉ": "writing-compelling-dialogue",
        "cho thấy đừng kể show dont tell và dựng cảnh scene craft kịch tính hóa so với tóm tắt chi tiết cụ thể giác quan hé lộ cảm xúc qua hành động cảnh so với tóm tắt neo cảnh vào không gian": "show-dont-tell-and-scene-craft",
        "nhịp truyện và căng thẳng narrative pacing tension điều tốc độ qua cảnh tóm tắt xây và giải căng thẳng cược và đồng hồ đếm ngược mồi câu và trang cuốn cliffhanger ngắt chương": "narrative-pacing-and-tension",
        "ngôi kể và người kể chuyện point of view narration ngôi thứ nhất thứ ba hạn tri so với toàn tri khoảng cách trần thuật người kể không đáng tin lỗi nhảy đầu head hopping chọn thì": "point-of-view-and-narration",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_game_genre_batch5_topics_match():
    from app import skills
    cases = {
        "thu thập và thuần hóa quái vật monster taming creature collector bắt và sưu tầm hệ khắc chế nguyên tố type matchup chỉ số tiến hóa xây đội chiến đấu quái theo lượt ám ảnh sưu tầm": "creature-collector-and-monster-taming",
        "phiêu lưu văn bản và tiểu thuyết tương tác text adventure interactive fiction parser so với lựa chọn mô hình thế giới ngôn ngữ lệnh câu đố bằng chữ văn xuôi là giao diện phân nhánh": "text-adventure-and-interactive-fiction",
        "god game và mô phỏng simulation điều khiển gián tiếp các tác nhân tự chủ hành vi nổi lên từ luật đơn giản vòng quan sát và tác động nhẹ chiều sâu hệ thống mục tiêu người chơi": "god-game-and-simulation",
        "tycoon và mô phỏng kinh doanh business sim xây và vận hành doanh nghiệp có lãi lõi kinh tế thu chi lợi nhuận giá và sự hài lòng khách hàng vòng tăng trưởng chiều sâu tối ưu": "tycoon-and-business-sim",
        "thiết kế game gõ phím typing game biến gõ bàn phím thành lối chơi nhắm từ và câu độ khó theo tốc độ chính xác so với tốc độ tiến trình kỹ năng dạy gõ mười ngón": "typing-game-design",
        "thiết kế game đố vui và trắc nghiệm trivia quiz viết câu hỏi và độ khó độ rộng danh mục và công bằng định dạng câu hỏi tính điểm và cơ chế rủi ro nhịp độ áp lực thời gian": "trivia-and-quiz-game-design",
        "thiết kế roguelike deckbuilder slay the spire kết hợp lượt roguelike với xây bài trong lượt thưởng thẻ và làm mỏng bộ bài vật phẩm cổ vật relic thiết kế trận bản đồ độ khó leo thang khám phá build": "roguelike-deckbuilder-design",
        "thiết kế escape room và câu đố môi trường environmental puzzle chuỗi câu đố liên kết tìm kiếm quan sát khóa và chìa thiết kế manh mối độ khó công bằng khoảnh khắc bừng sáng hệ gợi ý": "escape-room-and-environmental-puzzle",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_game_genre_batch4_topics_match():
    from app import skills
    cases = {
        "dungeon crawler và blobber ô lưới di chuyển theo ô lập bản đồ điều hướng cả tổ đội là một khối party blob thiết kế trận và bẫy chiến lợi phẩm căng thẳng khám phá": "dungeon-crawler-and-blobber-design",
        "thiết kế jrpg và quy ước rpg nhật câu chuyện tuyến tính và nhân vật hệ chiến đấu theo lượt atb hệ tổ đội và nghề chạm địch ngẫu nhiên cày cấp đề cao tự sự": "jrpg-design-and-conventions",
        "thiết kế crpg rpg phương tây choice driven nhân vật do người chơi tạo và nhập vai lựa chọn phân nhánh có ý nghĩa và hậu quả phản ứng thế giới hội thoại kiểm tra kỹ năng nhiều lời giải": "crpg-and-choice-driven-design",
        "thiết kế hero shooter bắn súng theo tướng dàn tướng độc đáo kỹ năng và vai trò riêng đội hình và cộng hưởng kinh tế chiêu cuối ultimate khắc chế counter pick chế độ theo mục tiêu": "hero-shooter-design",
        "thiết kế extraction shooter bắn súng rút lui vòng lặp vào raid và thoát ra với rủi ro mất đồ sợ mất đồ gear fear căng thẳng pvpve điểm thoát rủi ro tiến hay rút": "extraction-shooter-design",
        "thiết kế tactical shooter bắn súng chiến thuật thời gian hạ gục thấp sát thương cao chân thực nhịp chậm gunplay giật tản đạn kinh tế theo vòng giao tiếp phối hợp kiểm soát bản đồ tiện ích": "tactical-shooter-design",
        "thiết kế puzzle platformer đố nhảy kết hợp cơ chế di chuyển với câu đố một cơ chế lõi đào sâu dạy qua thiết kế màn giới thiệu phát triển xoắn kết hợp cân bằng thao tác và tư duy": "puzzle-platformer-design",
        "thiết kế dating sim và cơ chế lãng mạn romance thước quan hệ thiện cảm nhân vật hẹn hò với tuyến riêng xây chỉ số quản lý thời gian hệ quà tặng hội thoại tuyến tình cảm mở khóa": "dating-sim-and-romance-design",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_game_genre_batch3_topics_match():
    from app import skills
    cases = {
        "chiến thuật thời gian thực real time tactics rtt không xây căn cứ không kinh tế vị trí và vật che kỹ năng đơn vị điều khiển micro đánh sườn tầm nhìn tổ đội nhỏ tạm dừng lên kế hoạch": "real-time-tactics-design",
        "trò chơi thẻ bài sưu tầm tcg ccg constructed thi đấu tam giác tài nguyên nhịp lợi thế thẻ nguyên mẫu bộ bài và meta thiết kế mana đường cong tương tác chồng stack xoay vòng set": "tcg-and-competitive-card-design",
        "match 3 và cơ chế câu đố lưới ghép và chuỗi đổ combo cascade tạo mảnh đặc biệt mục tiêu ràng buộc màn kinh tế lượt đi phản hồi đã mắt juicy": "match-3-and-puzzle-mechanic",
        "thiết kế game chạy vô tận endless runner tự động chạy điều khiển một chạm sinh chướng ngại thủ tục tăng tốc độ khó đuổi điểm suýt trúng near miss chơi lại tức thì": "endless-runner-design",
        "thiết kế vampire survivors bullet heaven survivors-like tự động tấn công chống bầy quái khổng lồ lên cấp vũ khí bị động tiến hóa cộng hưởng build bầy quái leo thang roguelite": "vampire-survivors-and-bullet-heaven",
        "thiết kế nhiều người bất đối xứng asymmetric multiplayer các phe luật mục tiêu sức mạnh khác nhau one vs many cân bằng vai trò không cân bằng thông tin bất đối xứng phản đòn": "asymmetric-multiplayer-design",
        "thiết kế mô phỏng bay và phương tiện flight vehicle sim phổ độ chân thực sim tới arcade vật lý bay trục điều khiển đồng hồ hệ thống phản hồi phân tầng hỗ trợ assists nhiệm vụ sandbox": "flight-and-vehicle-sim-design",
        "thiết kế board game và trò chơi bàn tabletop cơ chế lõi đặt thợ kiểm soát vùng engine building tương tác người chơi may rủi kỹ năng thời gian chờ co giãn số người rõ ràng luật dễ dạy": "board-game-and-tabletop-design",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_game_genre_batch2_topics_match():
    from app import skills
    cases = {
        "thiết kế beat em up và game đánh đấm brawler khống chế đám đông nhiều kẻ địch crowd control combo tung hứng juggle di chuyển mặt phẳng 2.5d đợt encounter": "beat-em-up-and-brawler-design",
        "thiết kế hack and slash character action stylish hệ combo sâu hủy đòn cancel thước phong cách xếp hạng style rank bộ chiêu né tạo lợi thế dodge offset just frame": "hack-and-slash-character-action",
        "thiết kế game phiêu lưu trỏ và nhấp point and click adventure câu đố túi đồ kết hợp vật phẩm inventory logic tránh moon logic cây hội thoại bẫy dò pixel": "point-and-click-adventure-design",
        "thiết kế walking sim và khám phá tự sự narrative exploration ưu tiên câu chuyện ít cơ chế kể chuyện qua môi trường dẫn dắt không chỉ dấu không khí nơi chốn": "walking-sim-and-narrative-exploration",
        "thiết kế party game trò chơi tiệc tùng luật dễ tiếp cận chơi ngay nhiều người cùng máy động lực xã hội cơ chế bám đuổi hỗn loạn catch up đa dạng minigame": "party-game-design",
        "thiết kế game thể thao sports mô phỏng so với arcade tính chân thực tái hiện cảm giác luật điều khiển ánh xạ hành động ai đồng đội chế độ sự nghiệp franchise": "sports-game-design",
        "thiết kế game sandbox và xây dựng sáng tạo building công cụ construction hệ khối voxel mô-đun mục tiêu do người chơi tạo player driven vật lý hệ thống nổi lên chia sẻ cộng đồng": "sandbox-building-game-design",
        "thiết kế twin stick và arena shooter bắn góc nhìn trên điều khiển tách rời di chuyển và ngắm bầy kẻ địch nhịp sinh spawn cảm giác vũ khí kéo địch kiting đợt leo thang": "twin-stick-and-arena-shooter-design",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_more_game_genre_topics_match():
    from app import skills
    cases = {
        "thiết kế game bắn súng fps shooter cảm giác bắn gunplay thời gian hạ gục ttk vai trò cân bằng vũ khí di chuyển ngắm bắn bản đồ tầm nhìn sightline giật đạn": "fps-and-shooter-design",
        "thiết kế bullet hell và game bắn máy bay shmup danmaku mẫu đạn hộp va chạm nhỏ dễ đọc rừng đạn lướt đạn graze tính điểm mẫu trùm boss pattern": "bullet-hell-and-shmup-design",
        "thiết kế combat souls-like và action rpg chiến đấu chậm dựa thể lực stamina cam kết hồi phục đòn địch báo trước telegraph né đỡ phản đòn dodge parry lửa trại": "souls-like-and-action-rpg-combat",
        "thiết kế game kinh dị sinh tồn survival horror khan hiếm tài nguyên yếu ớt căng thẳng dread hơn jump scare không khí ánh sáng âm thanh phòng an toàn safe room": "survival-horror-design",
        "thiết kế game chiến thuật theo lượt turn based tactics di chuyển ô lưới vị trí vật che đánh sườn flanking tỉ lệ trúng rng lớp đơn vị kỹ năng chết vĩnh viễn": "turn-based-tactics-design",
        "thiết kế auto battler và auto chess vòng lặp mua sắm kinh tế đặt quân chiến đấu tự động cộng hưởng tộc hệ synergy vàng lãi suất interest reroll đặt vị trí": "auto-battler-and-autochess-design",
        "thiết kế game nông trại và mô phỏng cuộc sống farming life sim cozy chu kỳ ngày mùa kinh tế năng lượng thời gian trồng trọt quan hệ xã hội nhịp thư giãn": "farming-and-life-sim-design",
        "thiết kế game gacha và dịch vụ trực tuyến live service quay ngẫu nhiên sưu tầm nhân vật tiền tệ premium hệ thống pity banner fomo daily giữ chân cá voi whale đạo đức": "gacha-and-live-service-design",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_medical_law_science_topics_match():
    from app import skills
    cases = {
        "lập luận lâm sàng và chẩn đoán clinical reasoning khai thác bệnh sử khám chẩn đoán phân biệt differential xác suất trước sau test độ nhạy đặc hiệu red flag thiên kiến": "clinical-reasoning-and-diagnosis",
        "dược lý và tác dụng thuốc pharmacology dược động học hấp thu phân bố chuyển hóa thải trừ adme dược lực học thụ thể chủ vận đối vận liều đáp ứng bán thải": "pharmacology-and-drug-action",
        "luật hợp đồng cơ bản contract law hình thành đề nghị chấp nhận đối giá consideration vi phạm và biện pháp khắc phục breach remedies điều khoản bồi hoàn giới hạn trách nhiệm": "contract-law-fundamentals",
        "luật sở hữu trí tuệ intellectual property bản quyền copyright bằng sáng chế patent nhãn hiệu trademark bí mật thương mại cấp phép mã nguồn mở": "intellectual-property-law",
        "phản ứng hóa học và tỉ lượng stoichiometry cân bằng phương trình mol khối lượng mol chất giới hạn hiệu suất oxi hóa khử redox bảo toàn khối lượng": "chemical-reactions-and-stoichiometry",
        "cơ học quỹ đạo và bay vũ trụ orbital mechanics định luật kepler vận tốc quỹ đạo phương trình tên lửa delta-v chuyển quỹ đạo hohmann vận tốc thoát": "orbital-mechanics-and-spaceflight",
        "thiết kế visual novel và truyện phân nhánh branching narrative tuyến truyện lựa chọn có ý nghĩa hậu quả biến cờ trạng thái nhiều kết thúc ảo giác lựa chọn": "visual-novel-and-branching-narrative",
        "thiết kế mmo và nền kinh tế ảo virtual economy thế giới bền vững persistent sharding endgame guild nguồn tạo điểm hút faucet sink lạm phát nhà đấu giá": "mmo-design-and-virtual-economies",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_science_bio_finance_topics_match():
    from app import skills
    cases = {
        "thiết kế battle royale sinh tồn cuối cùng vòng bo thu hẹp shrinking zone storm người cuối trụ lại phân bố loot nhặt đồ mật độ người chơi bản đồ lớn": "battle-royale-design",
        "thiết kế 4x và đại chiến lược grand strategy khám phá mở rộng khai thác tiêu diệt theo lượt quản lý đế chế cây công nghệ thể chế kinh tế sản lượng yields ngoại giao": "fourx-grand-strategy-design",
        "sinh học phân tử và biểu hiện gen cấu trúc adn dna chuỗi xoắn kép giáo lý trung tâm dna rna protein phiên mã dịch mã mã di truyền codon điều hòa": "molecular-biology-and-gene-expression",
        "tin sinh học và phân tích chuỗi bioinformatics gióng hàng alignment blast định dạng fasta fastq sam vcf quy trình ngs variant calling lắp ráp bộ gen": "bioinformatics-and-sequence-analysis",
        "trực giác cơ học lượng tử quantum mechanics lưỡng tính sóng hạt chồng chập superposition hàm sóng xác suất nguyên lý bất định rối lượng tử entanglement": "quantum-mechanics-intuition",
        "nhiệt động lực học và entropy bốn định luật bảo toàn năng lượng entropy và mũi tên thời gian nhiệt và công nhiệt độ thuận nghịch năng lượng tự do": "thermodynamics-and-entropy",
        "thu nhập cố định và định giá trái phiếu bond coupon mệnh giá kỳ hạn giá trị hiện tại lợi suất đáo hạn ytm duration convexity đường cong lợi suất": "fixed-income-and-bond-pricing",
        "rủi ro tín dụng và mô hình vỡ nợ credit risk xác suất vỡ nợ pd tổn thất lgd dư nợ ead tổn thất kỳ vọng chấm điểm tín dụng chênh lệch spread": "credit-risk-and-default-modeling",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_game_genre_and_education_topics_match():
    from app import skills
    cases = {
        "thiết kế moba đấu trường đội lane creep lính tháp mục tiêu kinh tế last hit vàng bộ kỹ năng tướng vai trò giao tranh tổng team fight snowball": "moba-design",
        "hệ thống rpg và tiến trình nhân vật chỉ số thuộc tính stats kinh nghiệm lên cấp cây kỹ năng build loot trang bị công thức sát thương scaling class": "rpg-systems-and-progression",
        "thiết kế thế giới mở và sandbox tự do phi tuyến điểm quan tâm mật độ points of interest kể chuyện qua môi trường lối chơi nổi lên emergent tránh filler": "open-world-and-sandbox-design",
        "immersive sim và thiết kế hệ thống các hệ thống tương tác thay vì kịch bản scripted nhiều lời giải luật thế giới nhất quán lối chơi nổi lên emergent agency": "immersive-sim-and-systemic-design",
        "thiết kế game nhàn rỗi và tăng dần idle incremental clicker tăng trưởng hàm mũ số lớn vòng mua máy phát generators prestige reset offline": "idle-and-incremental-game-design",
        "thiết kế game xây thành phố và quản lý management sim chuỗi cung ứng vòng sản xuất nhu cầu dịch vụ quy hoạch phân vùng zoning cân bằng kinh tế upkeep": "city-builder-and-management-design",
        "thiết kế giảng dạy instructional design thiết kế ngược từ kết quả backward design mục tiêu học tập thang bloom quy trình addie căn chỉnh alignment": "instructional-design-and-backward-design",
        "thiết kế game giáo dục và gamification động lực nội tại ngoại lai lồng học vào cơ chế lõi điểm huy hiệu bảng xếp hạng tránh pointsification": "educational-game-and-gamification-design",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_game_genre_topics_match():
    from app import skills
    cases = {
        "thiết kế rts chiến thuật thời gian thực kinh tế thu thập tài nguyên vĩ mô vi mô macro micro build order tech tree khắc chế sương mù chiến tranh": "rts-and-strategy-design",
        "thiết kế metroidvania khám phá bản đồ liên thông mở khóa bằng năng lực ability gate khóa và chìa quay lại vùng cũ backtracking": "metroidvania-design",
        "thiết kế game nhịp điệu rhythm cửa sổ thời gian bậc chấm perfect miss soạn beatmap theo nhạc độ trễ âm thanh hiệu chỉnh calibration": "rhythm-game-design",
        "thiết kế game đối kháng fighting thế giằng co neutral footsies lớp đoán ý mixup khoảng cách spacing hệ combo hủy đòn cancel thanh nộ meter": "fighting-game-design",
        "thiết kế game đua xe racing mô hình điều khiển arcade so với sim thiết kế đường đua racing line drift bám đường grip rubber banding": "racing-game-design",
        "thiết kế game sinh tồn và chế tạo survival crafting vòng lặp tài nguyên khan hiếm nhu cầu đói khát cây chế tạo xây căn cứ base building": "survival-and-crafting-design",
        "thiết kế game lén lút stealth nhận thức kẻ địch nón nhìn vision cone nghe âm thanh trạng thái cảnh giác alert vật che lối đi đánh lạc hướng": "stealth-game-design",
        "thiết kế game giải đố puzzle cơ chế lõi không gian luật dạy không cần tutorial khoảnh khắc bừng sáng aha giới thiệu phát triển xoắn kết hợp": "puzzle-game-design",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_animation_and_game_genre_topics_match():
    from app import skills
    cases = {
        "three.js và web 3d đồ thị cảnh scene mesh material tải mô hình gltf phát animation clip animationmixer instancing": "three-js-and-web-3d",
        "hoạt hình xương và bọc da skeletal skinning bind pose linear blend skinning trộn crossfade blend tree animation state machine root motion": "skeletal-animation-and-skinning",
        "css animation và transition animate transform và opacity mượt 60fps compositor tránh reflow will-change prefers-reduced-motion": "css-animations-and-transitions",
        "chuyển cảnh phần tử dùng chung shared element hero transition kỹ thuật flip view transitions api hero animation flutter": "shared-element-and-hero-transitions",
        "hiệu năng animation web ngân sách khung hình 60fps luồng chính compositor tránh giật layout thrash requestanimationframe chuyển động có mục đích": "web-animation-performance",
        "thiết kế roguelike roguelite lượt chơi run permadeath sinh màn ngẫu nhiên meta progression đa dạng build cộng hưởng synergy": "roguelike-design-and-procedural-runs",
        "thiết kế deckbuilder và game thẻ bài kinh tế thẻ chi phí làm mỏng bộ bài cộng hưởng synergy archetype cỗ máy rút bài": "deckbuilder-and-card-game-design",
        "thiết kế tower defense và đợt tấn công wave đường đi mê cung pathing vai trò tháp cây nâng cấp loại kẻ địch khắc chế nhịp leo thang": "tower-defense-and-wave-design",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_simulation_and_domain_topics_match():
    from app import skills
    cases = {
        "động học ngược inverse kinematics ik giải góc khớp đặt đầu cuối end effector giải hai khớp analytic ccd fabrik": "inverse-kinematics-for-animation",
        "mười hai nguyên tắc hoạt hình disney co giãn squash and stretch lấy đà anticipation chậm vào ra slow in out cung arc": "animation-twelve-principles",
        "nhạc lý cơ bản sáng tác âm giai scale hợp âm ba triad vòng hòa thanh progression căng và giải quyết tension resolution": "music-theory-basics",
        "hệ điều khiển và bộ pid vòng phản hồi feedback loop tỉ lệ tích phân vi phân p i d chỉnh tham số tuning vọt lố overshoot": "control-systems-and-pid",
        "robot học động học và điều khiển hệ trục tọa độ frame transform bậc tự do degrees of freedom lập kế hoạch chuyển động tránh vật cản": "robotics-kinematics-and-control",
        "mô phỏng vải và vật thể mềm soft body hệ khối lượng lò xo mass-spring tích phân verlet động lực dựa vị trí position based dynamics pbd": "cloth-and-soft-body-simulation",
        "mô phỏng chất lỏng cơ bản fluid simulation lưới eulerian so với hạt sph đối lưu advection chiếu áp suất giữ không nén": "fluid-simulation-basics",
        "dự báo chuỗi thời gian time series forecasting xu hướng mùa vụ tính dừng stationarity sai phân arima backtesting rolling origin": "time-series-forecasting",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_game_systems_topics_match():
    from app import skills
    cases = {
        "âm thanh game và nhạc thích ứng phân lớp biến thể ducking nhạc tương tác theo lối chơi stem stinger": "game-audio-and-adaptive-music",
        "hiệu ứng shader và hậu xử lý post-processing render ra texture bloom quầng sáng tone mapping color grading vignette": "shader-effects-and-post-processing",
        "nhập liệu game hỗ trợ tay cầm controller ánh xạ hành động gán lại phím rebinding vùng chết dead zone": "input-and-controller-support",
        "điều khiển cảm ứng và game mobile cần điều khiển ảo virtual joystick cử chỉ vuốt tầm với ngón cái mục tiêu chạm to": "touch-controls-and-mobile-games",
        "atlas sprite đóng gói texture gộp nhiều ảnh vào một texture giảm draw call bin packing đệm chống rỉ màu bleeding": "sprite-atlas-and-texture-packing",
        "wave function collapse wfc điền lưới tile theo luật kề adjacency chọn ô entropy thấp nhất lan truyền propagate": "wave-function-collapse",
        "hệ thống chiến đấu và hitbox vùng nhận đòn hurtbox dữ liệu khung hình startup active recovery hitstun combo i-frame": "combat-systems-and-hitboxes",
        "hệ thống túi đồ và vật phẩm inventory item xếp chồng stack trang bị equipment độ hiếm rarity bảng rơi đồ loot table": "inventory-and-item-systems",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_ux_ui_deep_topics_match():
    from app import skills
    cases = {
        "khung dây wireframe và tạo mẫu prototype mức độ chi tiết fidelity lo-fi hi-fi mẫu bấm được clickable": "wireframing-and-prototyping",
        "tìm kiếm và lọc search filter gợi ý autocomplete lọc faceted trạng thái không có kết quả lưu vào url": "search-and-filter-ux",
        "trải nghiệm người dùng mới onboarding đưa tới aha moment thiết lập dần trạng thái rỗng dạy cách dùng": "onboarding-and-first-run-experience",
        "thiết kế thông báo toast snackbar mức độ nghiêm trọng tự ẩn có nút hoàn tác undo tránh quá tải": "notification-and-toast-design",
        "bảng dữ liệu data table căn cột số phải chữ trái tiêu đề dính sticky ảo hóa hàng virtualization chọn hàng bulk": "data-tables-and-grids",
        "bảng lệnh command palette cmd ctrl k tìm kiếm mờ fuzzy phím tắt shortcut quản lý tiêu điểm focus": "command-palette-and-keyboard-ux",
        "tiết lộ dần progressive disclosure ẩn phần nâng cao nút xem thêm accordion mặc định hợp lý": "progressive-disclosure",
        "thông báo lỗi và khôi phục nói lỗi gì và cách sửa không đổ lỗi cho người dùng đặt inline giữ lại dữ liệu đã nhập": "error-messages-and-recovery",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_game_dev_deep_topics_match():
    from app import skills
    cases = {
        "làm tilemap bản đồ ô lưới ghép tile tự động autotiling bitmask và thiết kế màn chơi level pacing": "tilemap-and-level-design",
        "thêm ánh sáng 2d động đổ bóng shadow từ vật cản light map trộn cộng additive normal map giả chiều sâu": "2d-lighting-and-shadows",
        "hệ thống hạt particle system làm lửa khói nổ tia lửa trail với gpu instancing hàng nghìn hạt vfx": "particle-systems-and-vfx",
        "điều khiển nhân vật platformer di chuyển nhảy với coyote time đệm nhảy jump buffer độ cao nhảy thay đổi": "character-controller-and-platformer-movement",
        "kiến trúc ecs entity component system thực thể là id component là dữ liệu ưu tiên kết hợp thay vì kế thừa": "entity-component-system-architecture",
        "hệ thống lưu game save serialize ghi trạng thái ra đĩa ghi nguyên tử tránh hỏng file save phiên bản schema": "save-systems-and-serialization",
        "thiết kế giao diện game hud thanh máu đạn điểm số minimap ui diegetic điều hướng tay cầm gamepad": "game-ui-and-hud-design",
        "hệ thống hội thoại dialogue game cây phân nhánh lựa chọn điều kiện cờ flag biến trạng thái quest": "dialogue-and-narrative-systems",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_agent_llm_batch_topics_match():
    from app import skills
    cases = {
        "thuật toán index vector hnsw ivf tìm lân cận gần đúng ann đánh đổi recall": "vector-index-algorithms",
        "mẫu workflow agent prompt chaining routing parallelization orchestrator workers": "agentic-workflow-patterns",
        "prompt caching tái dùng tiền tố ổn định giảm chi phí và độ trễ token lặp": "prompt-caching-and-context-reuse",
        "semantic cache cho llm theo ngữ nghĩa câu hỏi tương tự trên ngưỡng trả đáp án cũ": "semantic-caching-for-llms",
        "quan sát và trace agent nhiều bước span mỗi lần gọi llm và tool token chi phí": "agent-observability-and-tracing",
        "quản lý phiên bản prompt như artifact tách khỏi code a/b test rollback": "prompt-versioning-and-management",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_mit_frontend_batch_topics_match():
    from app import skills
    cases = {
        "crdt hợp nhất không xung đột local-first đồng bộ nhiều thiết bị yjs automerge": "crdts-and-local-first-sync",
        "máy trạng thái cho giao diện statechart xstate tránh mớ boolean cờ isloading": "state-machines-for-uis",
        "signals phản ứng mịn fine-grained solidjs preact signal không re-render cả cây": "signals-and-fine-grained-reactivity",
        "an toàn kiểu đầu cuối end-to-end trpc chia sẻ kiểu client server zod": "end-to-end-type-safety",
        "ứng dụng hướng hypermedia htmx trả html fragment thay vì json": "hypermedia-driven-apps",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_agent_protocols_and_rl_topics_match():
    from app import skills
    cases = {
        "giao thức giao tiếp giữa các agent a2a anp mcp phối hợp nhiều agent": "agent-communication-protocols",
        "agentic rl huấn luyện agent bằng grpo phần thưởng theo kết quả tác vụ verifiable": "agentic-rl-and-grpo",
    }
    for text, expected in cases.items():
        s = skills.find_matching_skill(text)
        assert s is not None and s.slug == expected, (text, s.slug if s else None)


def test_skills_repos_and_systems_topics_match():
    from app import skills
    cases = {
        "quản lý vòng đời kỹ năng agent và đánh giá chất lượng skill theo kết quả thực tế": "agent-skill-lifecycle-management",
        "tối ưu skill như huấn luyện với validation gate cho prompt": "skill-optimization-as-training",
        "dựng mô hình 3d thủ tục từ ảnh tham chiếu ra three.js": "image-to-procedural-3d",
        "tái tạo cảnh 3d thời gian thực từ video point cloud": "streaming-3d-reconstruction",
        "dịch pdf giữ nguyên bố cục và bảo toàn công thức code": "layout-preserving-document-translation",
        "kiến trúc multiplayer server giữ trạng thái thẩm quyền chống gian lận": "server-authoritative-multiplayer",
        "sinh id duy nhất phân tán snowflake sắp xếp được theo thời gian": "distributed-id-generation",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_crawl_and_pm_topics_match():
    from app import skills
    cases = {
        "crawl web cho llm rag render javascript tách nội dung chính markdown sạch giữ trích dẫn nguồn": "web-crawling-for-llms",
        "jobs to be done jtbd khách hàng thuê sản phẩm để tiến bộ trong công việc milkshake không phải tính năng": "jobs-to-be-done",
        "opportunity solution tree teresa torres nối outcome cơ hội giải pháp experiment tránh feature-first": "opportunity-solution-tree",
        "north star metric một chỉ số giá trị cốt lõi input metric vì sao doanh thu vanity là north star tệ": "north-star-metric",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_graphrag_and_trend_topics_match():
    from app import skills
    cases = {
        "graphrag rag dựa knowledge graph trích entity quan hệ community summarization global local search": "how-graphrag-works",
        "xếp hạng theo engagement upvote view odds cửa sổ recency gộp câu chuyện qua nhiều nền tảng trending": "engagement-and-recency-ranking",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_mlops_topics_match():
    from app import skills
    cases = {
        "feature store online offline chống training-serving skew tái dùng feature phục vụ nhất quán": "feature-stores",
        "training-serving skew model thấy feature khác lúc production rò rỉ dữ liệu leakage": "training-serving-skew",
        "theo dõi thí nghiệm ml mlflow ghi param metric artifact phiên bản dữ liệu tái lập": "experiment-tracking-and-reproducibility",
        "giám sát model ml data drift concept drift nhãn thật đến chậm khi nào retrain": "ml-model-monitoring-and-drift",
        "phục vụ triển khai model batch vs online real-time shadow canary champion challenger": "model-serving-and-deployment",
        "model registry versioning lineage staging production archived promote rollback": "model-registry-and-versioning",
        "gán nhãn dữ liệu chất lượng đồng thuận giữa người gán active learning giảm chi phí": "data-labeling-and-annotation",
        "đánh giá model precision recall f1 auc bẫy accuracy mất cân bằng calibration": "ml-model-evaluation-metrics",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_api_design_topics_match():
    from app import skills
    cases = {
        "nguyên tắc thiết kế rest api resource danh từ http method status code stateless": "rest-api-design-principles",
        "chiến lược versioning api url header media-type breaking change tương thích ngược deprecate": "api-versioning-strategies",
        "thiết kế xử lý lỗi api status code thân lỗi máy đọc problem details rfc 7807 không lộ nội bộ": "api-error-handling-design",
        "openapi swagger hợp đồng api đặc tả sinh code client server mock contract-first": "openapi-and-api-contracts",
        "thiết kế rate limit api 429 retry-after header limit remaining reset quota tier": "rate-limiting-api-design",
        "hateoas richardson maturity model các cấp rest hypermedia link trong response": "hateoas-and-rest-maturity",
        "backend for frontend bff một backend riêng cho web mobile gộp dữ liệu theo client": "backend-for-frontend-pattern",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_advanced_testing_topics_match():
    from app import skills
    cases = {
        "mutation testing cố tình chèn lỗi mutant xem test có bắt được đo chất lượng bộ test": "mutation-testing",
        "snapshot testing lưu output làm ảnh chụp báo lỗi khi khác rubber-stamp giòn": "snapshot-testing",
        "test flaky đậu rớt ngẫu nhiên timing async phụ thuộc thứ tự trạng thái chung": "flaky-test-diagnosis",
        "fuzzing ném input ngẫu nhiên méo mó tìm crash coverage-guided parser input không tin cậy": "fuzzing-basics",
        "quản lý dữ liệu test factory builder fixture cô lập dọn dẹp tránh dữ liệu chung": "test-data-management",
        "characterization test golden master ghi lại hành vi hiện tại code cũ lưới an toàn refactor": "characterization-testing",
        "kiểm thử code bất đồng bộ đồng thời await thay vì sleep fake clock race detector": "testing-async-and-concurrent-code",
        "độ phủ code line branch condition path coverage vì sao 100 phần trăm không có nghĩa test tốt": "coverage-metrics-and-limits",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_devops_iac_topics_match():
    from app import skills
    cases = {
        "terraform state file ánh xạ config với hạ tầng thật remote state locking drift không sửa tay": "how-terraform-state-works",
        "autoscaling mở rộng ngang dọc hpa cpu cluster autoscaler node tránh thrashing": "autoscaling-strategies",
        "helm đóng gói kubernetes chart template values release nâng cấp rollback vs kustomize": "helm-and-kubernetes-packaging",
        "kubernetes operator crd custom resource controller reconcile loop tự động vận hành": "kubernetes-operators-and-crds",
        "môi trường preview tạm thời cho mỗi pull request cô lập tự hủy seed dữ liệu chi phí": "ephemeral-preview-environments",
        "bảo mật chuỗi cung ứng phần mềm sbom ký artifact provenance slsa dependency bị chiếm": "software-supply-chain-security",
        "build tái lập tất định cùng input cùng output ghim dependency cache theo nội dung hermetic": "reproducible-builds-and-caching",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_agent_reach_topics_match():
    from app import skills
    cases = {
        "lớp năng lực trừu tượng tách cái gì khỏi cách làm danh sách backend fallback đổi thứ tự không viết lại code": "capability-abstraction-and-backend-routing",
        "cho agent khả năng đọc web nền tảng ngoài nguồn không cần cấu hình rss phụ đề fallback health check": "giving-agents-external-reach",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_gsap_and_frameworks_match():
    from app import skills
    cases = {
        "gsap greensock tween timeline easing animate web stagger": "gsap-animation",
        "scrolltrigger gsap hoạt hình theo cuộn pin scrub reveal on scroll": "gsap-scrolltrigger",
        "django orm queryset select_related n+1 migration": "django-patterns",
        "angular rxjs observable async pipe dependency injection onpush": "angular-patterns",
        "dotnet c# asp.net core async await dependency injection linq ef core": "dotnet-csharp-patterns",
        "spring boot java controller service repository jpa transactional dto": "spring-boot-patterns",
        "vue 3 composition api ref reactive computed composable": "vue-patterns",
        "flutter dart widget stateless stateful provider riverpod const": "flutter-patterns",
        "react native core component view text stylesheet flexbox flatList": "react-native-patterns",
        "kotlin android jetpack compose viewmodel coroutines flow lifecycle": "kotlin-android",
        "laravel php eloquent orm eager loading blade mass assignment queue": "laravel-php",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_repos_and_game_topics_match():
    from app import skills
    cases = {
        "email smtp imap mx record spf dkim dmarc": "how-email-works",
        "webrtc gọi video trình duyệt peer to peer stun turn ice": "how-webrtc-works",
        "captcha chống bot recaptcha thử thách hình ảnh behavioral": "how-captcha-works",
        "single sign-on sso identity provider saml oidc": "how-single-sign-on-works",
        "quy trình sparc specification pseudocode architecture refinement": "sparc-methodology",
        "goap goal oriented action planning ai game npc precondition effect": "goap-planning",
        "thiết kế game mda framework core loop động lực người chơi flow": "game-design-fundamentals",
        "game feel juice phản hồi tức thì screen shake hit-stop": "game-feel-and-juice",
        "cân bằng game đường cong độ khó economy sink source": "game-balancing",
        "tạo video ngắn ai tiktok script visual voiceover phụ đề": "ai-short-video-generation",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_milestone300_topics_match():
    from app import skills
    cases = {
        "qr code module đen trắng finder pattern sửa lỗi reed-solomon": "how-qr-codes-work",
        "gps định vị vệ tinh trilateration cần 4 vệ tinh": "how-gps-works",
        "dịch máy neural machine translation transformer attention": "how-machine-translation-works",
        "nhận dạng giọng nói speech to text spectrogram whisper": "how-speech-recognition-works",
        "định dạng ảnh jpeg png webp svg raster vector": "how-image-formats-work",
        "package manager npm pip giải quyết phụ thuộc transitive": "how-package-managers-work",
        "serverless faas lambda cold start tự scale trả theo lần chạy": "how-serverless-works",
        "reinforcement learning agent reward policy exploration exploitation": "reinforcement-learning-basics",
        "data pipeline etl elt batch streaming orchestration incremental": "data-pipelines-etl",
        "giải quyết xung đột mâu thuẫn nhóm tách người khỏi vấn đề": "conflict-resolution",
        "giao việc ủy quyền không ôm hết việc tránh micromanage": "delegation",
        "brainstorming động não tạo ý tưởng scamper phân kỳ hội tụ": "brainstorming-techniques",
        "mentoring coaching cố vấn huấn luyện hỏi thay vì bảo grow": "mentoring-and-coaching",
        "nhận diện khuôn mặt faceprint embedding bias quyền riêng tư": "how-face-recognition-works",
        "bộ lọc spam bayesian sender reputation spf dkim dmarc": "how-spam-filters-work",
        "graph database node edge neo4j cypher dữ liệu nhiều liên kết": "graph-databases",
        "stream processing xử lý luồng event time watermark windowing": "stream-processing",
        "trí tuệ cảm xúc eq tự nhận thức kiểm soát cảm xúc đồng cảm": "emotional-intelligence",
        "hệ thống ghi chú zettelkasten para tóm tắt bằng lời mình": "note-taking-systems",
        "feature engineering tạo đặc trưng encode categorical tránh data leakage": "feature-engineering",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_systems_and_life_topics_match():
    from app import skills
    cases = {
        "http method get post status code 404 500 header": "how-http-works",
        "cookie session duy trì đăng nhập httponly secure samesite": "how-cookies-and-sessions-work",
        "vpn mạng riêng ảo tunnel mã hóa ẩn ip có an toàn không": "how-vpns-work",
        "2fa xác thực hai yếu tố totp authenticator security key fido": "how-2fa-works",
        "thanh toán thẻ gateway authorization capture pci tokenization": "how-payment-processing-works",
        "video streaming adaptive bitrate hls dash buffering codec": "how-video-streaming-works",
        "webassembly wasm compile rust sang wasm sandbox tốc độ": "how-webassembly-works",
        "bloom filter kiểm tra tồn tại tiết kiệm bộ nhớ false positive": "how-bloom-filters-work",
        "consistent hashing băm nhất quán hash ring virtual node": "how-consistent-hashing-works",
        "nlp xử lý ngôn ngữ tự nhiên tokenization tf-idf word embedding": "nlp-basics",
        "computer vision thị giác máy tính cnn convolution object detection": "computer-vision-basics",
        "mlops đưa model lên production drift dữ liệu retrain monitor": "mlops-basics",
        "copywriting viết quảng cáo lợi ích vs tính năng aida pas headline": "copywriting-basics",
        "xây dựng thói quen tốt cue routine reward habit stacking": "habit-formation",
        "đặt mục tiêu okr objective key result smart goal": "goal-setting-okrs",
        "tài chính cá nhân lập ngân sách quỹ khẩn cấp trả nợ lãi cao lãi kép": "personal-finance-basics",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_professional_and_ml_topics_match():
    from app import skills
    cases = {
        "async await event loop microtask không chặn luồng": "how-async-await-works",
        "machine learning cơ bản supervised overfitting train test split": "machine-learning-basics",
        "hệ thống gợi ý collaborative filtering cold start": "how-recommendation-systems-work",
        "quản lý thời gian sắp xếp ưu tiên ma trận eisenhower": "time-management-and-prioritization",
        "ra quyết định khó reversible irreversible tránh phân tích tê liệt": "decision-making-frameworks",
        "cách góp ý feedback cho người khác và nhận phê bình sbi": "giving-and-receiving-feedback",
        "họp hiệu quả agenda action item có nên họp không": "effective-meetings",
        "viết email công việc hiệu quả subject bluf": "writing-effective-emails",
        "thuyết trình kể chuyện làm slide một ý một slide pitch": "presentations-and-storytelling",
        "đàm phán batna lợi ích vs lập trường win-win": "negotiation-basics",
        "học cách học active recall spaced repetition feynman": "learning-how-to-learn",
        "tư duy phản biện đánh giá lập luận ngụy biện thiên kiến bias": "critical-thinking",
        "sql join inner left outer bị nhân đôi dòng fan-out": "sql-joins-explained",
        "trực quan hóa dữ liệu chọn loại biểu đồ tránh chart gây hiểu lầm": "data-visualization-principles",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_network_crypto_ai_topics_match():
    from app import skills
    cases = {
        "tcp hoạt động thế nào three-way handshake retransmit flow control": "how-tcp-works",
        "dns phân giải tên miền recursive resolver bản ghi a cname ttl": "how-dns-works",
        "https tls handshake chứng chỉ ca đối xứng bất đối xứng": "how-https-tls-works",
        "hash table bảng băm hàm băm va chạm collision load factor": "how-hash-tables-work",
        "cdn edge server cache cache-control purge": "how-cdns-work",
        "load balancer cân bằng tải round-robin health check sticky session": "how-load-balancers-work",
        "nén dữ liệu lossless lossy huffman gzip jpeg": "how-compression-works",
        "mã hóa bất đối xứng khóa công khai riêng tư chữ ký số": "how-public-key-crypto-works",
        "jwt json web token header payload signature ký hmac": "how-jwt-works",
        "làm chatbot rag trên tài liệu riêng chunking embedding retrieval": "rag-fundamentals",
        "embedding vector tìm kiếm ngữ nghĩa cosine similarity vector database": "vector-embeddings",
        "khi nào fine-tune hay rag hay prompt để chỉnh llm": "fine-tuning-vs-rag-vs-prompting",
        "thống kê cơ bản mean median độ lệch chuẩn tương quan nhân quả": "statistics-fundamentals",
        "xác suất bayes cập nhật niềm tin base rate dương tính giả": "probability-and-bayes",
        "làm sạch dữ liệu xử lý missing outlier trùng lặp chuẩn hóa": "data-cleaning",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_backend_devops_topics_match():
    from app import skills
    cases = {
        "thiết kế xử lý lỗi exception vs result, fail fast không nuốt lỗi": "error-handling-patterns",
        "phân trang api dùng cursor keyset thay offset và lọc filter": "api-pagination-and-filtering",
        "giới hạn tần suất api bằng token bucket sliding window chống spam": "rate-limiting-algorithms",
        "thiết kế webhook gửi sự kiện có retry và ký hmac xác thực": "webhooks-design",
        "làm thao tác idempotent với idempotency key tránh xử lý trùng": "idempotency",
        "thử lại retry với exponential backoff và circuit breaker": "retries-and-resilience",
        "truy vết phân tán trace span correlation id tìm nút thắt latency": "distributed-tracing",
        "quản lý secret api key không commit dùng env và vault": "secrets-management",
        "kiểm thử tải load test đo latency p99 percentile tìm bottleneck": "load-testing",
        "chaos engineering tiêm lỗi có chủ đích kiểm thử độ bền": "chaos-engineering",
        "event sourcing cqrs lưu trạng thái bằng sự kiện tách read write": "event-sourcing-cqrs",
        "clean architecture phân tầng dependency rule hexagonal": "clean-architecture",
        "nhận biết code smell god class trùng lặp khi nào refactor": "code-smells",
        "quản lý nợ kỹ thuật khi nào refactor hay làm tính năng": "technical-debt",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_algorithm_patterns_match():
    from app import skills
    cases = {
        "tìm chuỗi con liên tiếp dài nhất dùng sliding window cửa sổ trượt": "sliding-window-pattern",
        "dùng hai con trỏ two pointers tìm cặp tổng trên mảng sorted": "two-pointers-pattern",
        "tìm kiếm nhị phân binary search vị trí đầu cuối": "binary-search-patterns",
        "liệt kê hoán vị tổ hợp bằng backtracking quay lui": "backtracking-pattern",
        "quy hoạch động dp knapsack với memoization": "dynamic-programming-patterns",
        "thuật toán tham lam greedy interval scheduling": "greedy-algorithms",
        "union find disjoint set đếm thành phần liên thông": "union-find",
        "dùng monotonic stack ngăn xếp đơn điệu next greater element": "monotonic-stack",
        "duyệt đồ thị bfs dfs tìm đường ngắn nhất không trọng số đếm đảo": "graph-traversal",
        "tìm đường đi ngắn nhất có trọng số dijkstra bellman-ford": "shortest-paths",
        "sắp xếp topo thứ tự phụ thuộc dependency kahn": "topological-sort",
        "dùng trie cây tiền tố làm autocomplete": "trie-prefix-tree",
        "dùng heap priority queue tìm top k phần tử lớn nhất": "heap-priority-queue",
        "thao tác bit bitwise xor tìm số xuất hiện lẻ lần và bitmask": "bit-manipulation",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_how_x_works_and_llm_topics_match():
    from app import skills
    cases = {
        "công cụ tìm kiếm hoạt động thế nào, inverted index và tf-idf bm25": "how-search-engines-work",
        "cpu bộ xử lý hoạt động thế nào, fetch decode execute pipeline cache": "how-cpus-work",
        "render đồ họa 3d thế nào, rasterization ray tracing shader z-buffer": "how-3d-rendering-works",
        "game engine hoạt động thế nào, game loop delta time entity component system": "how-game-engines-work",
        "react vue hoạt động thế nào, virtual dom diffing reactivity": "how-frontend-frameworks-work",
        "malloc cấp phát bộ nhớ thế nào, free list phân mảnh": "how-memory-allocators-work",
        "emulator giả lập máy thế nào, thông dịch opcode fetch decode": "how-emulators-work",
        "consensus phân tán thế nào, raft bầu leader replication quorum": "how-distributed-consensus-works",
        "trình soạn thảo lưu văn bản thế nào, gap buffer rope piece table": "how-text-editors-work",
        "bittorrent chia sẻ p2p thế nào, torrent seeder dht swarm": "how-bittorrent-works",
        "physics engine mô phỏng thế nào, tích phân euler va chạm collision": "how-physics-engines-work",
        "chạy llm trên máy local, ollama gguf quantization theo vram": "running-llms-locally",
        "serve llm nhiều người dùng, vllm continuous batching kv cache": "llm-inference-optimization",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_how_x_works_topics_match():
    from app import skills
    cases = {
        "git lưu trữ dữ liệu thế nào, cơ chế object blob tree commit bên trong": "how-git-works-internally",
        "database hoạt động bên trong ra sao, b-tree wal mvcc buffer pool": "how-databases-work",
        "container hoạt động thế nào, namespace cgroup overlay filesystem": "how-docker-containers-work",
        "web server xử lý request thế nào, event loop vs thread pool socket": "how-web-servers-work",
        "trình duyệt render trang thế nào, dom cssom layout reflow paint": "how-browsers-work",
        "compiler và interpreter hoạt động thế nào, lexer parser ast bytecode": "how-compilers-work",
        "shell chạy lệnh thế nào bên trong, fork exec pipe file descriptor": "how-shells-work",
        "regex engine khớp thế nào, nfa dfa backtracking redos": "how-regex-engines-work",
        "blockchain hoạt động thế nào, proof of work merkle hàm băm chuỗi khối": "how-blockchain-works",
        "mạng nơ-ron học thế nào, backpropagation gradient descent activation": "how-neural-networks-work",
        "hệ điều hành hoạt động thế nào, process scheduler virtual memory syscall": "how-operating-systems-work",
        "llm mô hình ngôn ngữ hoạt động thế nào, transformer attention dự đoán token": "how-llms-work",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_architecture_and_product_topics_match():
    from app import skills
    cases = {
        "lỡ tay git reset hard mất commit, dùng reflog khôi phục": "git-advanced",
        "thiết kế schema graphql, xử lý n+1 bằng dataloader": "graphql-design",
        "nên tách microservice không, ranh giới service và database per service": "microservices-and-boundaries",
        "quản lý dependency, semver lockfile và cập nhật thư viện an toàn": "dependency-management",
        "dùng feature flag triển khai dần canary và kill switch": "feature-flags-and-rollouts",
        "mô hình dữ liệu nosql mongodb, chọn partition key": "nosql-data-modeling",
        "app bị rò rỉ bộ nhớ memory leak, tốn ram tăng dần": "memory-management",
        "làm chat realtime, chọn websocket hay sse": "websockets-and-realtime",
        "hỗ trợ đa ngôn ngữ i18n, số nhiều plural theo locale": "i18n-and-localization",
        "hệ thống production sập, ứng phó sự cố và mitigate rollback": "incident-response",
        "kiểm chứng ý tưởng bằng mvp trước khi xây": "mvp-and-validation",
        "phỏng vấn khách hàng hỏi đúng câu hỏi không dẫn dắt để research": "user-research",
        "chạy a/b testing đúng cách, sample size và tránh p-hacking": "ab-testing",
        "nên test cái gì và bao nhiêu, test pyramid": "testing-strategy",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_infra_and_language_topics_match():
    from app import skills
    cases = {
        "pod bị crashloopbackoff trên kubernetes, sửa manifest sao": "kubernetes-basics",
        "lỗi kết nối mạng timeout, dns không resolve được host": "networking-fundamentals",
        "làm đăng nhập bảo mật, hash mật khẩu và session vs jwt": "authentication-and-authorization",
        "thêm cache nhưng dữ liệu cũ stale, invalidate cache sao cho đúng": "caching-strategies",
        "dùng message queue kafka, xử lý tin nhắn trùng lặp at-least-once": "message-queues-and-events",
        "viết python chuẩn pythonic, tránh mutable default argument": "python-best-practices",
        "component react bị render lại nhiều lần, useeffect chạy vô hạn": "react-patterns",
        "dàn trang css responsive dùng flexbox grid, căn giữa phần tử": "css-layout",
        "lập trình hàm pure function và immutable map filter reduce": "functional-programming",
        "thiết kế hướng đối tượng theo solid, composition over inheritance": "object-oriented-design",
        "lưu thời gian sai múi giờ, xử lý timezone utc và dst": "datetime-and-timezones",
        "chữ bị lỗi font mojibake ký tự lạ, encoding utf-8 sai": "encoding-and-unicode",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_core_engineering_topics_match():
    from app import skills
    cases = {
        "viết prompt cho ai sao cho ra kết quả đúng, tối ưu system prompt": "prompt-engineering",
        "query sql chạy chậm quá, đánh index và explain plan sao": "sql-query-optimization",
        "viết regex biểu thức chính quy để bắt chuỗi email": "regular-expressions",
        "viết dockerfile multi-stage cho image nhẹ hơn": "docker-containers",
        "dùng grep sed awk xử lý text trên terminal linux": "linux-command-line",
        "thiết kế schema database chuẩn hóa với khóa ngoại quan hệ": "database-schema-design",
        "thiết kế hệ thống chịu tải cao với load balancing và caching": "system-design-fundamentals",
        "code đa luồng bị race condition và deadlock, xử lý sao": "concurrency-and-parallelism",
        "phân tích dữ liệu bằng pandas, groupby và merge dataframe": "data-analysis-pandas",
        "tối ưu độ phức tạp big-o, chọn cấu trúc dữ liệu phù hợp": "algorithms-and-complexity",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_writing_process_topics_match():
    from app import skills
    cases = {
        "trả lời ngắn gọn thôi, bớt dài dòng lan man đi": "communicating-concisely",
        "viết sop quy trình chuẩn hướng dẫn từng bước cho nhân viên": "sop-writing",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_finance_and_product_topics_match():
    from app import skills
    cases = {
        "quản lý dòng tiền và unit economics ltv cac cho startup": "bootstrapped-cfo",
        "viết prd cho tính năng mới, ưu tiên p0 sao cho đúng": "prd-writing",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_career_topics_match():
    from app import skills
    cases = {
        "cv của tôi bị loại tự động qua ats, tối ưu từ khóa sao": "resume-ats-optimizer",
        "viết lại gạch đầu dòng cv cho mạnh, định lượng thành tích": "resume-bullet-writer",
        "chuẩn bị phỏng vấn xin việc, luyện star story trả lời câu hỏi hành vi": "interview-prep",
        "nên deal lương bao nhiêu, đàm phán lương khi nhận offer": "salary-negotiation",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_playwright_and_email_topics_match():
    from app import skills
    cases = {
        "test playwright bị flaky, nên dùng getByRole và web-first assertion": "playwright-testing",
        "cải thiện tỉ lệ mở email và deliverability cho chiến dịch email": "email-marketing",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_marketing_and_upgrade_topics_match():
    from app import skills
    cases = {
        "nâng cấp rails lên phiên bản mới mà giữ tùy biến, đừng chạy app:update": "rails-upgrade",
        "chiến lược marketing cho sản phẩm mới, tư duy marketing hệ thống": "marketing-principles",
        "làm rõ định vị thương hiệu và value proposition sản phẩm": "positioning-basics",
        "đánh giá landing page tối ưu chuyển đổi, headline chưa tốt": "homepage-audit",
        "viết chuỗi email cold outreach tiếp cận khách hàng linkedin": "cold-outreach-sequence",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_rails_and_creative_topics_match():
    from app import skills
    cases = {
        "review code rails 8 theo đúng quy ước native hotwire": "rails-conventions",
        "nghĩ ý tưởng sáng tạo cho chiến dịch quảng cáo dùng scamper": "creative-director",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_codebase_understanding_topics_match():
    from app import skills
    cases = {
        "giải thích hàm này làm gì và nó thuộc tầng nào liên kết với gì": "explaining-code-in-context",
        "viết tài liệu onboarding cho người mới vào team, lộ trình học codebase": "codebase-onboarding-guide",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_verification_topic_matches():
    from app import skills
    skill = skills.find_matching_skill("đừng tuyên bố hoàn thành khi chưa kiểm chứng, phải xác minh trước khi báo")
    assert skill is not None
    assert skill.slug == "verification-before-completion"


def test_skills_backend_and_method_topics_match():
    from app import skills
    cases = {
        "viết node.js bằng typescript dùng type stripping không cần build": "nodejs-typescript",
        "debug native crash segfault trong node.js core với v8": "nodejs-core",
        "xây rest api bằng fastify với schema validation và plugin": "fastify-best-practices",
        "chia nhỏ tính năng thành các lát dọc vertical slice có ticket": "vertical-slice-tickets",
        "review module terraform và tổ chức state backend": "terraform-and-opentofu",
        "có nhiều hướng giải chưa rõ chọn cái nào, review theo nguyên lý gốc": "first-principles-review",
        "xác định mục tiêu và điều kiện dừng trước khi bắt đầu làm": "goal-framing",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_context_engineering_topics_match():
    from app import skills
    cases = {
        "giải thích attention budget và cửa sổ ngữ cảnh hoạt động thế nào": "context-fundamentals",
        "phiên chat dài quá agent quên mất đã sửa file nào, nén ngữ cảnh sao": "context-compression",
        "giảm chi phí token bằng observation masking và kv-cache": "context-optimization",
        "xây dựng rubric nhiều chiều để đánh giá agent": "evaluation",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_dembrandt_advanced_topics_match():
    from app import skills
    cases = {
        "ảnh sản phẩm nên dùng dữ liệu thật hay lorem ipsum": "authentic-product-representation",
        "đặt bộ chọn ngôn ngữ và tiền tệ ở đâu trên header": "global-toolbar-controls",
        "mấy card này khác chiều cao so le nhau nhìn lệch": "repeated-component-alignment",
        "thêm breadcrumb cho điều hướng nhiều tầng": "ui-context-and-scope",
        "thiết kế tool điều phối kho vận cho nhân viên chuyên nghiệp": "operational-expert-tool-ui",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_accessibility_and_layout_topics_match():
    from app import skills
    cases = {
        "kiểm tra website này có đạt chuẩn wcag không": "wcag-accessibility",
        "nút bấm này thiếu trạng thái hover disabled": "button-states",
        "làm chế độ tối dark mode cho web": "color-mode-and-theme",
        "dùng shadow đổ bóng thế nào cho hợp lý": "elevation-and-depth",
        "nên dùng modal hay drawer cho tính năng này": "modal-and-overlay-patterns",
        "thiết kế thanh tab điều hướng": "tab-navigation",
        "vùng cuộn nội dung dài xử lý sao": "scroll-areas",
        "thanh header dính cố định khi cuộn trang": "sticky-and-fixed-elements",
        "giao diện quá chật hay quá thưa": "ui-density",
        "viết html ngữ nghĩa chuẩn seo": "semantic-html-and-seo",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_visual_and_data_topics_match():
    from app import skills
    cases = {
        "phân cấp thị giác nhấn mạnh nội dung quan trọng": "visual-emphasis-and-hierarchy",
        "thiết kế luồng người dùng dẫn dắt từng bước": "user-flows-and-guided-paths",
        "bố cục trang nhất quán xuyên suốt": "layout-paradigms-and-consistency",
        "các thành phần giao diện đồng bộ với nhau": "component-family-consistency",
        "ẩn dụ thực tế trong thiết kế giao diện": "real-world-metaphors",
        "tạo bảng màu tự động bằng thuật toán": "algorithmic-color-palette",
        "ngôn ngữ hình ảnh thương hiệu nhất quán": "brand-visual-language",
        "tối ưu core web vitals cho trang web": "performance-web-vitals",
        "chuyển động kể chuyện trong giao diện": "motion-and-storytelling",
        "nhiều view dữ liệu đồng bộ liên kết với nhau": "coordinated-data-views",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_languages_security_business_topics_match():
    from app import skills
    cases = {
        "giải thích ownership và borrow checker trong rust": "rust-ownership",
        "goroutine channel trong go bị leak thì sao": "go-concurrency",
        "discriminated union và utility type trong typescript": "typescript-advanced-types",
        "next.js app router server client component render thế nào": "nextjs-patterns",
        "làm sao mô hình hóa mối đe dọa stride khi thiết kế": "threat-modeling",
        "sveltekit runes state derived load function": "svelte-patterns",
        "xây dựng api node.js với express middleware": "express-and-node-apis",
        "owasp top 10 lỗ hổng web phổ biến broken access control": "owasp-top-10",
        "thêm content security policy csp và header bảo mật http": "security-headers",
        "tor hoạt động thế nào onion routing": "how-tor-works",
        "mã hóa dữ liệu lưu trữ encryption at rest envelope": "how-encryption-at-rest-works",
        "chỉ số saas mrr churn ltv cac": "saas-metrics",
        "phân tích phễu funnel và cohort retention": "cohort-and-funnel-analysis",
        "viết tài liệu kỹ thuật readme diataxis": "technical-writing",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_systems_how_x_topics_match():
    from app import skills
    cases = {
        "garbage collection hoạt động thế nào mark and sweep": "how-garbage-collection-works",
        "bộ nhớ ảo virtual memory page table mmu": "how-virtual-memory-works",
        "filesystem inode thư mục lưu file trên đĩa": "how-filesystems-work",
        "database transaction acid isolation level mvcc": "how-database-transactions-work",
        "b-tree chỉ mục database index cây cân bằng": "how-b-trees-work",
        "write ahead log wal durability crash recovery": "how-write-ahead-logging-works",
        "interpreter thông dịch bytecode máy ảo": "how-interpreters-work",
        "chữ ký số digital signature ký bằng khóa riêng": "how-digital-signatures-work",
        "certificate authority pki chain of trust tls": "how-certificate-authorities-work",
        "ntp đồng bộ thời gian clock skew": "how-ntp-time-sync-works",
        "raid gộp nhiều ổ đĩa striping mirroring parity": "how-raid-works",
        "mapreduce xử lý dữ liệu lớn map shuffle reduce": "how-mapreduce-works",
        "antivirus phát hiện malware signature behavioral": "how-antivirus-works",
        "text rendering hiển thị chữ font glyph shaping": "how-text-rendering-works",
        "audio codec nén âm thanh mp3 opus psychoacoustic": "how-audio-codecs-work",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_hardware_how_x_topics_match():
    from app import skills
    cases = {
        "ssd nand flash wear leveling trim hoạt động": "how-ssds-work",
        "ổ cứng từ hdd platter seek time sequential random": "how-hard-drives-work",
        "dram ram tụ điện refresh row buffer": "how-dram-works",
        "cpu cache l1 l2 l3 locality cache miss false sharing": "how-cpu-caches-work",
        "máy tính khởi động bios uefi bootloader nạp kernel": "how-computers-boot",
        "virtualization ảo hóa hypervisor máy ảo vm vs container": "how-virtualization-works",
        "cpu scheduling lập lịch tiến trình context switch round robin": "how-cpu-scheduling-works",
        "transistor logic gate cổng logic công tắc nhị phân": "how-transistors-and-logic-gates-work",
        "usb enumeration endpoint device class power delivery": "how-usb-works",
        "wifi mạng không dây băng tần kênh csma nhiễu": "how-wifi-works",
        "bluetooth ble frequency hopping pairing gatt": "how-bluetooth-works",
        "màn hình cảm ứng capacitive điện dung multi-touch": "how-touchscreens-work",
        "màn hình lcd oled pixel subpixel refresh rate": "how-displays-work",
        "máy ảnh số cảm biến cmos bayer raw jpeg phơi sáng": "how-digital-cameras-work",
        "gpu xử lý song song simt warp tăng tốc ml": "how-gpus-work",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_networking_how_x_topics_match():
    from app import skills
    cases = {
        "nat network address translation chia sẻ ip riêng port forwarding": "how-nat-works",
        "firewall tường lửa lọc gói tin stateful default deny": "how-firewalls-work",
        "proxy forward reverse proxy nginx tls termination": "how-proxies-work",
        "định tuyến internet bgp autonomous system router bảng": "how-internet-routing-works",
        "arp ánh xạ ip sang mac broadcast arp spoofing": "how-arp-works",
        "dhcp cấp phát ip tự động dora lease gateway dns": "how-dhcp-works",
        "địa chỉ ip subnet mask cidr ipv4 ipv6 private": "how-ip-addressing-works",
        "switch mạng forward theo mac address table vlan": "how-network-switches-work",
        "ssh shell từ xa an toàn key host key public key tunnel": "how-ssh-works",
        "grpc protocol buffers protobuf rpc http2 streaming": "how-grpc-works",
        "quic http3 udp head of line blocking 0-rtt": "how-quic-and-http3-work",
        "mqtt iot publish subscribe broker topic qos": "how-mqtt-works",
        "mạng di động cellular cell base station handoff sim 5g": "how-cellular-networks-work",
        "multicast gửi một tới nhiều igmp group": "how-multicast-works",
        "internet vệ tinh geo leo starlink độ trễ ground station": "how-satellite-internet-works",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_ai_ml_how_x_topics_match():
    from app import skills
    cases = {
        "transformer self-attention query key value multi-head": "how-transformers-work",
        "backpropagation lan truyền ngược chain rule gradient": "how-backpropagation-works",
        "gradient descent learning rate sgd adam optimizer": "how-gradient-descent-works",
        "tokenizer token hóa subword bpe vocabulary": "how-tokenizers-work",
        "diffusion model sinh ảnh khử nhiễu text to image": "how-diffusion-models-work",
        "gan generator discriminator huấn luyện đối kháng mode collapse": "how-gans-work",
        "cnn convolutional mạng tích chập bộ lọc pooling": "how-convolutional-networks-work",
        "rnn lstm mạng hồi quy hidden state vanishing gradient": "how-recurrent-networks-work",
        "vector database tìm kiếm tương đồng ann hnsw semantic": "how-vector-databases-work",
        "decision tree cây quyết định random forest gradient boosting": "how-decision-trees-work",
        "clustering phân cụm k-means dbscan không giám sát": "how-clustering-works",
        "overfitting quá khớp regularization dropout bias variance": "how-overfitting-and-regularization-work",
        "quantization lượng tử hóa int8 int4 nén mô hình chạy llm nhẹ": "how-model-quantization-works",
        "dimensionality reduction giảm chiều pca t-sne umap": "how-dimensionality-reduction-works",
        "anomaly detection phát hiện bất thường isolation forest outlier": "how-anomaly-detection-works",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_data_and_web_how_x_topics_match():
    from app import skills
    cases = {
        "columnar storage lưu trữ theo cột parquet olap nén": "how-columnar-storage-works",
        "data warehouse kho dữ liệu olap oltp star schema etl": "how-data-warehouses-work",
        "database replication sao chép leader follower replication lag": "how-database-replication-works",
        "database sharding phân mảnh shard key horizontal partition": "how-database-sharding-works",
        "lsm tree memtable sstable compaction tối ưu ghi rocksdb": "how-lsm-trees-work",
        "change data capture cdc tail transaction log debezium": "how-change-data-capture-works",
        "time series database chuỗi thời gian metrics downsampling retention": "how-time-series-databases-work",
        "serialization tuần tự hóa json vs binary protobuf schema evolution": "how-json-serialization-works",
        "http caching cache-control etag 304 cache busting": "how-http-caching-works",
        "cors same-origin policy preflight options access-control-allow-origin": "how-cors-works",
        "server-sent events sse eventsource text/event-stream reconnect": "how-server-sent-events-work",
        "graphql single endpoint resolver n+1 dataloader query mutation": "how-graphql-works",
        "oauth2 authorization code pkce access token refresh scope": "how-oauth2-flows-work",
        "websocket protocol upgrade handshake framing full-duplex ping pong": "how-websocket-protocol-works",
        "service worker offline pwa cache proxy chặn request push": "how-service-workers-work",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_algorithms_and_cs_topics_match():
    from app import skills
    cases = {
        "cryptographic hash sha-256 md5 collision băm mật khẩu salt": "how-cryptographic-hashing-works",
        "string matching tìm chuỗi con kmp rabin-karp boyer-moore": "string-matching-algorithms",
        "divide and conquer chia để trị merge sort master theorem": "divide-and-conquer",
        "sorting algorithm sắp xếp quicksort merge heapsort stability": "sorting-algorithms",
        "random number generation sinh số ngẫu nhiên prng csprng seed": "how-random-number-generation-works",
        "floating point số thực dấu phẩy động 0.1 + 0.2 ieee 754 nan": "floating-point-arithmetic",
        "segment tree fenwick binary indexed range query o(log n)": "segment-and-fenwick-trees",
        "minimum spanning tree cây khung nhỏ nhất kruskal prim": "minimum-spanning-tree",
        "maximum flow min cut luồng cực đại ford-fulkerson bipartite matching": "maximum-flow-and-min-cut",
        "skip list danh sách bỏ qua nhiều tầng redis sorted set": "skip-lists",
        "reservoir sampling lấy mẫu ngẫu nhiên từ luồng k phần tử một lần duyệt": "reservoir-sampling",
        "sweep line đường quét interval overlap meeting rooms": "sweep-line-algorithms",
        "lru cache hash map doubly linked list eviction o(1)": "lru-cache-design",
        "number theory gcd euclid modular exponentiation sieve nguyên tố": "number-theory-for-programmers",
        "game theory lý thuyết trò chơi nash equilibrium prisoner dilemma minimax": "game-theory-basics",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_product_business_topics_match():
    from app import skills
    cases = {
        "product discovery khám phá sản phẩm phỏng vấn khách hàng jobs to be done": "product-discovery",
        "user story mapping bản đồ câu chuyện backbone slice mvp": "user-story-mapping",
        "roadmap prioritization ưu tiên lộ trình rice moscow": "roadmap-prioritization",
        "pricing strategy chiến lược định giá value-based willingness to pay": "pricing-strategy",
        "go-to-market gtm đưa sản phẩm ra thị trường icp product-led": "go-to-market-strategy",
        "competitive analysis phân tích đối thủ cạnh tranh ma trận": "competitive-analysis",
        "b2b sales bán hàng doanh nghiệp pipeline qualify champion": "b2b-sales-fundamentals",
        "content strategy chiến lược nội dung pillar phễu phân phối": "content-strategy",
        "seo content nghiên cứu từ khóa search intent backlink": "seo-content-strategy",
        "growth loop vòng lặp tăng trưởng viral k-factor": "growth-loops-and-virality",
        "community building xây dựng cộng đồng belonging seed moderation": "community-building",
        "stakeholder management quản lý các bên liên quan power interest buy-in": "stakeholder-management",
        "agile scrum kanban sprint standup retrospective wip": "agile-and-scrum",
        "project risk management quản lý rủi ro dự án likelihood impact register": "project-risk-management",
        "career growth engineer phát triển sự nghiệp ic vs management promotion": "career-growth-for-engineers",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_engineering_practices_topics_match():
    from app import skills
    cases = {
        "database migration di trú schema zero-downtime expand contract backfill": "database-migrations",
        "deployment strategy triển khai blue-green canary rolling rollback": "deployment-strategies",
        "slo sli error budget mục tiêu độ tin cậy sre": "slos-and-error-budgets",
        "capacity planning lập kế hoạch dung lượng peak headroom autoscaling": "capacity-planning",
        "connection pooling gộp kết nối database too many connections": "connection-pooling",
        "n+1 query problem orm lazy loading eager join dataloader": "n-plus-one-query-problem",
        "saga pattern giao dịch phân tán microservices compensating orchestration": "saga-pattern",
        "two-phase commit 2pc prepare commit coordinator blocking": "two-phase-commit",
        "distributed locking khóa phân tán redis ttl fencing token": "distributed-locking",
        "leader election bầu chọn leader raft split-brain quorum failover": "leader-election",
        "object storage lưu trữ đối tượng s3 bucket presigned url": "object-storage",
        "materialized view khung nhìn cụ thể hóa precompute refresh": "materialized-views",
        "contract testing kiểm thử hợp đồng api pact consumer provider": "contract-testing",
        "property-based testing sinh input ngẫu nhiên hypothesis shrinking": "property-based-testing",
        "test double mocking stub mock fake over-mocking cô lập": "test-doubles-and-mocking",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_devops_cloud_topics_match():
    from app import skills
    cases = {
        "gitops git nguồn sự thật argocd flux reconcile declarative": "gitops",
        "service mesh sidecar proxy mtls istio linkerd data plane": "service-mesh",
        "api gateway cổng api điểm vào duy nhất bff routing auth": "api-gateway-patterns",
        "event-driven architecture kiến trúc hướng sự kiện kafka pub sub decoupling": "event-driven-architecture",
        "cloud cost optimization tối ưu chi phí đám mây finops reserved spot": "cloud-cost-optimization",
        "disaster recovery backup rpo rto 3-2-1 restore": "disaster-recovery-and-backups",
        "immutable infrastructure hạ tầng bất biến cattle pets config drift golden image": "immutable-infrastructure",
        "configuration management quản lý cấu hình env var tách config ansible": "configuration-management",
        "container image optimization tối ưu image docker multi-stage layer caching distroless": "container-image-optimization",
        "kubernetes networking service clusterip ingress dns pod ip": "kubernetes-networking",
        "monitoring alerting golden signals alert fatigue actionable symptom": "monitoring-and-alerting",
        "runbook oncall ca trực rotation escalation toil blameless": "runbooks-and-oncall",
        "serverless architecture faas cold start stateless lambda": "serverless-architecture-patterns",
        "edge computing tính toán tại biên edge function giảm độ trễ cdn": "edge-computing",
        "policy as code opa rego admission controller guardrail tuân thủ": "policy-as-code",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_frontend_web_topics_match():
    from app import skills
    cases = {
        "css cascade specificity độ ưu tiên quy tắc nào thắng important": "css-cascade-and-specificity",
        "state management frontend local global server state redux signals": "state-management-patterns",
        "virtual dom reconciliation diff key trong list react": "virtual-dom-and-reconciliation",
        "browser storage localstorage cookie indexeddb lưu token": "browser-storage",
        "css architecture bem tailwind utility-first css modules scoped": "css-architecture",
        "web bundler webpack vite tree shaking code splitting": "web-build-tools-and-bundlers",
        "tối ưu ảnh web webp avif responsive srcset lazy loading": "web-image-optimization",
        "progressive enhancement html trước baseline graceful degradation": "progressive-enhancement",
        "webauthn passkey đăng nhập không mật khẩu fido2 chống phishing": "webauthn-and-passkeys",
        "data fetching react query swr caching optimistic loading state": "data-fetching-patterns",
        "debounce throttle giới hạn tần suất search scroll": "debouncing-and-throttling",
        "web worker luồng nền offload việc nặng giữ ui mượt postmessage": "web-workers",
        "rendering patterns csr ssr ssg isr hydration islands": "rendering-patterns",
        "micro-frontends chia nhỏ frontend module federation đội độc lập": "micro-frontends",
        "design tokens biến thiết kế semantic primitive theming dark mode": "design-tokens",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_osint_monitoring_topics_match():
    from app import skills
    cases = {
        "osint tình báo nguồn mở thu thập thông tin công khai chu trình": "osint-fundamentals",
        "news aggregation rss feed tổng hợp tin tức polling dedup": "news-aggregation-and-rss",
        "web scraping cào dữ liệu parse html trang động robots.txt": "web-scraping-fundamentals",
        "event detection alerting phát hiện sự kiện spike burst breaking news": "event-detection-and-alerting",
        "data source reliability độ tin cậy nguồn tin corroboration primary": "data-source-reliability",
        "information verification fact-check xác minh reverse image geolocation": "information-verification",
        "media monitoring social listening giám sát mention share of voice bot": "media-monitoring-and-social-listening",
        "sentiment analysis trend phân tích cảm xúc xu hướng aspect sarcasm": "sentiment-and-trend-analysis",
        "geospatial mapping geocoding tọa độ projection geojson spatial index": "geospatial-mapping-and-geocoding",
        "real-time dashboard bảng giám sát thời gian thực glanceable live update": "real-time-monitoring-dashboards",
        "entity resolution deduplication khử trùng lặp fuzzy matching blocking": "entity-resolution-and-deduplication",
        "geopolitical risk analysis rủi ro địa chính trị scenario indicators warnings": "geopolitical-risk-analysis",
        "crisis monitoring giám sát khủng hoảng cảnh báo sớm situational awareness": "crisis-monitoring",
        "monitoring pipeline ingest normalize enrich detect alert backpressure": "monitoring-pipeline-design",
        "data journalism storytelling kể chuyện bằng dữ liệu tránh biểu đồ gây hiểu lầm": "data-journalism-and-storytelling",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_llm_app_topics_match():
    from app import skills
    cases = {
        "agentic rag agent lập luận truy xuất nhiều bước multi-hop self-correct": "agentic-rag",
        "rag chunking chia nhỏ tài liệu kích thước chunk overlap semantic": "rag-chunking-strategies",
        "rag retrieval reranking hybrid search bm25 cross-encoder top-k": "rag-retrieval-and-reranking",
        "rag evaluation đánh giá faithfulness groundedness llm as judge": "rag-evaluation",
        "structured output llm json schema constrained decoding validate": "structured-output-from-llms",
        "llm function calling tool definition arguments agent loop": "llm-function-calling",
        "prompt injection defense tiêm lệnh indirect injection jailbreak least privilege": "prompt-injection-defense",
        "hallucination mitigation giảm ảo giác grounding citation không biết": "hallucination-mitigation",
        "llm guardrails safety lọc input output moderation pii": "llm-guardrails-and-safety",
        "llm cost latency optimization giảm chi phí model routing prompt caching streaming": "llm-cost-and-latency-optimization",
        "semantic caching cache theo ngữ nghĩa embedding similarity threshold": "semantic-caching",
        "agent planning react plan and execute reflection lập kế hoạch": "agent-planning-patterns",
        "document parsing rag trích xuất pdf bảng ocr layout ingest": "document-parsing-for-rag",
        "llm observability tracing chuỗi agent log token cost quality": "llm-observability",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_trading_topics_match():
    from app import skills
    cases = {
        "technical analysis indicators chỉ báo moving average rsi macd bollinger": "technical-analysis-indicators",
        "algorithmic trading strategies trend following mean reversion arbitrage": "algorithmic-trading-strategies",
        "backtesting trading lookahead survivorship overfitting slippage out-of-sample": "backtesting-trading-strategies",
        "risk management trading position sizing stop-loss kelly drawdown": "risk-management-in-trading",
        "order types execution market limit stop slippage twap vwap": "order-types-and-execution",
        "market microstructure order book bid ask spread liquidity depth": "market-microstructure",
        "portfolio theory diversification correlation sharpe efficient frontier": "portfolio-theory-and-diversification",
        "quantitative trading signals alpha beta factor model signal decay": "quantitative-trading-signals",
        "crypto defi trading dex amm liquidity pool impermanent loss gas mev": "crypto-and-defi-trading",
        "options derivatives call put greeks delta theta vega leverage": "options-and-derivatives-basics",
        "high frequency trading hft latency colocation speed arms race": "high-frequency-trading-concepts",
        "trading psychology tâm lý fear greed loss aversion revenge fomo": "trading-psychology",
        "trading bot architecture kiến trúc bot data signal risk execution kill-switch": "trading-bot-architecture",
        "market data ohlcv tick candle level 2 order book split adjusted": "market-data-fundamentals",
        "sentiment-driven trading news social signal already priced in manipulation": "sentiment-driven-trading",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_edtech_tutoring_topics_match():
    from app import skills
    cases = {
        "intelligent tutoring systems its mô hình miền học viên sư phạm adaptive": "intelligent-tutoring-systems",
        "socratic method dạy bằng câu hỏi thay vì trả lời dẫn dắt": "socratic-method-teaching",
        "spaced repetition srs đường cong quên anki khoảng cách tăng dần": "spaced-repetition-systems",
        "adaptive personalized learning điều chỉnh độ khó tốc độ lộ trình": "adaptive-and-personalized-learning",
        "formative assessment feedback đánh giá quá trình phản hồi kịp thời": "formative-assessment-and-feedback",
        "scaffolding zpd giàn giáo vùng phát triển gần fading": "scaffolding-and-zpd",
        "cognitive load theory tải nhận thức bộ nhớ làm việc worked example": "cognitive-load-theory",
        "mastery learning học theo thành thạo trước khi tiến remediation": "mastery-learning",
        "blooms taxonomy learning objectives thang bloom mục tiêu động từ đo được": "blooms-taxonomy-and-objectives",
        "knowledge tracing mô hình hóa kiến thức bayesian deep student model": "knowledge-tracing",
        "building an ai tutor gia sư llm dẫn dắt không cho đáp án grounding": "building-an-ai-tutor",
        "misconception diagnosis chẩn đoán quan niệm sai lỗi có hệ thống đối chất": "misconception-diagnosis",
        "learner motivation engagement động lực growth mindset gamification": "learner-motivation-and-engagement",
        "curriculum sequencing sắp xếp chương trình tiên quyết spiral interleaving": "curriculum-sequencing",
        "retrieval practice testing effect luyện truy hồi active recall đọc lại yếu": "retrieval-practice-and-testing-effect",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_generative_media_topics_match():
    from app import skills
    cases = {
        "how video generation works sinh video ai temporal consistency diffusion": "how-video-generation-works",
        "image to video animation làm ảnh tĩnh chuyển động motion conditioning": "image-to-video-and-animation",
        "motion capture pose estimation bắt chuyển động keypoint retargeting": "motion-capture-and-pose-estimation",
        "keyframe animation interpolation khung chính tween easing timing": "keyframe-animation-and-interpolation",
        "camera cinematography cỡ cảnh góc máy chuyển động bố cục": "camera-and-cinematography-basics",
        "storyboarding shot planning dựng storyboard lập kế hoạch cảnh continuity": "storyboarding-and-shot-planning",
        "video editing dựng phim cut pacing transition j-cut audio": "video-editing-fundamentals",
        "color grading chỉnh màu white balance lut mood consistency": "color-grading-basics",
        "prompt engineering visual media viết prompt sinh ảnh style lighting camera": "prompt-engineering-for-visual-media",
        "controllable image generation controlnet pose depth inpainting lora reference": "controllable-image-generation",
        "audio music generation sinh âm thanh nhạc tts voice cloning": "audio-and-music-generation",
        "lip-sync talking heads đồng bộ môi khuôn mặt nói viseme uncanny": "lip-sync-and-talking-heads",
        "3d scene representation nerf gaussian splatting dựng 3d từ ảnh": "3d-scene-representation",
        "generative media pipeline quy trình sản xuất video ai consistency cross-shot": "generative-media-pipeline",
        "ai avatar character animation rig skeleton vtuber điều khiển bằng motion": "ai-avatar-and-character-animation",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_spec_driven_and_agent_topics_match():
    from app import skills
    cases = {
        "writing a project constitution hiến pháp nguyên tắc nền tảng ràng buộc agent": "writing-a-project-constitution",
        "writing executable specifications viết đặc tả what why acceptance criteria": "writing-executable-specifications",
        "spec clarification ambiguity làm rõ đặc tả mơ hồ câu hỏi trước khi build": "spec-clarification-and-ambiguity-resolution",
        "technical planning from specs lập kế hoạch kỹ thuật kiến trúc từ đặc tả": "technical-planning-from-specs",
        "task decomposition implementation chia kế hoạch thành task vertical slice": "task-decomposition-for-implementation",
        "spec plan consistency analysis đối chiếu spec plan tasks phủ sóng traceability": "spec-plan-consistency-analysis",
        "working with ai coding agents làm việc agent lập trình review giữ kiểm soát": "working-with-ai-coding-agents",
        "prompting coding agents viết chỉ thị cho agent intent constraint done criteria": "prompting-coding-agents",
        "reviewing ai generated code review code ai hallucinate api plausible sai": "reviewing-ai-generated-code",
        "context for coding agents claude.md agents.md ngữ cảnh quy ước": "context-for-coding-agents",
        "iterative development with agents lặp instruct review refine phản hồi cụ thể": "iterative-development-with-agents",
        "preventing agent over-engineering ngăn agent làm quá phức tạp yagni gold-plating": "preventing-agent-over-engineering",
        "human in the loop development con người quyết định checkpoint hành động hệ trọng": "human-in-the-loop-development",
        "intent-driven development tách intent what why khỏi how nguồn sự thật": "intent-driven-development",
        "living documentation tài liệu sống đồng bộ code docs-as-code sinh từ nguồn": "living-documentation",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_build_your_own_x_extra_topics_match():
    from app import skills
    cases = {
        "template engine placeholder parse compile render autoescape": "how-template-engines-work",
        "uuid guid định danh 128 bit v4 v7 không cần phối hợp": "how-uuids-work",
        "diff algorithm longest common subsequence myers so sánh khác biệt": "how-diff-algorithms-work",
        "augmented reality ar slam visual inertial theo dõi vị trí": "how-augmented-reality-works",
        "voxel engine thế giới khối chunking meshing procedural minecraft": "how-voxel-engines-work",
        "chess engine cờ vua minimax alpha-beta evaluation transposition": "how-chess-engines-work",
        "spreadsheet excel công thức đồ thị phụ thuộc tính lại topological": "how-spreadsheet-engines-work",
        "parser biến token thành cây recursive descent grammar precedence": "how-parsers-work",
        "terminal emulator pty escape ansi sequence shell vs terminal": "how-terminal-emulators-work",
        "key-value store get put log index lsm redis": "how-key-value-stores-work",
        "ray tracing mô phỏng ánh sáng tia phản xạ khúc xạ bóng": "how-ray-tracing-works",
        "container runtime namespaces cgroups overlay filesystem tiến trình cô lập": "how-container-runtimes-work",
        "network stack phân tầng link ip tcp application encapsulation": "how-network-stacks-work",
        "syntax highlighting tô màu cú pháp lexing token textmate tree-sitter": "how-syntax-highlighting-works",
        "url shortener mã ngắn base62 redirect 301 302 cache đọc nhiều": "how-url-shorteners-work",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_advanced_llm_ml_topics_match():
    from app import skills
    cases = {
        "rlhf là gì reward model ppo và dpo căn chỉnh": "how-rlhf-works",
        "giải thích mixture of experts moe router sparse": "how-mixture-of-experts-works",
        "lora fine-tuning adapter low-rank qlora peft": "how-lora-fine-tuning-works",
        "knowledge distillation student teacher dark knowledge": "how-knowledge-distillation-works",
        "beam search giữ top-k chuỗi khi giải mã dịch máy": "how-beam-search-works",
        "temperature top-p top-k tham số sinh llm ngẫu nhiên": "how-llm-sampling-works",
        "speculative decoding draft model tăng tốc suy luận": "how-speculative-decoding-works",
        "word embedding word2vec glove vector nghĩa": "how-word-embeddings-work",
        "positional encoding rope sinusoidal mã hóa vị trí": "how-positional-encoding-works",
        "batch norm layer norm chuẩn hóa activation mạng sâu": "how-normalization-in-networks-works",
        "relu gelu sigmoid hàm kích hoạt phi tuyến": "how-activation-functions-work",
        "kv cache bộ nhớ suy luận llm prefill decode": "how-kv-cache-works",
        "multimodal đa phương thức vision language ảnh văn bản": "how-multimodal-models-work",
        "model context protocol mcp tools resources prompts": "how-model-context-protocol-works",
        "vision transformer vit chia ảnh thành patch token": "how-vision-transformers-work",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_messaging_reliability_topics_match():
    from app import skills
    cases = {
        "backpressure áp lực ngược producer nhanh consumer chậm bounded queue": "backpressure-and-flow-control",
        "dead letter queue dlq poison message hàng đợi thư chết": "dead-letter-queues",
        "transactional outbox tránh dual write phát sự kiện tin cậy": "transactional-outbox",
        "delivery semantics at-least-once exactly-once ngữ nghĩa phân phối": "delivery-semantics",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_fintech_payments_topics_match():
    from app import skills
    cases = {
        "double entry accounting kế toán kép ghi nợ ghi có bằng nhau": "how-double-entry-accounting-works",
        "thiết kế ledger sổ cái ví tiền entry bất biến append-only": "designing-a-ledger",
        "xử lý tiền tệ không dùng float lưu số nguyên cents đa tiền tệ": "money-and-currency-handling",
        "payment reconciliation đối soát khớp sổ với sao kê ngân hàng settlement": "payment-reconciliation",
        "subscription billing thanh toán định kỳ proration dunning đổi gói": "subscription-billing-design",
        "ach wire transfer chuyển khoản ngân hàng gom lô không đảo": "how-ach-and-wire-transfers-work",
        "thuế và hóa đơn sales tax vat gst hóa đơn tuần tự credit note": "tax-and-invoicing-basics",
        "pci dss không lưu số thẻ tokenization hosted field saq a": "pci-dss-basics",
        "3d secure sca xác thực mạnh chuyển trách nhiệm gian lận": "how-3d-secure-works",
        "fraud detection phát hiện gian lận velocity false positive review queue": "fraud-detection-basics",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_web_platform_topics_match():
    from app import skills
    cases = {
        "spf dkim dmarc xác thực email chống giả mạo spoofing": "how-spf-dkim-dmarc-work",
        "email deliverability khả năng vào inbox warm-up reputation list hygiene": "email-deliverability",
        "browser rendering critical rendering path dom cssom layout paint composite": "how-browser-rendering-works",
        "web workers luồng nền javascript tính toán nặng không đơ ui": "how-web-workers-work",
        "progressive web app pwa cài được chạy offline manifest service worker": "how-progressive-web-apps-work",
        "source map ánh xạ code minified về source gốc debug stack trace": "how-source-maps-work",
        "lazy loading tải lười code splitting hoãn tải ảnh ngoài màn hình": "how-lazy-loading-works",
        "content security policy csp chặn script nội tuyến chống xss nonce": "how-content-security-policy-works",
        "slowly changing dimensions scd type 2 lưu lịch sử phiên bản chiều": "slowly-changing-dimensions",
        "window function sql running total xếp hạng row_number over partition by": "window-functions-in-sql",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_mobile_topics_match():
    from app import skills
    cases = {
        "kiến trúc app di động tách ui khỏi logic mvvm mvi repository": "mobile-app-architecture",
        "push notification thông báo đẩy device token apns fcm": "how-push-notifications-work",
        "offline-first local db nguồn sự thật đồng bộ hai chiều conflict": "offline-first-mobile-sync",
        "hiệu năng app di động khởi động nhanh cold start cuộn mượt 60fps jank": "mobile-app-performance",
        "app store optimization aso từ khóa icon screenshot đánh giá ranking": "app-store-optimization",
        "deep linking universal link app link mở màn hình app deferred": "how-deep-linking-works",
        "điều hướng di động tab bar drawer stack back stack nút back": "mobile-navigation-patterns",
        "responsive vs adaptive nhiều kích thước màn hình tablet foldable notch": "responsive-vs-adaptive-mobile",
        "tiết kiệm pin và mạng gom việc nền doze giảm đánh thức radio": "battery-and-network-efficiency",
        "khôi phục trạng thái process death os giết app xoay màn hình instance state": "mobile-state-restoration",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_game_dev_topics_match():
    from app import skills
    cases = {
        "quản lý trạng thái game scene menu chơi tạm dừng game over state stack": "game-state-management",
        "sprite animation game sprite sheet atlas máy trạng thái idle run jump": "sprite-and-animation-systems",
        "pathfinding trong game tìm đường npc navmesh a-star steering tránh va chạm": "how-pathfinding-in-games-works",
        "netcode multiplayer server authoritative client prediction interpolation rollback": "game-networking-and-netcode",
        "procedural generation sinh nội dung địa hình noise perlin seed wave function collapse": "procedural-generation-in-games",
        "game camera theo người chơi damping dead zone look-ahead va chạm 3d": "how-game-cameras-work",
        "game ai npc behavior finite state machine behavior tree utility ai đáng tin": "game-ai-behavior",
        "object pooling tái sử dụng đạn particle tránh giật garbage collection": "object-pooling-in-games",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_privacy_crypto_topics_match():
    from app import skills
    cases = {
        "gdpr quyền riêng tư dữ liệu cơ sở pháp lý tối thiểu hóa quyền xóa": "gdpr-and-data-privacy",
        "xử lý pii tối thiểu hóa thu thập cô lập mã hóa loại khỏi log": "pii-handling-and-minimization",
        "lưu giữ và xóa dữ liệu retention hard soft delete xóa lan tỏa backup": "data-retention-and-deletion",
        "quản lý đồng ý consent opt-in granular ghi bằng chứng cookie rút lại": "consent-management",
        "ẩn danh và bí danh hóa anonymization pseudonymization tái định danh k-anonymity": "data-anonymization-and-pseudonymization",
        "mã hóa đối xứng vs bất đối xứng khóa công khai riêng tư hybrid": "symmetric-vs-asymmetric-encryption",
        "băm và lưu mật khẩu bcrypt argon2 salt hash chậm không plaintext": "password-hashing-and-storage",
        "nhật ký kiểm toán bảo mật audit log append-only chống giả mạo ai làm gì": "security-audit-logging",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_compiler_pl_topics_match():
    from app import skills
    cases = {
        "type checker kiểm tra kiểu tĩnh duyệt ast luật kiểu nominal structural soundness": "how-type-checkers-work",
        "type inference hindley milner suy luận kiểu unification let-polymorphism": "type-inference-hindley-milner",
        "compiler tối ưu code constant folding dead code inlining ssa register allocation": "how-compilers-optimize-code",
        "jit compiler biên dịch nóng lúc chạy profiling deoptimization tiered warmup": "how-jit-compilers-work",
        "calling convention abi truyền tham số thanh ghi stack frame ffi tương thích nhị phân": "calling-conventions-and-abi",
        "memory model happens-before data race reordering atomics memory barrier đồng bộ": "memory-models-and-happens-before",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_blockchain_web3_topics_match():
    from app import skills
    cases = {
        "smart contract hợp đồng thông minh code tự thực thi trên blockchain bất biến gas": "how-smart-contracts-work",
        "ethereum evm world computer tài khoản trạng thái toàn cục bytecode gas": "how-ethereum-and-evm-work",
        "ví crypto khóa riêng seed phrase custodial non-custodial mất khóa mất tiền": "how-crypto-wallets-and-keys-work",
        "token nft erc-20 erc-721 fungible non-fungible metadata off-chain": "how-nfts-and-tokens-work",
        "lỗ hổng smart contract reentrancy tràn số nguyên thao túng oracle front-running mev": "blockchain-security-pitfalls",
        "oracle blockchain đưa dữ liệu off-chain lên chain thao túng giá flash loan chainlink": "how-oracles-work",
        "stablecoin neo giá usd dự trữ fiat thế chấp crypto thuật toán depeg": "how-stablecoins-work",
        "defi cho vay vay mượn thế chấp vượt mức thanh lý yield composability money legos": "how-defi-primitives-work",
        "rollup layer 2 mở rộng blockchain optimistic fraud proof zk validity proof": "how-rollups-and-layer2-work",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_audio_dsp_topics_match():
    from app import skills
    cases = {
        "âm thanh số sample rate bit depth pcm dải động 44.1khz": "how-digital-audio-works",
        "biến đổi fourier phân tích tín hiệu thành sóng sin miền tần số fft phổ": "how-fourier-transform-works",
        "lấy mẫu nyquist aliasing tần số cao giả dạng thấp bộ lọc chống aliasing": "how-sampling-and-aliasing-works",
        "bộ lọc âm thanh eq low-pass high-pass cutoff resonance định hình tần số": "how-audio-filters-work",
        "tổng hợp âm thanh oscillator dạng sóng envelope adsr subtractive fm synthesizer": "how-audio-synthesis-works",
        "reverb delay echo chorus flanger phaser compressor hiệu ứng không gian": "how-reverb-and-effects-work",
        "đổi cao độ pitch shift kéo giãn thời gian time stretch phase vocoder formant": "how-pitch-and-time-work",
        "khử nhiễu âm thanh spectral subtraction hồ sơ nhiễu musical noise ml denoiser": "how-noise-reduction-works",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_image_generation_topics_match():
    from app import skills
    cases = {
        "workflow vẽ ảnh dạng node comfyui đồ thị load model encode prompt sample decode": "node-based-image-workflows",
        "viết prompt cho model vẽ ảnh cấu trúc chủ thể phong cách trọng số nhấn mạnh": "prompting-for-image-models",
        "sampler scheduler diffusion euler dpm++ ddim số bước steps karras": "image-model-samplers-and-schedulers",
        "negative prompt và cfg guidance scale bám prompt ảnh bị cháy fried": "negative-prompts-and-cfg",
        "nâng cấp phóng to ảnh super resolution esrgan hires fix tiled upscale": "image-upscaling-methods",
        "textual inversion embedding dạy model khái niệm mới vài ảnh so với lora dreambooth": "textual-inversion-and-embeddings",
        "độ phân giải tỉ lệ khung ảnh vẽ quá lớn nhân đôi chủ thể vẽ nhỏ rồi upscale": "image-generation-resolution-and-aspect",
        "seed tái lập vẽ ảnh cùng seed cùng cài đặt cùng ảnh cố định seed tinh chỉnh": "seed-and-reproducibility-image-gen",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_skills_advanced_data_engineering_topics_match():
    from app import skills
    cases = {
        "lakehouse iceberg delta hudi acid transaction time travel table format trên parquet": "data-lakehouse-and-table-formats",
        "query optimizer cost based kế hoạch thực thi thống kê bảng thứ tự join explain": "how-query-optimizers-work",
        "mpp database xử lý song song nhiều node shuffle di chuyển dữ liệu distribution key skew": "how-mpp-databases-work",
        "chiến lược phân vùng dữ liệu partition theo thời gian range hash pruning tránh quá nhiều": "data-partitioning-strategies",
        "data contract schema evolution tương thích ngược producer consumer chống vỡ pipeline": "data-contracts-and-schema-evolution",
        "data lineage observability theo dõi dữ liệu đến từ đâu freshness volume schema distribution": "data-lineage-and-observability",
        "reverse etl đồng bộ từ warehouse ra crm ads email operational analytics idempotent": "reverse-etl",
        "vectorized query execution xử lý theo lô vector cột simd nhanh thay vì từng dòng": "how-vectorized-query-execution-works",
        "data mesh phi tập trung domain sở hữu dữ liệu như sản phẩm self-serve governance": "data-mesh-architecture",
    }
    for text, expected_slug in cases.items():
        skill = skills.find_matching_skill(text)
        assert skill is not None, text
        assert skill.slug == expected_slug, (text, skill.slug)


def test_find_matching_skills_multi_for_rich_request():
    from app import skills
    # Yêu cầu chi tiết dùng đúng từ vựng skill → kéo NHIỀU kỹ năng liên quan.
    q = ("dựng game platformer: tilemap bản đồ ô lưới và thiết kế màn chơi; điều khiển nhân vật "
         "character controller có coyote time đệm nhảy jump buffer; hệ thống hạt particle system "
         "làm lửa; ánh sáng 2d động; giao diện game hud thanh máu điểm số")
    got = {s.slug for s in skills.find_matching_skills(q, limit=4)}
    assert "character-controller-and-platformer-movement" in got
    assert len(got) >= 2  # ít nhất 2 kỹ năng game được nạp cùng lúc
    # Câu mơ hồ / một chủ đề → chỉ 1 kỹ năng (không nhồi nhiễu).
    assert len(skills.find_matching_skills("làm sao viết test trước khi code", limit=4)) == 1
    # Không liên quan → rỗng.
    assert skills.find_matching_skills("hôm nay trời đẹp không đi chơi đâu", limit=4) == []


def test_semantic_match_ranks_by_cosine_and_threshold(monkeypatch):
    import asyncio
    from app import skill_search, skills
    a, b = skills.all_skills()[0].slug, skills.all_skills()[1].slug
    skill_search._index.clear()
    skill_search._index[a] = {"hash": "x", "vec": [1.0, 0.0]}
    skill_search._index[b] = {"hash": "x", "vec": [0.0, 1.0]}
    skill_search._loaded = True
    monkeypatch.setattr(skill_search, "embeddings_enabled", lambda: True)
    async def qembed(_):  # câu hỏi gần vector của skill a
        return [0.9, 0.1]
    monkeypatch.setattr(skill_search, "embed_text", qembed)
    got = asyncio.run(skill_search.semantic_match("bất kỳ", limit=1, min_sim=0.5))
    assert got and got[0].slug == a
    # Ngưỡng quá cao → loại hết (không đủ giống thì không nạp bừa).
    assert asyncio.run(skill_search.semantic_match("bất kỳ", limit=1, min_sim=0.999)) == []
    # Tắt embeddings (không có key) → rỗng, lùi về từ khóa.
    monkeypatch.setattr(skill_search, "embeddings_enabled", lambda: False)
    assert asyncio.run(skill_search.semantic_match("bất kỳ")) == []
    skill_search._index.clear()
    skill_search._loaded = False



# ── Xoay vòng nhiều API key (gộp hạn mức free nhiều tài khoản) ──────────

def test_multikey_config_parses_dedups_and_falls_back(monkeypatch):
    """_multikey đọc {PREFIX}_API_KEYS (nhiều key phẩy), loại rỗng+trùng, giữ thứ tự;
    thiếu thì fallback về {PREFIX}_API_KEY đơn."""
    from app import config
    monkeypatch.setenv("XPROV_API_KEYS", " k1 , k2 ,, k1 , k3 ")
    assert config._multikey("XPROV") == ["k1", "k2", "k3"]
    monkeypatch.delenv("XPROV_API_KEYS", raising=False)
    monkeypatch.setenv("XPROV_API_KEY", "solo")
    assert config._multikey("XPROV") == ["solo"]
    monkeypatch.delenv("XPROV_API_KEY", raising=False)
    assert config._multikey("XPROV") == []


def test_keyring_round_robin_advances_start():
    """rotation_order xoay điểm bắt đầu mỗi lần gọi → cân bằng tải qua các key."""
    from app import keyring
    keyring.reset()
    assert keyring.rotation_order("p", 3) == [0, 1, 2]
    assert keyring.rotation_order("p", 3) == [1, 2, 0]
    assert keyring.rotation_order("p", 3) == [2, 0, 1]
    assert keyring.rotation_order("p", 3) == [0, 1, 2]
    # Nhà cung cấp khác có bộ đếm riêng.
    assert keyring.rotation_order("q", 2) == [0, 1]
    # Biên: 1 key → [0]; 0 key → rỗng.
    assert keyring.rotation_order("p", 1) == [0]
    assert keyring.rotation_order("p", 0) == []
    keyring.reset()


def test_engine_keys_and_availability_reflect_multikey():
    """Engine với danh sách api_keys → _keys() trả đủ, available() True; rỗng → False."""
    from app.engines.openai_compatible import GroqEngine
    e = GroqEngine()
    e.api_keys = ["a", "b", "c"]
    e.api_key = "a"
    assert e._keys() == ["a", "b", "c"]
    assert e.available() is True
    e.api_keys = []
    e.api_key = ""
    assert e._keys() == []
    assert e.available() is False
    # Chỉ có api_key đơn (không có api_keys) vẫn hoạt động (tương thích ngược).
    e.api_key = "solo"
    assert e._keys() == ["solo"]
    assert e.available() is True


def test_multikey_rotation_skips_dead_key_then_succeeds():
    """stream_chat nhảy sang key kế khi key đầu 429/hỏng (rotate=True) và CHƯA phát
    chữ nào — key thứ hai trả lời thành công."""
    import asyncio
    from app.engines.openai_compatible import OpenAICompatibleEngine
    from app.engines.base import EngineError
    from app import keyring
    from app.schemas import RouteDecision

    keyring.reset()

    class FakeEngine(OpenAICompatibleEngine):
        name = "fake-rot"
        base_url = "http://x"
        used_keys: list[str] = []

        async def _stream_once(self, payload, key):
            type(self).used_keys.append(key)
            if key == "bad":               # key hết lượt → yêu cầu xoay
                raise EngineError("429", rotate=True)
            yield {"type": "text", "text": "ok"}

    FakeEngine.used_keys = []
    e = FakeEngine()
    e.api_keys = ["bad", "good"]
    route = RouteDecision(mode="fast", label="x", model="m", use_web_search=False)

    async def run():
        return [ev async for ev in e.stream_chat([{"role": "user", "content": "hi"}], route, "sys")]

    events = asyncio.run(run())
    assert FakeEngine.used_keys == ["bad", "good"]      # thử bad trước, rồi good
    assert any(ev.get("type") == "text" and ev["text"] == "ok" for ev in events)
    assert events[-1]["type"] == "final"
    keyring.reset()


def test_multikey_rotation_does_not_retry_after_text_emitted():
    """Nếu ĐÃ phát chữ rồi mới lỗi → KHÔNG xoay key (tránh lặp nội dung), ném lỗi ra."""
    import asyncio
    from app.engines.openai_compatible import OpenAICompatibleEngine
    from app.engines.base import EngineError
    from app import keyring
    from app.schemas import RouteDecision

    keyring.reset()

    class MidFailEngine(OpenAICompatibleEngine):
        name = "fake-midfail"
        base_url = "http://x"
        tries = 0

        async def _stream_once(self, payload, key):
            type(self).tries += 1
            yield {"type": "text", "text": "phần đầu "}
            raise EngineError("đứt giữa chừng", rotate=True)

    MidFailEngine.tries = 0
    e = MidFailEngine()
    e.api_keys = ["k1", "k2"]
    route = RouteDecision(mode="fast", label="x", model="m", use_web_search=False)

    async def run():
        out = []
        async for ev in e.stream_chat([{"role": "user", "content": "hi"}], route, "sys"):
            out.append(ev)
        return out

    try:
        asyncio.run(run())
        assert False, "phải ném EngineError"
    except EngineError:
        pass
    assert MidFailEngine.tries == 1     # KHÔNG thử key thứ hai sau khi đã phát chữ
    keyring.reset()


def test_multikey_non_rotate_error_does_not_rotate():
    """Lỗi KHÔNG phải do key (rotate=False, vd rỗng nội dung) → không phí key khác."""
    import asyncio
    from app.engines.openai_compatible import OpenAICompatibleEngine
    from app.engines.base import EngineError
    from app import keyring
    from app.schemas import RouteDecision

    keyring.reset()

    class EmptyEngine(OpenAICompatibleEngine):
        name = "fake-empty"
        base_url = "http://x"
        tries = 0

        async def _stream_once(self, payload, key):
            type(self).tries += 1
            raise EngineError("rỗng", rotate=False)
            yield  # pragma: no cover

    EmptyEngine.tries = 0
    e = EmptyEngine()
    e.api_keys = ["k1", "k2", "k3"]
    route = RouteDecision(mode="fast", label="x", model="m", use_web_search=False)

    async def run():
        return [ev async for ev in e.stream_chat([{"role": "user", "content": "hi"}], route, "sys")]

    try:
        asyncio.run(run())
        assert False
    except EngineError:
        pass
    assert EmptyEngine.tries == 1       # dừng ngay, không xoay
    keyring.reset()
